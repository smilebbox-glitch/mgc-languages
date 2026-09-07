from __future__ import annotations

import hashlib
import hmac
import io
import csv
import json
import os
import random
import re
import secrets
import shutil
import subprocess
import logging
import threading
import time
import uuid
from collections import Counter, defaultdict, deque
from functools import lru_cache
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from fastapi import Cookie, Depends, FastAPI, File, HTTPException, Query, Request, Response, UploadFile
from fastapi.responses import JSONResponse, RedirectResponse, StreamingResponse, PlainTextResponse
from fastapi.staticfiles import StaticFiles
from starlette.middleware.sessions import SessionMiddleware
from starlette.middleware.trustedhost import TrustedHostMiddleware
from starlette.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field
from sqlalchemy import Boolean, DateTime, ForeignKey, Integer, String, Text, UniqueConstraint, create_engine, delete, event, func, select, text
from sqlalchemy.orm import DeclarativeBase, Mapped, Session, mapped_column, sessionmaker
from sqlalchemy.exc import SQLAlchemyError

import cmudict


# v5.7.3: configuration extracted into a dependency-light module.
# Re-exported here to preserve the historical app.<SETTING> compatibility contract.
from mgc.config import *  # noqa: F401,F403,E402

logger = logging.getLogger("mgc.languages")
if not logging.getLogger().handlers:
    logging.basicConfig(level=os.getenv("LOG_LEVEL", "INFO").upper(), format="%(message)s")

if DATABASE_URL.startswith("postgres://"):
    DATABASE_URL = DATABASE_URL.replace("postgres://", "postgresql+psycopg://", 1)
elif DATABASE_URL.startswith("postgresql://") and "+psycopg" not in DATABASE_URL:
    DATABASE_URL = DATABASE_URL.replace("postgresql://", "postgresql+psycopg://", 1)

connect_args = {"check_same_thread": False} if DATABASE_URL.startswith("sqlite") else {
    "connect_timeout": DB_CONNECT_TIMEOUT_SECONDS,
    "options": f"-c statement_timeout={DB_STATEMENT_TIMEOUT_MS}",
}
engine_kwargs: dict[str, Any] = {"pool_pre_ping": True, "connect_args": connect_args}
if not DATABASE_URL.startswith("sqlite"):
    engine_kwargs.update({
        "pool_size": DB_POOL_SIZE,
        "max_overflow": DB_MAX_OVERFLOW,
        "pool_timeout": max(1, int(os.getenv("DB_POOL_TIMEOUT", "30"))),
        "pool_recycle": max(60, int(os.getenv("DB_POOL_RECYCLE", "1800"))),
    })
engine = create_engine(DATABASE_URL, **engine_kwargs)
SessionLocal = sessionmaker(bind=engine, expire_on_commit=False)

# Query telemetry is deliberately metadata-only: no SQL text, parameters or learning content are retained.
_DB_QUERY_LOCK = threading.Lock()
_DB_QUERY_WINDOW: deque[tuple[float, float, str, str, bool]] = deque(maxlen=20000)
_DB_QUERY_TOTAL = Counter()

def _query_fingerprint(statement: str) -> tuple[str, str]:
    compact = re.sub(r"\s+", " ", (statement or "").strip())
    operation = (compact.split(" ", 1)[0] if compact else "unknown").upper()[:16]
    # Literals are normalized before hashing to reduce accidental sensitivity and cardinality.
    normalized = re.sub(r"'(?:''|[^'])*'", "?", compact)
    normalized = re.sub(r'"(?:""|[^"])*"', '"?"', normalized)
    normalized = re.sub(r"\b\d+(?:\.\d+)?\b", "?", normalized)
    return hashlib.sha256(normalized.encode("utf-8", "replace")).hexdigest()[:16], operation

@event.listens_for(engine, "before_cursor_execute")
def _db_before_cursor_execute(conn, cursor, statement, parameters, context, executemany):  # noqa: ANN001
    context._mgc_query_started = time.perf_counter()

@event.listens_for(engine, "after_cursor_execute")
def _db_after_cursor_execute(conn, cursor, statement, parameters, context, executemany):  # noqa: ANN001
    started = getattr(context, "_mgc_query_started", None)
    if started is None:
        return
    duration_ms = (time.perf_counter() - started) * 1000.0
    fingerprint, operation = _query_fingerprint(statement)
    slow = duration_ms >= DB_SLOW_QUERY_THRESHOLD_MS
    with _DB_QUERY_LOCK:
        _DB_QUERY_WINDOW.append((time.time(), duration_ms, fingerprint, operation, slow))
        _DB_QUERY_TOTAL["queries"] += 1
        if slow:
            _DB_QUERY_TOTAL["slow_queries"] += 1

def db_query_telemetry_snapshot() -> dict[str, Any]:
    cutoff = time.time() - DB_QUERY_WINDOW_MINUTES * 60
    with _DB_QUERY_LOCK:
        while _DB_QUERY_WINDOW and _DB_QUERY_WINDOW[0][0] < cutoff:
            _DB_QUERY_WINDOW.popleft()
        rows = list(_DB_QUERY_WINDOW)
        lifetime = dict(_DB_QUERY_TOTAL)
    durations = [r[1] for r in rows]
    slow_rows = [r for r in rows if r[4]]
    fingerprints: dict[str, dict[str, Any]] = {}
    for _, duration_ms, fingerprint, operation, slow in slow_rows:
        item = fingerprints.setdefault(fingerprint, {"fingerprint": fingerprint, "operation": operation, "count": 0, "max_ms": 0.0})
        item["count"] += 1
        item["max_ms"] = max(float(item["max_ms"]), duration_ms)
    top = sorted(fingerprints.values(), key=lambda x: (-int(x["count"]), -float(x["max_ms"])))[:10]
    for item in top:
        item["max_ms"] = round(float(item["max_ms"]), 2)
    return {
        "window_minutes": DB_QUERY_WINDOW_MINUTES,
        "queries": len(rows),
        "slow_queries": len(slow_rows),
        "slow_threshold_ms": DB_SLOW_QUERY_THRESHOLD_MS,
        "p50_ms": _percentile(durations, 0.50) if durations else 0.0,
        "p95_ms": _percentile(durations, 0.95) if durations else 0.0,
        "p99_ms": _percentile(durations, 0.99) if durations else 0.0,
        "max_ms": round(max(durations), 2) if durations else 0.0,
        "lifetime_queries": int(lifetime.get("queries", 0)),
        "lifetime_slow_queries": int(lifetime.get("slow_queries", 0)),
        "top_slow_fingerprints": top,
        "privacy_note": "Only hashed normalized query fingerprints and timing metadata are retained; SQL text and parameters are not stored.",
    }

def db_pool_snapshot() -> dict[str, Any]:
    if DATABASE_URL.startswith("sqlite"):
        return {"driver": engine.dialect.name, "available": False, "reason": "pool saturation is a PostgreSQL pilot metric"}
    def val(name: str) -> Any:
        attr = getattr(engine.pool, name, None)
        try:
            return attr() if callable(attr) else attr
        except Exception:
            return None
    size = val("size")
    checked = val("checkedout")
    overflow = val("overflow")
    capacity = DB_POOL_SIZE + DB_MAX_OVERFLOW
    saturation = round((float(checked or 0) / capacity) * 100.0, 2) if capacity else 0.0
    return {"driver": engine.dialect.name, "available": True, "size": size, "max_overflow": DB_MAX_OVERFLOW, "checked_out": checked, "overflow": overflow, "capacity": capacity, "saturation_percent": saturation, "alert_percent": DB_POOL_ALERT_PERCENT}


class Base(DeclarativeBase):
    pass


class User(Base):
    __tablename__ = "users"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    username: Mapped[str] = mapped_column(String(80), unique=True, index=True)
    display_name: Mapped[str] = mapped_column(String(120))
    password_hash: Mapped[str] = mapped_column(String(300))
    role: Mapped[str] = mapped_column(String(30), default="user")
    department: Mapped[str] = mapped_column(String(160), default="General", index=True)
    preferred_language: Mapped[str] = mapped_column(String(10), default="chinese")
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class LoginSession(Base):
    __tablename__ = "login_sessions"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    token_hash: Mapped[str] = mapped_column(String(64), unique=True, index=True)
    user_id: Mapped[int] = mapped_column(ForeignKey("users.id", ondelete="CASCADE"), index=True)
    expires_at: Mapped[datetime] = mapped_column(DateTime(timezone=True))


class TermProgress(Base):
    __tablename__ = "term_progress"
    __table_args__ = (UniqueConstraint("user_id", "language", "term_id", name="uq_progress_language_term"),)
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    user_id: Mapped[int] = mapped_column(ForeignKey("users.id", ondelete="CASCADE"), index=True)
    language: Mapped[str] = mapped_column(String(10), index=True)
    term_id: Mapped[str] = mapped_column(String(80), index=True)
    status: Mapped[str] = mapped_column(String(20), default="learning")
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class ExamResult(Base):
    __tablename__ = "exam_results"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    user_id: Mapped[int] = mapped_column(ForeignKey("users.id", ondelete="CASCADE"), index=True)
    language: Mapped[str] = mapped_column(String(10), index=True)
    score: Mapped[int] = mapped_column(Integer)
    total: Mapped[int] = mapped_column(Integer)
    details: Mapped[str] = mapped_column(Text, default="{}")
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class CourseDayResult(Base):
    __tablename__ = "course_day_results"
    __table_args__ = (UniqueConstraint("user_id", "language", "day", name="uq_course_day_result"),)
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    user_id: Mapped[int] = mapped_column(ForeignKey("users.id", ondelete="CASCADE"), index=True)
    language: Mapped[str] = mapped_column(String(10), index=True)
    day: Mapped[int] = mapped_column(Integer)
    score: Mapped[int] = mapped_column(Integer)
    total: Mapped[int] = mapped_column(Integer, default=5)
    completed_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class GamificationProfile(Base):
    __tablename__ = "gamification_profiles"
    user_id: Mapped[int] = mapped_column(ForeignKey("users.id", ondelete="CASCADE"), primary_key=True)
    lifetime_xp: Mapped[int] = mapped_column(Integer, default=0)
    spendable_xp: Mapped[int] = mapped_column(Integer, default=0)
    weekly_xp: Mapped[int] = mapped_column(Integer, default=0)
    week_key: Mapped[str] = mapped_column(String(12), default="")
    last_activity_at: Mapped[datetime | None] = mapped_column(DateTime(timezone=True), nullable=True)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class XPEvent(Base):
    __tablename__ = "xp_events"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    user_id: Mapped[int] = mapped_column(ForeignKey("users.id", ondelete="CASCADE"), index=True)
    event_type: Mapped[str] = mapped_column(String(60), index=True)
    source_id: Mapped[str] = mapped_column(String(160), default="", index=True)
    language: Mapped[str] = mapped_column(String(10), default="")
    topic: Mapped[str] = mapped_column(String(160), default="")
    xp_delta: Mapped[int] = mapped_column(Integer)
    lifetime_delta: Mapped[int] = mapped_column(Integer, default=0)
    raw_xp: Mapped[int] = mapped_column(Integer, default=0)
    multiplier_pct: Mapped[int] = mapped_column(Integer, default=100)
    idempotency_key: Mapped[str] = mapped_column(String(180), unique=True, index=True)
    metadata_json: Mapped[str] = mapped_column(Text, default="{}")
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc), index=True)


class PracticeResult(Base):
    __tablename__ = "practice_results"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    user_id: Mapped[int] = mapped_column(ForeignKey("users.id", ondelete="CASCADE"), index=True)
    session_id: Mapped[str] = mapped_column(String(120), unique=True, index=True)
    kind: Mapped[str] = mapped_column(String(40), index=True)
    language: Mapped[str] = mapped_column(String(10), index=True)
    topic: Mapped[str] = mapped_column(String(160), default="")
    score: Mapped[int] = mapped_column(Integer)
    total: Mapped[int] = mapped_column(Integer)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc), index=True)


class GameSession(Base):
    __tablename__ = "game_sessions"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    public_id: Mapped[str] = mapped_column(String(80), unique=True, index=True)
    user_id: Mapped[int] = mapped_column(ForeignKey("users.id", ondelete="CASCADE"), index=True)
    game_type: Mapped[str] = mapped_column(String(40), index=True)
    language: Mapped[str] = mapped_column(String(10), index=True)
    topic: Mapped[str] = mapped_column(String(160), default="")
    payload_json: Mapped[str] = mapped_column(Text, default="{}")
    status: Mapped[str] = mapped_column(String(20), default="open")
    score: Mapped[int] = mapped_column(Integer, default=0)
    total: Mapped[int] = mapped_column(Integer, default=0)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    completed_at: Mapped[datetime | None] = mapped_column(DateTime(timezone=True), nullable=True)


class CustomTerm(Base):
    __tablename__ = "custom_terms"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    public_id: Mapped[str] = mapped_column(String(80), unique=True, index=True)
    language: Mapped[str] = mapped_column(String(10), index=True)
    shop: Mapped[str] = mapped_column(String(160), default="")
    topic: Mapped[str] = mapped_column(String(160), index=True)
    subtopic: Mapped[str] = mapped_column(String(160), default="")
    level: Mapped[str] = mapped_column(String(10), default="A1")
    term: Mapped[str] = mapped_column(String(500))
    pronunciation: Mapped[str] = mapped_column(String(500), default="")
    reading: Mapped[str] = mapped_column(String(500), default="")
    translation: Mapped[str] = mapped_column(String(500))
    example: Mapped[str] = mapped_column(Text, default="")
    example_translation: Mapped[str] = mapped_column(Text, default="")
    tags: Mapped[str] = mapped_column(Text, default="")
    source_type: Mapped[str] = mapped_column(String(60), default="manual", index=True)
    source_ref: Mapped[str] = mapped_column(Text, default="")
    status: Mapped[str] = mapped_column(String(20), default="draft", index=True)
    created_by: Mapped[int] = mapped_column(ForeignKey("users.id", ondelete="SET NULL"), nullable=True)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class TermRevision(Base):
    __tablename__ = "term_revisions"
    __table_args__ = (UniqueConstraint("term_id", "revision_no", name="uq_term_revision_no"),)
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    term_id: Mapped[int] = mapped_column(ForeignKey("custom_terms.id", ondelete="CASCADE"), index=True)
    revision_no: Mapped[int] = mapped_column(Integer)
    action: Mapped[str] = mapped_column(String(30), default="update")
    snapshot_json: Mapped[str] = mapped_column(Text)
    actor_user_id: Mapped[int | None] = mapped_column(ForeignKey("users.id", ondelete="SET NULL"), nullable=True)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc), index=True)


class TermReview(Base):
    __tablename__ = "term_reviews"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    term_id: Mapped[int] = mapped_column(ForeignKey("custom_terms.id", ondelete="CASCADE"), index=True)
    revision_no: Mapped[int] = mapped_column(Integer)
    reviewer_user_id: Mapped[int | None] = mapped_column(ForeignKey("users.id", ondelete="SET NULL"), nullable=True)
    decision: Mapped[str] = mapped_column(String(30), index=True)
    note: Mapped[str] = mapped_column(Text, default="")
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc), index=True)


class SRSCard(Base):
    __tablename__ = "srs_cards"
    __table_args__ = (UniqueConstraint("user_id", "language", "term_id", name="uq_srs_user_language_term"),)
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    user_id: Mapped[int] = mapped_column(ForeignKey("users.id", ondelete="CASCADE"), index=True)
    language: Mapped[str] = mapped_column(String(10), index=True)
    term_id: Mapped[str] = mapped_column(String(80), index=True)
    topic: Mapped[str] = mapped_column(String(160), default="", index=True)
    repetitions: Mapped[int] = mapped_column(Integer, default=0)
    lapses: Mapped[int] = mapped_column(Integer, default=0)
    ease_factor_pct: Mapped[int] = mapped_column(Integer, default=250)
    interval_days: Mapped[int] = mapped_column(Integer, default=0)
    due_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc), index=True)
    last_quality: Mapped[int] = mapped_column(Integer, default=0)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class QuestionAttempt(Base):
    __tablename__ = "question_attempts"
    __table_args__ = (UniqueConstraint("user_id", "session_id", "question_id", name="uq_question_attempt_session"),)
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    user_id: Mapped[int] = mapped_column(ForeignKey("users.id", ondelete="CASCADE"), index=True)
    session_id: Mapped[str] = mapped_column(String(120), default="", index=True)
    question_id: Mapped[str] = mapped_column(String(160), index=True)
    term_id: Mapped[str] = mapped_column(String(80), default="", index=True)
    language: Mapped[str] = mapped_column(String(10), index=True)
    topic: Mapped[str] = mapped_column(String(160), default="", index=True)
    kind: Mapped[str] = mapped_column(String(50), default="quiz", index=True)
    correct: Mapped[bool] = mapped_column(Boolean, default=False, index=True)
    response_ms: Mapped[int] = mapped_column(Integer, default=0)
    selected_hash: Mapped[str] = mapped_column(String(64), default="")
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc), index=True)


class LearningPreference(Base):
    __tablename__ = "learning_preferences"
    user_id: Mapped[int] = mapped_column(ForeignKey("users.id", ondelete="CASCADE"), primary_key=True)
    show_pinyin: Mapped[bool] = mapped_column(Boolean, default=True)
    show_reading: Mapped[bool] = mapped_column(Boolean, default=True)
    server_audio_enabled: Mapped[bool] = mapped_column(Boolean, default=True)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class NotificationPreference(Base):
    __tablename__ = "notification_preferences"
    user_id: Mapped[int] = mapped_column(ForeignKey("users.id", ondelete="CASCADE"), primary_key=True)
    mode: Mapped[str] = mapped_column(String(20), default="normal")
    window_start: Mapped[str] = mapped_column(String(5), default="09:00")
    window_end: Mapped[str] = mapped_column(String(5), default="19:00")
    browser_enabled: Mapped[bool] = mapped_column(Boolean, default=False)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class LearningNudge(Base):
    __tablename__ = "learning_nudges"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    user_id: Mapped[int] = mapped_column(ForeignKey("users.id", ondelete="CASCADE"), index=True)
    title: Mapped[str] = mapped_column(String(200))
    body: Mapped[str] = mapped_column(Text)
    target_view: Mapped[str] = mapped_column(String(40), default="home")
    reason: Mapped[str] = mapped_column(String(80), default="review")
    read: Mapped[bool] = mapped_column(Boolean, default=False, index=True)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc), index=True)


class AuditLog(Base):
    __tablename__ = "audit_logs"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    event_type: Mapped[str] = mapped_column(String(100), index=True)
    actor_user_id: Mapped[int | None] = mapped_column(ForeignKey("users.id", ondelete="SET NULL"), nullable=True, index=True)
    target_type: Mapped[str] = mapped_column(String(80), default="")
    target_id: Mapped[str] = mapped_column(String(180), default="")
    request_id: Mapped[str] = mapped_column(String(80), default="", index=True)
    source_ip_hash: Mapped[str] = mapped_column(String(64), default="")
    metadata_json: Mapped[str] = mapped_column(Text, default="{}")
    previous_hash: Mapped[str] = mapped_column(String(64), default="")
    event_hash: Mapped[str] = mapped_column(String(64), default="", index=True)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc), index=True)


class AuditAnchor(Base):
    __tablename__ = "audit_anchors"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, default=1)
    last_deleted_id: Mapped[int] = mapped_column(Integer, default=0)
    last_deleted_hash: Mapped[str] = mapped_column(String(64), default="")
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class OperationalEvent(Base):
    __tablename__ = "operational_events"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    event_type: Mapped[str] = mapped_column(String(100), index=True)
    severity: Mapped[str] = mapped_column(String(20), default="warning", index=True)
    component: Mapped[str] = mapped_column(String(60), default="app", index=True)
    status_code: Mapped[int] = mapped_column(Integer, default=0)
    request_id: Mapped[str] = mapped_column(String(80), default="", index=True)
    path: Mapped[str] = mapped_column(String(240), default="")
    detail: Mapped[str] = mapped_column(String(500), default="")
    metadata_json: Mapped[str] = mapped_column(Text, default="{}")
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc), index=True)


class PilotAlert(Base):
    __tablename__ = "pilot_alerts"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    alert_key: Mapped[str] = mapped_column(String(120), unique=True, index=True)
    severity: Mapped[str] = mapped_column(String(20), default="warning", index=True)
    status: Mapped[str] = mapped_column(String(20), default="open", index=True)
    component: Mapped[str] = mapped_column(String(60), default="app", index=True)
    title: Mapped[str] = mapped_column(String(220), default="")
    detail: Mapped[str] = mapped_column(String(700), default="")
    metadata_json: Mapped[str] = mapped_column(Text, default="{}")
    first_seen_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc), index=True)
    last_seen_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc), index=True)
    acknowledged_by: Mapped[int | None] = mapped_column(ForeignKey("users.id", ondelete="SET NULL"), nullable=True)
    acknowledged_at: Mapped[datetime | None] = mapped_column(DateTime(timezone=True), nullable=True)
    resolved_at: Mapped[datetime | None] = mapped_column(DateTime(timezone=True), nullable=True)


class PilotGroup(Base):
    __tablename__ = "pilot_groups"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    name: Mapped[str] = mapped_column(String(160), unique=True, index=True)
    department: Mapped[str] = mapped_column(String(160), default="General", index=True)
    description: Mapped[str] = mapped_column(Text, default="")
    status: Mapped[str] = mapped_column(String(20), default="draft", index=True)
    wave: Mapped[int] = mapped_column(Integer, default=1, index=True)
    starts_at: Mapped[datetime | None] = mapped_column(DateTime(timezone=True), nullable=True)
    ends_at: Mapped[datetime | None] = mapped_column(DateTime(timezone=True), nullable=True)
    created_by: Mapped[int | None] = mapped_column(ForeignKey("users.id", ondelete="SET NULL"), nullable=True)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class PilotGroupMember(Base):
    __tablename__ = "pilot_group_members"
    __table_args__ = (UniqueConstraint("group_id", "user_id", name="uq_pilot_group_member"),)
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    group_id: Mapped[int] = mapped_column(ForeignKey("pilot_groups.id", ondelete="CASCADE"), index=True)
    user_id: Mapped[int] = mapped_column(ForeignKey("users.id", ondelete="CASCADE"), index=True)
    assigned_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class FeatureFlag(Base):
    __tablename__ = "feature_flags"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    flag_key: Mapped[str] = mapped_column(String(80), unique=True, index=True)
    title: Mapped[str] = mapped_column(String(160), default="")
    description: Mapped[str] = mapped_column(Text, default="")
    default_enabled: Mapped[bool] = mapped_column(Boolean, default=True)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class PilotGroupFeature(Base):
    __tablename__ = "pilot_group_features"
    __table_args__ = (UniqueConstraint("group_id", "flag_key", name="uq_pilot_group_feature"),)
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    group_id: Mapped[int] = mapped_column(ForeignKey("pilot_groups.id", ondelete="CASCADE"), index=True)
    flag_key: Mapped[str] = mapped_column(String(80), index=True)
    enabled: Mapped[bool] = mapped_column(Boolean, default=True)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class PilotTrackAssignment(Base):
    __tablename__ = "pilot_track_assignments"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    group_id: Mapped[int] = mapped_column(ForeignKey("pilot_groups.id", ondelete="CASCADE"), index=True)
    language: Mapped[str] = mapped_column(String(10), index=True)
    track_name: Mapped[str] = mapped_column(String(180))
    topic: Mapped[str] = mapped_column(String(180), default="")
    target_level: Mapped[str] = mapped_column(String(10), default="A1")
    due_date: Mapped[str] = mapped_column(String(10), default="")
    created_by: Mapped[int | None] = mapped_column(ForeignKey("users.id", ondelete="SET NULL"), nullable=True)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class PilotDailyUsage(Base):
    __tablename__ = "pilot_daily_usage"
    __table_args__ = (UniqueConstraint("user_id", "day_key", name="uq_pilot_daily_usage"),)
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    user_id: Mapped[int] = mapped_column(ForeignKey("users.id", ondelete="CASCADE"), index=True)
    day_key: Mapped[str] = mapped_column(String(10), index=True)
    xp_awarded: Mapped[int] = mapped_column(Integer, default=0)
    game_starts: Mapped[int] = mapped_column(Integer, default=0)
    tts_requests: Mapped[int] = mapped_column(Integer, default=0)
    practice_submissions: Mapped[int] = mapped_column(Integer, default=0)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


if AUTO_CREATE_SCHEMA:
    Base.metadata.create_all(bind=engine)


def _pilot_day_key() -> str:
    return date.today().isoformat()


def pilot_daily_usage(db: Session, user_id: int) -> PilotDailyUsage:
    key = _pilot_day_key()
    row = db.scalar(select(PilotDailyUsage).where(PilotDailyUsage.user_id == user_id, PilotDailyUsage.day_key == key))
    if not row:
        row = PilotDailyUsage(user_id=user_id, day_key=key)
        db.add(row)
        db.flush()
    return row


def _consume_daily_quota(db: Session, user_id: int, field: str, limit: int, amount: int = 1) -> PilotDailyUsage:
    row = pilot_daily_usage(db, user_id)
    current = int(getattr(row, field, 0) or 0)
    if current + amount > limit:
        raise HTTPException(429, "Дневной лимит пилота достигнут. Продолжить можно завтра.", headers={"Retry-After":"3600"})
    setattr(row, field, current + amount)
    row.updated_at = datetime.now(timezone.utc)
    return row


def _group_active(group: PilotGroup, now: datetime | None = None) -> bool:
    now = now or datetime.now(timezone.utc)
    if group.status != "active":
        return False
    start = group.starts_at
    end = group.ends_at
    if start and start.tzinfo is None: start = start.replace(tzinfo=timezone.utc)
    if end and end.tzinfo is None: end = end.replace(tzinfo=timezone.utc)
    return not (start and now < start) and not (end and now > end)


def user_pilot_groups(db: Session, user_id: int, active_only: bool = False) -> list[PilotGroup]:
    rows = db.scalars(select(PilotGroup).join(PilotGroupMember, PilotGroupMember.group_id == PilotGroup.id).where(PilotGroupMember.user_id == user_id).order_by(PilotGroup.wave, PilotGroup.name)).all()
    return [g for g in rows if _group_active(g)] if active_only else rows


def feature_flag_catalog(db: Session) -> dict[str, dict[str, Any]]:
    catalog = {k: dict(v) for k, v in BUILTIN_FEATURE_FLAGS.items()}
    for row in db.scalars(select(FeatureFlag)).all():
        catalog[row.flag_key] = {"title":row.title,"description":row.description,"default_enabled":bool(row.default_enabled)}
    return catalog


def user_feature_flags(db: Session, user_id: int) -> dict[str, bool]:
    catalog = feature_flag_catalog(db)
    flags = {k: bool(v.get("default_enabled", True)) for k, v in catalog.items()}
    groups = user_pilot_groups(db, user_id, active_only=True)
    if groups:
        group_ids = [g.id for g in groups]
        overrides = db.scalars(select(PilotGroupFeature).where(PilotGroupFeature.group_id.in_(group_ids))).all()
        by_flag: dict[str, list[bool]] = defaultdict(list)
        for row in overrides: by_flag[row.flag_key].append(bool(row.enabled))
        for key, values in by_flag.items():
            # Conservative rollout rule: an explicit disable in any active group wins.
            # This prevents overlapping pilot memberships from accidentally widening access.
            flags[key] = all(values)
    return flags


def require_feature(db: Session, user_id: int, flag_key: str) -> None:
    if not user_feature_flags(db, user_id).get(flag_key, True):
        raise HTTPException(403, "Функция пока не включена для вашей волны пилота")


def _pilot_usage_view(db: Session, user_id: int) -> dict[str, Any]:
    row = db.scalar(select(PilotDailyUsage).where(PilotDailyUsage.user_id == user_id, PilotDailyUsage.day_key == _pilot_day_key()))
    used = {
        "xp_awarded": int(row.xp_awarded or 0) if row else 0,
        "game_starts": int(row.game_starts or 0) if row else 0,
        "tts_requests": int(row.tts_requests or 0) if row else 0,
        "practice_submissions": int(row.practice_submissions or 0) if row else 0,
    }
    return {"day":_pilot_day_key(),"used":used,"remaining":{
        "xp":max(0,PILOT_DAILY_XP_CAP-used["xp_awarded"]),
        "games":max(0,PILOT_DAILY_GAME_START_CAP-used["game_starts"]),
        "tts":max(0,PILOT_DAILY_TTS_CAP-used["tts_requests"]),
        "practice":max(0,PILOT_DAILY_PRACTICE_CAP-used["practice_submissions"]),
    }}


def pilot_me_view(db: Session, user: User) -> dict[str, Any]:
    groups = user_pilot_groups(db, user.id)
    active_ids = [g.id for g in groups if _group_active(g)]
    assignments = []
    if active_ids:
        assignments = db.scalars(select(PilotTrackAssignment).where(PilotTrackAssignment.group_id.in_(active_ids)).order_by(PilotTrackAssignment.created_at.desc())).all()
    return {
        "groups":[{"id":g.id,"name":g.name,"department":g.department,"status":g.status,"wave":g.wave,"active":_group_active(g)} for g in groups],
        "wave":min((g.wave for g in groups if _group_active(g)), default=None),
        "features":user_feature_flags(db, user.id),
        "assignments":[{"id":a.id,"language":a.language,"track_name":a.track_name,"topic":a.topic,"target_level":a.target_level,"due_date":a.due_date} for a in assignments],
        "quotas":{"daily_xp":PILOT_DAILY_XP_CAP,"daily_games":PILOT_DAILY_GAME_START_CAP,"daily_tts":PILOT_DAILY_TTS_CAP,"daily_practice":PILOT_DAILY_PRACTICE_CAP},
        "usage_today": _pilot_usage_view(db, user.id),
        "chinese_standard":{"name":"Путунхуа (普通话)","label":"стандартный китайский","primary_learning_track":True,"dialects_reference_only":True,
                            "message":"Вы изучаете Путунхуа (普通话) — стандартный китайский. Диалекты в разделе «Информация о китайском» даны только для понимания реальной языковой среды."},
    }


def make_password_hash(password: str) -> str:
    iterations = 260_000
    salt = secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode(), salt, iterations)
    return "$".join([str(iterations), salt.hex(), digest.hex()])


def verify_password(password: str, stored: str) -> bool:
    try:
        iterations_s, salt_s, digest_s = stored.split("$", 2)
        digest = hashlib.pbkdf2_hmac("sha256", password.encode(), bytes.fromhex(salt_s), int(iterations_s))
        return hmac.compare_digest(digest.hex(), digest_s)
    except (ValueError, TypeError):
        return False


def load_json(name: str) -> dict[str, Any]:
    with (DATA_DIR / name).open("r", encoding="utf-8") as fh:
        return json.load(fh)


CHINESE_RAW = load_json("chinese.json")
ENGLISH_RAW = load_json("english.json")
APP_CONTENT = load_json("app_content.json")
EXPERIENCE = load_json("experience.json")
CHINESE_FOUNDATIONS = load_json("chinese_foundations.json")
LEVELS = ["A1", "A2", "B1", "B2", "C1"]
LANGUAGES = {"english", "chinese"}
TOPIC_ROWS = EXPERIENCE["topics"]
ROLEPLAYS = APP_CONTENT["roleplays"] + EXPERIENCE["added_roleplays"]


def topic_for(*parts: str) -> str:
    value = " ".join(part or "" for part in parts).lower()
    rules = [
        ("Сварка кузова", ["сварка кузова", "welding", "weld seam", "spot weld", "焊"]),
        ("Окраска", ["покраск", "окраск", "paint", "coating", "катафор", "漆", "涂装"]),
        ("Штамповка", ["штампов", "stamping", "press shop", "冲压"]),
        ("3D-печать", ["3d-печать", "3d печать", "3d print", "additive", "增材", "打印"]),
        ("Литьё пластмасс", ["литьё пластмасс", "литье пластмасс", "injection molding", "injection moulding", "注塑"]),
        ("Оснастка и пресс-формы", ["tool shop", "tooling", "пресс-форм", "fixture", "jig", "模具", "工装"]),
        ("Сиденья и интерьер", ["сидень", "seat", "interior trim", "座椅"]),
        ("Пассивная безопасность", ["пассивн", "passive safety", "airbag", "seat belt", "подушк", "ремень безопас", "安全气囊"]),
        ("Двигатель и трансмиссия", ["двигатель", "трансмис", "powertrain", "engine", "gearbox", "transmission"]),
        ("Электрика, ПО и ADAS", ["ecu", "embedded", "diagnostic", "ev /", "battery", "adas", "software", "functional safety", "cybersecurity", "электр", "电池"]),
        ("IT, AI и автоматизация", ["digital", "it и автомат", "it /", "ai", "mes", "data", "цифров"]),
        ("Гарантия и сервис", ["warranty", "after-sales", "гарант", "service claim"]),
        ("Качество и APQP", ["качест", "quality", "apqp", "ppap", "iatf", "problem solving", "8d", "fmea", "spc"]),
        ("R&D и инженерия", ["r&d", "engineering", "испыт", "разработ", "validation"]),
        ("Проекты и запуск", ["project", "launch", "sop", "change management"]),
        ("Закупки и поставщики", ["закуп", "поставщик", "procurement", "commercial", "contract", "legal", "localization", "trade", "tax"]),
        ("Логистика JIT/JIS", ["логист", "supply chain", "вэд", "jit", "jis", "customs"]),
        ("Безопасность и омологация", ["безопас", "охрана труда", "homologation", "сертификац"]),
        ("Переговоры и культура", ["переговор", "деловой", "культур", "business culture", "живой китайский"]),
        ("Финансы и JV", ["финанс", "finance", "financial", "economics", "corporate", "jv", "m&a", "plant management", "sales", "hr", "руководител", "capex", "budget"]),
        ("Производство и сборка", ["производ", "сборк", "assembly", "production", "industrial engineering", "lean"]),
    ]
    for topic, needles in rules:
        if any(needle in value for needle in needles):
            return topic
    return "Производство и сборка"


_CMU = cmudict.dict()
_ARPA_IPA = {
    "AA":"ɑ","AE":"æ","AH":"ʌ","AO":"ɔ","AW":"aʊ","AY":"aɪ","EH":"ɛ","ER":"ɝ","EY":"eɪ",
    "IH":"ɪ","IY":"i","OW":"oʊ","OY":"ɔɪ","UH":"ʊ","UW":"u",
    "B":"b","CH":"tʃ","D":"d","DH":"ð","F":"f","G":"ɡ","HH":"h","JH":"dʒ","K":"k","L":"l",
    "M":"m","N":"n","NG":"ŋ","P":"p","R":"r","S":"s","SH":"ʃ","T":"t","TH":"θ","V":"v","W":"w",
    "Y":"j","Z":"z","ZH":"ʒ"
}
_ARPA_RU = {
    "AA":"а","AE":"э","AH":"а","AO":"о","AW":"ау","AY":"ай","EH":"э","ER":"эр","EY":"эй",
    "IH":"и","IY":"и","OW":"оу","OY":"ой","UH":"у","UW":"у",
    "B":"б","CH":"ч","D":"д","DH":"з","F":"ф","G":"г","HH":"х","JH":"дж","K":"к","L":"л",
    "M":"м","N":"н","NG":"нг","P":"п","R":"р","S":"с","SH":"ш","T":"т","TH":"с","V":"в","W":"у",
    "Y":"й","Z":"з","ZH":"ж"
}
_EN_LETTER_RU = {
    "A":"эй","B":"би","C":"си","D":"ди","E":"и","F":"эф","G":"джи","H":"эйч","I":"ай",
    "J":"джей","K":"кей","L":"эл","M":"эм","N":"эн","O":"оу","P":"пи","Q":"кью","R":"ар",
    "S":"эс","T":"ти","U":"ю","V":"ви","W":"дабл-ю","X":"экс","Y":"уай","Z":"зи"
}
_EN_LETTER_IPA = {
    "A":"eɪ","B":"biː","C":"siː","D":"diː","E":"iː","F":"ɛf","G":"dʒiː","H":"eɪtʃ","I":"aɪ",
    "J":"dʒeɪ","K":"keɪ","L":"ɛl","M":"ɛm","N":"ɛn","O":"oʊ","P":"piː","Q":"kjuː","R":"ɑːr",
    "S":"ɛs","T":"tiː","U":"juː","V":"viː","W":"ˈdʌbəljuː","X":"ɛks","Y":"waɪ","Z":"ziː"
}
_TONE_CHAR_MAP = str.maketrans({
    "ā":"a","á":"a","ǎ":"a","à":"a","ē":"e","é":"e","ě":"e","è":"e",
    "ī":"i","í":"i","ǐ":"i","ì":"i","ō":"o","ó":"o","ǒ":"o","ò":"o",
    "ū":"u","ú":"u","ǔ":"u","ù":"u","ǖ":"ü","ǘ":"ü","ǚ":"ü","ǜ":"ü",
    "Ā":"A","Á":"A","Ǎ":"A","À":"A","Ē":"E","É":"E","Ě":"E","È":"E",
    "Ī":"I","Í":"I","Ǐ":"I","Ì":"I","Ō":"O","Ó":"O","Ǒ":"O","Ò":"O",
    "Ū":"U","Ú":"U","Ǔ":"U","Ù":"U","Ǖ":"Ü","Ǘ":"Ü","Ǚ":"Ü","Ǜ":"Ü",
})
_PINYIN_INITIALS = {
    "zh":"чж","ch":"ч","sh":"ш","b":"б","p":"п","m":"м","f":"ф","d":"д","t":"т","n":"н","l":"л",
    "g":"г","k":"к","h":"х","j":"цз","q":"ц","x":"с","r":"ж","z":"цз","c":"ц","s":"с","y":"й","w":"в"
}
_PINYIN_FINALS = {
    "iang":"ян","iong":"юн","uang":"уан","ang":"ан","eng":"эн","ong":"ун","iao":"яо","ian":"ян",
    "uan":"уань","uai":"уай","uei":"уэй","ui":"уэй","iou":"ю","iu":"ю","in":"инь","ing":"ин","un":"унь",
    "ai":"ай","ei":"эй","ao":"ао","ou":"оу","an":"ань","en":"энь","er":"эр","ia":"я","ie":"е","ua":"уа","uo":"о",
    "ue":"юэ","üe":"юэ","üan":"юань","ün":"юнь","a":"а","o":"о","e":"э","i":"и","u":"у","ü":"юй"
}


def _arpa_parts(phoneme: str) -> tuple[str, str]:
    match = re.match(r"^([A-Z]+)([012]?)$", phoneme)
    return (match.group(1), match.group(2)) if match else (phoneme, "")


@lru_cache(maxsize=8192)
def english_pronunciation(text_value: str) -> dict[str, str]:
    words = re.findall(r"[A-Za-z]+(?:'[A-Za-z]+)?|[0-9]+", text_value or "")
    if not words:
        return {"ipa":"", "reading":""}
    ipa_words: list[str] = []
    ru_words: list[str] = []
    for word in words:
        variants = _CMU.get(word.lower())
        if not variants:
            if word.isupper() and 2 <= len(word) <= 8 and word.isalpha():
                ipa_words.append(" ".join(_EN_LETTER_IPA.get(ch, ch.lower()) for ch in word))
                ru_words.append("-".join(_EN_LETTER_RU.get(ch, ch.lower()) for ch in word))
            else:
                ipa_words.append(word.lower())
                ru_words.append(word.lower())
            continue
        ipa_parts: list[str] = []
        ru_parts: list[str] = []
        for phone in variants[0]:
            base, stress = _arpa_parts(phone)
            ipa_piece = _ARPA_IPA.get(base, base.lower())
            if base == "AH" and stress == "0": ipa_piece = "ə"
            if base == "ER" and stress == "0": ipa_piece = "ɚ"
            if stress == "1": ipa_piece = "ˈ" + ipa_piece
            elif stress == "2": ipa_piece = "ˌ" + ipa_piece
            ipa_parts.append(ipa_piece)
            ru_parts.append(_ARPA_RU.get(base, base.lower()))
        ipa_words.append("".join(ipa_parts))
        ru_words.append("".join(ru_parts))
    return {"ipa":"/" + " ".join(ipa_words) + "/", "reading":"-".join(ru_words)}


_PINYIN_SYLLABLES = sorted(
    set(_PINYIN_FINALS) | {initial + final for initial in _PINYIN_INITIALS for final in _PINYIN_FINALS},
    key=len, reverse=True,
)


def _segment_pinyin_token(token: str) -> list[str]:
    pieces: list[str] = []
    rest = token
    while rest:
        match = next((candidate for candidate in _PINYIN_SYLLABLES if rest.startswith(candidate)), None)
        if not match:
            # Keep one character instead of losing information; this is only an approximate Russian aid.
            pieces.append(rest[0])
            rest = rest[1:]
        else:
            pieces.append(match)
            rest = rest[len(match):]
    return pieces


def pinyin_to_ru_approx(value: str) -> str:
    stripped = (value or "").translate(_TONE_CHAR_MAP).replace("v", "ü").lower()
    raw_tokens = re.findall(r"[a-zü]+", stripped)
    syllables: list[str] = []
    for token in raw_tokens:
        syllables.extend(_segment_pinyin_token(token))
    result: list[str] = []
    for syllable in syllables:
        initial = ""
        for candidate in sorted(_PINYIN_INITIALS, key=len, reverse=True):
            if syllable.startswith(candidate):
                initial = candidate
                break
        final = syllable[len(initial):]
        if initial in {"j","q","x","y"} and final.startswith("u"):
            final = "ü" + final[1:]
        result.append(_PINYIN_INITIALS.get(initial, "") + _PINYIN_FINALS.get(final, final))
    return "-".join(result)


def normalize_chinese(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": f"zh-{row['id']}",
        "topic": topic_for(row.get("category", ""), row.get("subcategory", ""), row.get("ru", ""), row.get("note", "")),
        "category": row.get("category", ""),
        "subcategory": row.get("subcategory", ""),
        "term": row.get("zh", ""),
        "pronunciation": row.get("pinyin", ""),
        "reading": row.get("ru_read", ""),
        "translation": row.get("ru", ""),
        "level": row.get("level", "B1"),
        "example": row.get("example_zh") or row.get("zh", ""),
        "example_pronunciation": row.get("example_pinyin", ""),
        "example_translation": row.get("example_ru") or row.get("ru", ""),
        "note": row.get("note", ""),
    }


def normalize_english(row: dict[str, Any]) -> dict[str, Any]:
    speech = english_pronunciation(row["term"])
    example = row.get("example_target", "")
    return {
        "id": row["id"],
        "topic": topic_for(row.get("topic", ""), row.get("category", ""), row.get("subcategory", ""), row.get("term", ""), row.get("ru", "")),
        "category": row.get("category", row["topic"]),
        "subcategory": row.get("subcategory", row["topic"]),
        "term": row["term"],
        "pronunciation": speech["ipa"],
        "reading": speech["reading"],
        "translation": row["ru"],
        "level": row["level"],
        "example": example,
        "example_pronunciation": english_pronunciation(example)["ipa"] if example else "",
        "example_translation": row.get("example_ru", ""),
        "note": "",
    }


def normalize_extra(row: dict[str, Any], language: str) -> dict[str, Any]:
    if language == "english":
        return {
            "id": f"en-{row['id']}", "topic": row["topic"], "category": row["topic"],
            "subcategory": row["topic"], "term": row["en"], "pronunciation": english_pronunciation(row["en"])["ipa"], "reading": english_pronunciation(row["en"])["reading"],
            "translation": row["en_ru"], "level": row["level"], "example": row["en"],
            "example_pronunciation": english_pronunciation(row["en"])["ipa"], "example_translation": row["en_ru"], "note": "",
        }
    return {
        "id": f"zh-{row['id']}", "topic": row["topic"], "category": row["topic"],
        "subcategory": row["topic"], "term": row["zh"], "pronunciation": row["pinyin"], "reading": pinyin_to_ru_approx(row["pinyin"]),
        "translation": row["zh_ru"], "level": row["level"], "example": row["zh"],
        "example_pronunciation": row["pinyin"], "example_translation": row["zh_ru"], "note": "",
    }


TERMS = {
    "chinese": [normalize_chinese(row) for row in CHINESE_RAW["items"]] + [normalize_extra(row, "chinese") for row in EXPERIENCE["extra_terms"]],
    "english": [normalize_english(row) for row in ENGLISH_RAW["items"]] + [normalize_extra(row, "english") for row in EXPERIENCE["extra_terms"]],
}


def validate_language(language: str) -> str:
    if language not in LANGUAGES:
        raise HTTPException(404, "Язык не найден")
    return language


def db_session():
    db = SessionLocal()
    try:
        yield db
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


def token_digest(token: str) -> str:
    return hashlib.sha256(token.encode()).hexdigest()


def current_user(mgc_session: str | None = Cookie(default=None), db: Session = Depends(db_session)) -> User:
    if not mgc_session:
        raise HTTPException(401, "Требуется вход")
    session = db.scalar(select(LoginSession).where(LoginSession.token_hash == token_digest(mgc_session)))
    if not session:
        raise HTTPException(401, "Сессия не найдена")
    expires = session.expires_at
    if expires.tzinfo is None:
        expires = expires.replace(tzinfo=timezone.utc)
    if expires <= datetime.now(timezone.utc):
        db.delete(session)
        db.commit()
        raise HTTPException(401, "Сессия истекла")
    user = db.get(User, session.user_id)
    if not user:
        raise HTTPException(401, "Пользователь не найден")
    apply_rls_context(db, user)
    return user


def apply_rls_context(db: Session, user: User | None = None, *, system_admin: bool = False) -> None:
    """Bind application identity to the current PostgreSQL transaction for FORCE RLS policies."""
    if not RLS_ENABLED or not DATABASE_URL.startswith("postgresql"):
        return
    role = "admin" if system_admin else (user.role if user else "")
    user_id = str(user.id) if user else "0"
    department = user.department if user else "__system__"
    db.execute(text("SELECT set_config('app.user_id', :v, true)"), {"v": user_id})
    db.execute(text("SELECT set_config('app.role', :v, true)"), {"v": role})
    db.execute(text("SELECT set_config('app.department', :v, true)"), {"v": department})


class AuthPayload(BaseModel):
    username: str = Field(min_length=3, max_length=80)
    password: str = Field(min_length=6, max_length=200)
    display_name: str | None = Field(default=None, max_length=120)


class ProgressPayload(BaseModel):
    term_id: str
    status: str = Field(pattern="^(learning|known)$")


class PreferencePayload(BaseModel):
    language: str


class LearningSettingsPayload(BaseModel):
    show_pinyin: bool = True
    show_reading: bool = True
    server_audio_enabled: bool = True


class PronunciationPayload(BaseModel):
    language: str
    text: str = Field(min_length=1, max_length=TTS_MAX_CHARS)
    rate: float = Field(default=0.9, ge=0.55, le=1.25)


class UserDepartmentPayload(BaseModel):
    department: str = Field(min_length=1, max_length=160)


class AssistantPayload(BaseModel):
    language: str
    question: str = Field(min_length=2, max_length=1000)


class ExamPayload(BaseModel):
    language: str
    score: int = Field(ge=0, le=50)
    total: int = Field(default=50, ge=1, le=50)
    answers: list[dict[str, Any]] = Field(default_factory=list)


class CourseDayPayload(BaseModel):
    language: str
    day: int = Field(ge=1, le=30)
    score: int = Field(ge=0, le=5)
    total: int = Field(default=5, ge=5, le=5)


class PracticePayload(BaseModel):
    session_id: str = Field(min_length=6, max_length=120)
    kind: str = Field(pattern="^(quiz|scenario|pair|course_day|exam|tone_lab)$")
    language: str
    topic: str = Field(default="", max_length=160)
    score: int = Field(ge=0, le=100)
    total: int = Field(ge=1, le=100)


class SpendPayload(BaseModel):
    reward_id: str = Field(min_length=2, max_length=80)
    language: str
    context: dict[str, Any] = Field(default_factory=dict)


class GameFinishPayload(BaseModel):
    answers: list[Any] = Field(default_factory=list, max_length=50)


class NotificationSettingsPayload(BaseModel):
    mode: str = Field(pattern="^(off|minimal|normal|active)$")
    window_start: str = Field(pattern=r"^([01]\d|2[0-3]):[0-5]\d$")
    window_end: str = Field(pattern=r"^([01]\d|2[0-3]):[0-5]\d$")
    browser_enabled: bool = False


class CustomTermPayload(BaseModel):
    language: str
    shop: str = Field(default="", max_length=160)
    topic: str = Field(min_length=1, max_length=160)
    subtopic: str = Field(default="", max_length=160)
    level: str = Field(default="A1", pattern="^(A1|A2|B1|B2|C1)$")
    term: str = Field(min_length=1, max_length=500)
    pronunciation: str = Field(default="", max_length=500)
    reading: str = Field(default="", max_length=500)
    translation: str = Field(min_length=1, max_length=500)
    example: str = Field(default="", max_length=3000)
    example_translation: str = Field(default="", max_length=3000)
    tags: str = Field(default="", max_length=1000)
    status: str = Field(default="published", pattern="^(draft|review|published|archived)$")


class TermReviewPayload(BaseModel):
    note: str = Field(default="", max_length=2000)


class SRSReviewPayload(BaseModel):
    language: str
    term_id: str = Field(min_length=1, max_length=80)
    quality: int = Field(ge=0, le=5)


class QuestionAttemptPayload(BaseModel):
    session_id: str = Field(default="", max_length=120)
    question_id: str = Field(min_length=1, max_length=160)
    term_id: str = Field(default="", max_length=80)
    language: str
    topic: str = Field(default="", max_length=160)
    kind: str = Field(default="quiz", max_length=50)
    correct: bool
    response_ms: int = Field(default=0, ge=0, le=600000)
    selected: str = Field(default="", max_length=1000)


class UserRolePayload(BaseModel):
    role: str = Field(pattern="^(user|manager|editor|admin)$")


class AlertAckPayload(BaseModel):
    note: str = Field(default="", max_length=500)


class PilotGroupPayload(BaseModel):
    name: str = Field(min_length=2, max_length=160)
    department: str = Field(default="General", min_length=1, max_length=160)
    description: str = Field(default="", max_length=2000)
    status: str = Field(default="draft", pattern="^(draft|active|paused|completed)$")
    wave: int = Field(default=1, ge=1, le=100)
    starts_at: datetime | None = None
    ends_at: datetime | None = None


class PilotGroupUpdatePayload(BaseModel):
    name: str | None = Field(default=None, min_length=2, max_length=160)
    department: str | None = Field(default=None, min_length=1, max_length=160)
    description: str | None = Field(default=None, max_length=2000)
    status: str | None = Field(default=None, pattern="^(draft|active|paused|completed)$")
    wave: int | None = Field(default=None, ge=1, le=100)
    starts_at: datetime | None = None
    ends_at: datetime | None = None


class PilotGroupStatusPayload(BaseModel):
    status: str = Field(pattern="^(draft|active|paused|completed)$")


class PilotFeaturePayload(BaseModel):
    enabled: bool


class FeatureFlagPayload(BaseModel):
    flag_key: str = Field(pattern=r"^[a-z0-9_\-]{2,80}$")
    title: str = Field(min_length=2, max_length=160)
    description: str = Field(default="", max_length=2000)
    default_enabled: bool = True


class TrackAssignmentPayload(BaseModel):
    language: str
    track_name: str = Field(min_length=2, max_length=180)
    topic: str = Field(default="", max_length=180)
    target_level: str = Field(default="A1", pattern="^(A1|A2|B1|B2|C1)$")
    due_date: str = Field(default="", pattern=r"^$|^\d{4}-\d{2}-\d{2}$")


app = FastAPI(title="MGC Languages", version=APP_VERSION, docs_url="/docs" if APP_ENV != "production" else None, redoc_url=None)

OTEL_STATUS = {"enabled": OTEL_ENABLED, "available": False, "configured": False, "error": ""}
if OTEL_ENABLED:
    try:
        from opentelemetry import trace
        from opentelemetry.sdk.resources import Resource
        from opentelemetry.sdk.trace import TracerProvider
        from opentelemetry.sdk.trace.export import BatchSpanProcessor
        from opentelemetry.instrumentation.fastapi import FastAPIInstrumentor
        from opentelemetry.instrumentation.sqlalchemy import SQLAlchemyInstrumentor
        provider = TracerProvider(resource=Resource.create({"service.name": OTEL_SERVICE_NAME, "service.version": APP_VERSION, "deployment.environment": APP_ENV}))
        if OTEL_EXPORTER_OTLP_ENDPOINT:
            from opentelemetry.exporter.otlp.proto.http.trace_exporter import OTLPSpanExporter
            provider.add_span_processor(BatchSpanProcessor(OTLPSpanExporter(endpoint=OTEL_EXPORTER_OTLP_ENDPOINT, timeout=OTEL_EXPORT_TIMEOUT_SECONDS)))
            OTEL_STATUS["configured"] = True
        trace.set_tracer_provider(provider)
        FastAPIInstrumentor.instrument_app(app, excluded_urls="/health/live,/metrics")
        SQLAlchemyInstrumentor().instrument(engine=engine)
        OTEL_STATUS["available"] = True
    except Exception as exc:
        OTEL_STATUS["error"] = type(exc).__name__
        logger.warning(json.dumps({"event":"otel_initialization_failed","error":str(exc)[:200]}, ensure_ascii=False))

@app.exception_handler(SQLAlchemyError)
async def sqlalchemy_unavailable_handler(request: Request, exc: SQLAlchemyError):
    request_id = getattr(getattr(request, "state", None), "request_id", "")
    logger.error(json.dumps({"event":"database_error","request_id":request_id,"path":request.url.path,"error":type(exc).__name__}, ensure_ascii=False))
    with _RELIABILITY_LOCK:
        _RELIABILITY_RUNTIME["database_failures_total"] += 1
    return JSONResponse(
        {"detail":"Сервис временно не может обратиться к базе данных. Повторите позже.","code":"database_unavailable","retryable":True,"request_id":request_id},
        status_code=503, headers={"Retry-After":"5","X-Request-ID":request_id},
    )

app.add_middleware(SessionMiddleware, secret_key=OIDC_STATE_SECRET, same_site="lax", https_only=COOKIE_SECURE, max_age=600)
if TRUSTED_HOSTS:
    app.add_middleware(TrustedHostMiddleware, allowed_hosts=TRUSTED_HOSTS)
if CORS_ORIGINS:
    app.add_middleware(CORSMiddleware, allow_origins=CORS_ORIGINS, allow_credentials=True, allow_methods=["GET", "POST", "PUT", "PATCH", "DELETE"], allow_headers=["Content-Type", "X-CSRF-Token", "X-Request-ID"])

_RATE_LOCK = threading.Lock()
_RATE_BUCKETS: dict[str, deque[float]] = defaultdict(deque)
_METRIC_LOCK = threading.Lock()
_METRIC_REQUESTS: Counter = Counter()
_METRIC_LATENCY: Counter = Counter()
_RELIABILITY_LOCK = threading.Lock()
_RELIABILITY_RUNTIME: Counter = Counter()
_OP_EVENT_LOCK = threading.Lock()
_OP_EVENT_LAST: dict[str, float] = {}
_SLO_LOCK = threading.Lock()
_HTTP_WINDOW: deque[tuple[float, float, int, str]] = deque(maxlen=20000)
_RECOVERY_LOCK = threading.Lock()
_RECOVERY_STATE: dict[str, Any] = {
    "state":"healthy", "reason":"startup", "since":datetime.now(timezone.utc),
    "recover_after":None, "last_transition":datetime.now(timezone.utc),
}

def _client_key(request: Request) -> str:
    forwarded = request.headers.get("x-forwarded-for", "").split(",")[0].strip()
    raw = forwarded or (request.client.host if request.client else "unknown")
    return hashlib.sha256(raw.encode()).hexdigest()[:20]

def rate_limit(request: Request, bucket: str, limit: int, window_seconds: int) -> None:
    key = f"{bucket}:{_client_key(request)}"
    now = time.monotonic()
    with _RATE_LOCK:
        q = _RATE_BUCKETS[key]
        while q and now - q[0] > window_seconds:
            q.popleft()
        if len(q) >= limit:
            retry = max(1, int(window_seconds - (now - q[0])))
            raise HTTPException(429, "Слишком много запросов. Повторите позже.", headers={"Retry-After": str(retry)})
        q.append(now)

def _normalized_metric_path(path: str) -> str:
    path = re.sub(r"/\d+(?=/|$)", "/{id}", path)
    path = re.sub(r"/[0-9a-f]{16,}(?=/|$)", "/{token}", path, flags=re.I)
    return path

def _percentile(values: list[float], percentile: float) -> float:
    if not values:
        return 0.0
    ordered = sorted(values)
    index = max(0, min(len(ordered) - 1, int(round((len(ordered) - 1) * percentile))))
    return round(float(ordered[index]), 2)


def slo_snapshot() -> dict[str, Any]:
    now = time.time()
    cutoff = now - SLO_WINDOW_MINUTES * 60
    with _SLO_LOCK:
        while _HTTP_WINDOW and _HTTP_WINDOW[0][0] < cutoff:
            _HTTP_WINDOW.popleft()
        rows = list(_HTTP_WINDOW)
    durations = [row[1] for row in rows]
    server_errors = sum(1 for _, _, status, _ in rows if status >= 500)
    total = len(rows)
    availability = round((total - server_errors) * 100 / total, 3) if total else 100.0
    error_rate = round(server_errors * 100 / total, 3) if total else 0.0
    p95 = _percentile(durations, 0.95)
    allowed_error_rate = max(0.001, 100.0 - SLO_AVAILABILITY_TARGET_PERCENT)
    burn_rate = round(error_rate / allowed_error_rate, 2) if total else 0.0
    status = "insufficient_data" if total < 20 else ("met" if availability >= SLO_AVAILABILITY_TARGET_PERCENT and error_rate <= SLO_ERROR_RATE_TARGET_PERCENT and p95 <= SLO_P95_TARGET_MS else "missed")
    return {
        "window_minutes":SLO_WINDOW_MINUTES, "service_window":PILOT_SERVICE_WINDOW, "sla_mode":PILOT_SLA_MODE,
        "scope":"process-local", "instance_id":INSTANCE_ID, "aggregation_note":"For multi-instance rollout aggregate Prometheus metrics centrally; the in-app window is per process.",
        "samples":total, "server_errors":server_errors, "availability_percent":availability, "error_rate_percent":error_rate,
        "p50_ms":_percentile(durations,0.50), "p95_ms":p95, "p99_ms":_percentile(durations,0.99),
        "targets":{"availability_percent":SLO_AVAILABILITY_TARGET_PERCENT,"error_rate_percent":SLO_ERROR_RATE_TARGET_PERCENT,"p95_ms":SLO_P95_TARGET_MS},
        "error_budget_burn_rate":burn_rate, "status":status,
    }


def _recovery_view() -> dict[str, Any]:
    with _RECOVERY_LOCK:
        data = dict(_RECOVERY_STATE)
    for key in ("since", "recover_after", "last_transition"):
        if isinstance(data.get(key), datetime):
            data[key] = data[key].isoformat()
    return data


def update_recovery_state(*, ready_ok: bool, checks: dict[str, Any] | None = None, tts_health: dict[str, Any] | None = None) -> dict[str, Any]:
    checks = checks or {}
    tts_health = tts_health or {}
    db_bad = checks.get("database") == "failed" or (checks.get("schema_head") or {}).get("status") == "failed"
    config_bad = any(checks.get(k) == "failed" for k in ("state_secret","metrics_protection","trusted_hosts","cors_policy","cookie_policy","tts_legacy_get"))
    tts_degraded = _tts_circuit_open() or not tts_health.get("server_available", True)
    if not ready_ok and (db_bad or config_bad):
        desired, reason = "unavailable", "readiness_failed"
    elif tts_degraded:
        desired, reason = "degraded", "tts_fallback"
    else:
        desired, reason = "healthy", "all_checks_ok"
    now = datetime.now(timezone.utc)
    with _RECOVERY_LOCK:
        current = _RECOVERY_STATE["state"]
        if desired == "healthy" and current in {"unavailable","degraded"}:
            _RECOVERY_STATE.update({"state":"recovering","reason":"stability_window","since":now,"recover_after":now + timedelta(seconds=RECOVERY_STABLE_SECONDS),"last_transition":now})
        elif current == "recovering" and desired == "healthy":
            recover_after = _RECOVERY_STATE.get("recover_after")
            if isinstance(recover_after, datetime) and now >= recover_after:
                _RECOVERY_STATE.update({"state":"healthy","reason":reason,"since":now,"recover_after":None,"last_transition":now})
        elif desired != "healthy" and current != desired:
            _RECOVERY_STATE.update({"state":desired,"reason":reason,"since":now,"recover_after":None,"last_transition":now})
        elif current == "healthy" and desired == "healthy":
            _RECOVERY_STATE["reason"] = reason
    return _recovery_view()


def audit_event(db: Session, event_type: str, *, actor_user_id: int | None = None, target_type: str = "", target_id: str = "", request: Request | None = None, metadata: dict[str, Any] | None = None) -> None:
    metadata_json = json.dumps(metadata or {}, ensure_ascii=False, sort_keys=True, default=str)
    if DATABASE_URL.startswith("postgresql"):
        db.execute(text("SELECT pg_advisory_xact_lock(:lock_id)"), {"lock_id": AUDIT_CHAIN_LOCK_ID})
    previous = db.scalar(select(AuditLog).order_by(AuditLog.id.desc()).limit(1))
    if not previous:
        anchor = db.get(AuditAnchor, 1)
        previous_hash = anchor.last_deleted_hash if anchor else ""
    else:
        previous_hash = previous.event_hash or ""
    request_id = getattr(getattr(request, "state", None), "request_id", "") if request else ""
    source_ip_hash = _client_key(request) if request else ""
    canonical = "|".join([previous_hash, event_type, str(actor_user_id or ""), target_type, target_id, request_id, source_ip_hash, metadata_json])
    event_hash = hashlib.sha256(canonical.encode("utf-8")).hexdigest()
    db.add(AuditLog(event_type=event_type, actor_user_id=actor_user_id, target_type=target_type, target_id=target_id, request_id=request_id, source_ip_hash=source_ip_hash, metadata_json=metadata_json, previous_hash=previous_hash, event_hash=event_hash))

def operational_event(event_type: str, *, severity: str = "warning", component: str = "app", status_code: int = 0, request: Request | None = None, path: str = "", detail: str = "", metadata: dict[str, Any] | None = None, throttle_seconds: int = 0) -> None:
    """Best-effort persistent operational telemetry. Failure to write telemetry never breaks the request path."""
    request_id = getattr(getattr(request, "state", None), "request_id", "") if request else ""
    event_path = path or (request.url.path if request else "")
    if throttle_seconds > 0:
        throttle_key = event_type + "|" + component + "|" + json.dumps(metadata or {}, ensure_ascii=False, sort_keys=True, default=str)
        now_mono = time.monotonic()
        with _OP_EVENT_LOCK:
            last = _OP_EVENT_LAST.get(throttle_key)
            if last is not None and now_mono - last < throttle_seconds:
                return
            _OP_EVENT_LAST[throttle_key] = now_mono
    try:
        with SessionLocal() as event_db:
            event_db.add(OperationalEvent(
                event_type=event_type, severity=severity[:20], component=component[:60],
                status_code=int(status_code or 0), request_id=request_id[:80], path=event_path[:240],
                detail=(detail or "")[:500], metadata_json=json.dumps(metadata or {}, ensure_ascii=False, default=str),
            ))
            event_db.commit()
    except Exception as exc:
        logger.warning(json.dumps({"event":"operational_event_write_failed","type":event_type,"error":str(exc)[:160]}, ensure_ascii=False))


def cleanup_runtime_data(db: Session, *, now: datetime | None = None, dry_run: bool = False) -> dict[str, Any]:
    apply_rls_context(db, system_admin=True)
    now = now or datetime.now(timezone.utc)
    cutoffs = {
        "audit": now - timedelta(days=AUDIT_RETENTION_DAYS),
        "operational": now - timedelta(days=OPERATIONAL_EVENT_RETENTION_DAYS),
        "nudges": now - timedelta(days=NUDGE_RETENTION_DAYS),
        "alerts": now - timedelta(days=ALERT_RETENTION_DAYS),
        "pilot_usage": now - timedelta(days=PILOT_USAGE_RETENTION_DAYS),
    }
    predicates = {
        "expired_sessions": (LoginSession, LoginSession.expires_at <= now),
        "old_audit": (AuditLog, AuditLog.created_at < cutoffs["audit"]),
        "old_operational_events": (OperationalEvent, OperationalEvent.created_at < cutoffs["operational"]),
        "old_nudges": (LearningNudge, LearningNudge.created_at < cutoffs["nudges"]),
        "old_resolved_alerts": (PilotAlert, PilotAlert.resolved_at.is_not(None) & (PilotAlert.resolved_at < cutoffs["alerts"])),
        "old_pilot_usage": (PilotDailyUsage, PilotDailyUsage.day_key < cutoffs["pilot_usage"].date().isoformat()),
    }
    counts: dict[str, int] = {}
    for key, (model, predicate) in predicates.items():
        counts[key] = int(db.scalar(select(func.count()).select_from(model).where(predicate)) or 0)
    if not dry_run:
        db.execute(delete(LoginSession).where(LoginSession.expires_at <= now).execution_options(synchronize_session=False))
        last_removed = db.scalar(select(AuditLog).where(AuditLog.created_at < cutoffs["audit"]).order_by(AuditLog.id.desc()).limit(1))
        if last_removed:
            anchor = db.get(AuditAnchor, 1)
            if not anchor:
                anchor = AuditAnchor(id=1); db.add(anchor)
            anchor.last_deleted_id = last_removed.id
            anchor.last_deleted_hash = last_removed.event_hash or last_removed.previous_hash or ""
            anchor.updated_at = now
        db.execute(delete(AuditLog).where(AuditLog.created_at < cutoffs["audit"]).execution_options(synchronize_session=False))
        db.execute(delete(OperationalEvent).where(OperationalEvent.created_at < cutoffs["operational"]).execution_options(synchronize_session=False))
        db.execute(delete(LearningNudge).where(LearningNudge.created_at < cutoffs["nudges"]).execution_options(synchronize_session=False))
        db.execute(delete(PilotAlert).where(PilotAlert.resolved_at.is_not(None), PilotAlert.resolved_at < cutoffs["alerts"]).execution_options(synchronize_session=False))
        db.execute(delete(PilotDailyUsage).where(PilotDailyUsage.day_key < cutoffs["pilot_usage"].date().isoformat()).execution_options(synchronize_session=False))
        db.commit()
    return {
        "dry_run": dry_run, "counts": counts,
        "retention_days": {"audit":AUDIT_RETENTION_DAYS,"operational_events":OPERATIONAL_EVENT_RETENTION_DAYS,"nudges":NUDGE_RETENTION_DAYS,"resolved_alerts":ALERT_RETENTION_DAYS,"pilot_usage":PILOT_USAGE_RETENTION_DAYS},
        "run_at": now.isoformat(),
    }

CSRF_EXEMPT = {"/api/login", "/api/register", "/api/auth/oidc/callback"}

@app.middleware("http")
async def security_and_observability(request: Request, call_next):
    started = time.perf_counter()
    request_id = (request.headers.get("x-request-id") or uuid.uuid4().hex)[:80]
    request.state.request_id = request_id
    if request.method in {"POST", "PUT", "PATCH", "DELETE"} and request.url.path.startswith("/api/") and request.url.path not in CSRF_EXEMPT:
        session_cookie = request.cookies.get("mgc_session")
        if session_cookie:
            cookie_token = request.cookies.get("mgc_csrf", "")
            header_token = request.headers.get("x-csrf-token", "")
            if not cookie_token or not header_token or not hmac.compare_digest(cookie_token, header_token):
                return JSONResponse({"detail": "CSRF-проверка не пройдена"}, status_code=403, headers={"X-Request-ID": request_id})
    try:
        response = await call_next(request)
    except Exception as exc:
        logger.exception(json.dumps({"event":"unhandled_exception","request_id":request_id,"path":request.url.path}, ensure_ascii=False))
        operational_event("request.exception", severity="error", component="app", status_code=500, request=request, detail=type(exc).__name__)
        raise
    duration = time.perf_counter() - started
    if response.status_code >= 500 and not request.url.path.startswith(("/health", "/ready", "/metrics")):
        with _RELIABILITY_LOCK:
            _RELIABILITY_RUNTIME["http_5xx_total"] += 1
        operational_event("request.5xx", severity="error", component="http", status_code=response.status_code, request=request)
    metric_path = _normalized_metric_path(request.url.path)
    with _METRIC_LOCK:
        _METRIC_REQUESTS[(request.method, metric_path, response.status_code)] += 1
        _METRIC_LATENCY[(request.method, metric_path)] += duration
    if request.url.path.startswith("/api/") and request.url.path not in {"/api/admin/slo", "/api/admin/it-dashboard"}:
        with _SLO_LOCK:
            _HTTP_WINDOW.append((time.time(), duration * 1000.0, int(response.status_code), metric_path))
    response.headers["X-Request-ID"] = request_id
    response.headers["X-MGC-Version"] = APP_VERSION
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "no-referrer"
    response.headers["Permissions-Policy"] = "camera=(), geolocation=(), microphone=()"
    response.headers["Content-Security-Policy"] = "default-src 'self'; img-src 'self' data:; style-src 'self' 'unsafe-inline'; script-src 'self'; connect-src 'self'; media-src 'self' blob:; object-src 'none'; base-uri 'self'; frame-ancestors 'none'"
    if COOKIE_SECURE:
        response.headers["Strict-Transport-Security"] = "max-age=31536000; includeSubDomains"
    if request.url.path.startswith("/api/"):
        response.headers["Cache-Control"] = "no-store"
    logger.info(json.dumps({"event":"http_request","request_id":request_id,"method":request.method,"path":metric_path,"status":response.status_code,"duration_ms":round(duration*1000,2)}, ensure_ascii=False))
    return response


def user_view(user: User):
    return {
        "id": user.id,
        "username": user.username,
        "display_name": user.display_name,
        "role": user.role,
        "department": user.department,
        "preferred_language": user.preferred_language,
    }



ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]
LEVEL_TITLES = [
    "Новичок", "Ученик", "Практик", "Специалист", "Уверенный",
    "Профессионал", "Наставник", "Эксперт", "Мастер", "Грандмастер",
]
REWARD_CATALOG = {
    "hint_small": {"title": "Небольшая подсказка", "price": 10, "description": "Намёк без полного ответа"},
    "show_pinyin": {"title": "Pinyin / транскрипция", "price": 10, "description": "Показать произношение"},
    "slow_audio": {"title": "Медленное аудио", "price": 15, "description": "Произнести фразу медленнее"},
    "eliminate_option": {"title": "Убрать неверный вариант", "price": 20, "description": "Скрыть один неправильный ответ"},
    "sentence_start": {"title": "Начало фразы", "price": 25, "description": "Показать начало рабочего предложения"},
    "explain_word": {"title": "Объяснить слово", "price": 25, "description": "Значение и короткий пример"},
    "pronunciation_breakdown": {"title": "Разбор произношения", "price": 30, "description": "Произношение по частям"},
    "work_example": {"title": "Рабочий пример", "price": 30, "description": "Пример из автомобильного контекста"},
    "mistake_explain": {"title": "Почему ответ неверный?", "price": 40, "description": "Разбор ошибки"},
    "smart_revision": {"title": "Smart Revision", "price": 75, "description": "Персональная подборка на повторение"},
    "focus_mode": {"title": "Focus Mode", "price": 100, "description": "10 терминов из слабых мест"},
    "workshop_pack": {"title": "Тренировка по цеху", "price": 100, "description": "Подборка по выбранной теме/цеху"},
    "roleplay": {"title": "AI Role Play", "price": 150, "description": "Расширенная ролевая тренировка"},
    "meeting_simulator": {"title": "Meeting Simulator", "price": 200, "description": "Сценарий рабочего совещания"},
}


def current_week_key(now: datetime | None = None) -> str:
    now = now or datetime.now(timezone.utc)
    iso = now.isocalendar()
    return f"{iso.year}-W{iso.week:02d}"


def level_info(lifetime_xp: int) -> dict[str, Any]:
    level = min(100, max(1, lifetime_xp // 500 + 1))
    if level == 100:
        title = "Легенда"
        progress = 500
        needed = 500
        next_xp = None
    else:
        tier = min(9, (level - 1) // 10)
        within = (level - 1) % 10
        title = f"{LEVEL_TITLES[tier]} {ROMAN[within]}"
        floor = (level - 1) * 500
        progress = lifetime_xp - floor
        needed = 500
        next_xp = level * 500
    return {"level": level, "title": title, "progress_xp": progress, "level_xp": needed, "next_level_xp": next_xp}


def get_profile(db: Session, user_id: int) -> GamificationProfile:
    profile = db.get(GamificationProfile, user_id)
    if not profile:
        profile = GamificationProfile(user_id=user_id, week_key=current_week_key())
        db.add(profile)
        db.flush()
    week = current_week_key()
    if profile.week_key != week:
        profile.week_key = week
        profile.weekly_xp = 0
    return profile


def gamification_view(db: Session, user_id: int) -> dict[str, Any]:
    profile = get_profile(db, user_id)
    info = level_info(profile.lifetime_xp)
    return {
        **info,
        "lifetime_xp": profile.lifetime_xp,
        "spendable_xp": profile.spendable_xp,
        "weekly_xp": profile.weekly_xp,
        "last_activity_at": profile.last_activity_at.isoformat() if profile.last_activity_at else None,
    }


def award_xp(
    db: Session, user_id: int, event_type: str, source_id: str, raw_xp: int,
    *, language: str = "", topic: str = "", idempotency_key: str | None = None,
    metadata: dict[str, Any] | None = None, anti_farm: bool = True,
) -> dict[str, Any]:
    key = idempotency_key or f"{user_id}:{event_type}:{source_id}:{secrets.token_hex(6)}"
    existing = db.scalar(select(XPEvent).where(XPEvent.idempotency_key == key))
    if existing:
        return {"awarded": 0, "duplicate": True, "profile": gamification_view(db, user_id)}
    multiplier = 100
    if anti_farm and source_id:
        since = datetime.now(timezone.utc) - timedelta(hours=24)
        repeat_count = len(db.scalars(select(XPEvent).where(
            XPEvent.user_id == user_id,
            XPEvent.event_type == event_type,
            XPEvent.source_id == source_id,
            XPEvent.created_at >= since,
            XPEvent.lifetime_delta > 0,
        )).all())
        if repeat_count >= 6:
            multiplier = 10
        elif repeat_count >= 3:
            multiplier = 50
    awarded = max(1, round(raw_xp * multiplier / 100)) if raw_xp > 0 else 0
    usage = pilot_daily_usage(db, user_id)
    remaining = max(0, PILOT_DAILY_XP_CAP - int(usage.xp_awarded or 0))
    awarded = min(awarded, remaining)
    usage.xp_awarded = int(usage.xp_awarded or 0) + awarded
    usage.updated_at = datetime.now(timezone.utc)
    profile = get_profile(db, user_id)
    profile.lifetime_xp += awarded
    profile.spendable_xp += awarded
    profile.weekly_xp += awarded
    profile.last_activity_at = datetime.now(timezone.utc)
    profile.updated_at = datetime.now(timezone.utc)
    db.add(XPEvent(
        user_id=user_id, event_type=event_type, source_id=source_id, language=language,
        topic=topic, xp_delta=awarded, lifetime_delta=awarded, raw_xp=raw_xp,
        multiplier_pct=multiplier, idempotency_key=key,
        metadata_json=json.dumps(metadata or {}, ensure_ascii=False),
    ))
    db.flush()
    return {"awarded": awarded, "duplicate": False, "multiplier_pct": multiplier, "profile": gamification_view(db, user_id)}


def spend_xp(db: Session, user_id: int, reward_id: str, *, language: str = "", metadata: dict[str, Any] | None = None) -> dict[str, Any]:
    reward = REWARD_CATALOG.get(reward_id)
    if not reward:
        raise HTTPException(404, "Неизвестная помощь")
    profile = get_profile(db, user_id)
    price = int(reward["price"])
    if profile.spendable_xp < price:
        raise HTTPException(409, f"Недостаточно XP: нужно {price}, доступно {profile.spendable_xp}")
    profile.spendable_xp -= price
    profile.updated_at = datetime.now(timezone.utc)
    db.add(XPEvent(
        user_id=user_id, event_type="spend", source_id=reward_id, language=language,
        topic="", xp_delta=-price, lifetime_delta=0, raw_xp=price, multiplier_pct=100,
        idempotency_key=f"spend:{user_id}:{reward_id}:{secrets.token_hex(10)}",
        metadata_json=json.dumps(metadata or {}, ensure_ascii=False),
    ))
    db.flush()
    return {"spent": price, "reward": reward, "profile": gamification_view(db, user_id)}


def require_roles(*roles: str):
    def dependency(user: User = Depends(current_user)) -> User:
        if user.role not in roles:
            raise HTTPException(403, "Недостаточно прав")
        return user
    return dependency


def custom_term_view(row: CustomTerm) -> dict[str, Any]:
    return {
        "id": row.public_id, "custom": True, "language": row.language, "shop": row.shop,
        "topic": row.topic, "category": row.shop or row.topic, "subcategory": row.subtopic or row.topic,
        "level": row.level, "term": row.term,
        "pronunciation": row.pronunciation or (english_pronunciation(row.term)["ipa"] if row.language == "english" else ""),
        "reading": row.reading or (english_pronunciation(row.term)["reading"] if row.language == "english" else pinyin_to_ru_approx(row.pronunciation)),
        "translation": row.translation, "example": row.example or row.term, "example_pronunciation": row.pronunciation,
        "example_translation": row.example_translation or row.translation, "note": row.tags,
        "status": row.status, "tags": row.tags, "source_type": row.source_type, "source_ref": row.source_ref,
    }


def terms_for(language: str, db: Session) -> list[dict[str, Any]]:
    rows = list(TERMS[language])
    custom = db.scalars(select(CustomTerm).where(
        CustomTerm.language == language, CustomTerm.status == "published"
    ).order_by(CustomTerm.id)).all()
    rows.extend(custom_term_view(row) for row in custom)
    return rows


def term_by_id(language: str, term_id: str, db: Session) -> dict[str, Any] | None:
    return next((row for row in terms_for(language, db) if row["id"] == term_id), None)


def ensure_bootstrap_admin() -> None:
    if AUTH_MODE != "local":
        return
    username = os.getenv("MGC_ADMIN_USERNAME", "").strip().lower()
    password = os.getenv("MGC_ADMIN_PASSWORD", "")
    if not username or not password:
        return
    if APP_ENV in {"pilot", "production"} and password in {"CHANGE_ME_WITH_A_LONG_PASSWORD", "CHANGE_ME", "admin", "password"}:
        logger.warning(json.dumps({"event":"insecure_bootstrap_admin_password","env":APP_ENV}))
        return
    with SessionLocal() as db:
        try:
            user = db.scalar(select(User).where(User.username == username))
        except Exception:
            db.rollback()
            logger.info(json.dumps({"event":"bootstrap_admin_deferred","reason":"schema_not_ready"}))
            return
        if user:
            if user.role != "admin":
                user.role = "admin"
                db.commit()
            return
        user = User(username=username, display_name=os.getenv("MGC_ADMIN_DISPLAY_NAME", "MGC Admin").strip() or "MGC Admin", password_hash=make_password_hash(password), role="admin")
        db.add(user)
        db.commit()


ensure_bootstrap_admin()


def create_login_session(db: Session, user: User, response: Response):
    raw = secrets.token_urlsafe(36)
    csrf = secrets.token_urlsafe(24)
    now = datetime.now(timezone.utc)
    expiry = now + timedelta(hours=SESSION_TTL_HOURS)
    db.execute(delete(LoginSession).where(LoginSession.expires_at <= now).execution_options(synchronize_session=False))
    db.add(LoginSession(token_hash=token_digest(raw), user_id=user.id, expires_at=expiry))
    db.commit()
    max_age = SESSION_TTL_HOURS * 3600
    response.set_cookie("mgc_session", raw, max_age=max_age, httponly=True, samesite=COOKIE_SAMESITE, secure=COOKIE_SECURE, path="/")
    response.set_cookie("mgc_csrf", csrf, max_age=max_age, httponly=False, samesite=COOKIE_SAMESITE, secure=COOKIE_SECURE, path="/")
    return {"ok": True, "user": user_view(user), "csrf": csrf}


def readiness_checks(db: Session) -> tuple[bool, dict[str, Any]]:
    checks: dict[str, Any] = {"database":"unknown", "auth_mode":AUTH_MODE, "environment":APP_ENV}
    ok = True
    try:
        db_started = time.perf_counter()
        db.execute(text("SELECT 1"))
        db_latency_ms = round((time.perf_counter() - db_started) * 1000, 2)
        checks["database"] = "ok"
        checks["database_latency_ms"] = db_latency_ms
        if db_latency_ms > DB_READY_MAX_LATENCY_MS:
            checks["database_latency"] = {"status":"failed","value_ms":db_latency_ms,"max_ms":DB_READY_MAX_LATENCY_MS}
            ok = False
        else:
            checks["database_latency"] = {"status":"ok","value_ms":db_latency_ms,"max_ms":DB_READY_MAX_LATENCY_MS}
    except Exception as exc:
        checks["database"] = "failed"
        checks["database_error"] = type(exc).__name__
        ok = False
    if READY_REQUIRE_POSTGRES and DATABASE_URL.startswith("sqlite"):
        checks["postgres_required"] = "failed"
        ok = False
    else:
        checks["postgres_required"] = "ok"
    if READY_REQUIRE_SCHEMA_HEAD:
        try:
            schema_head = db.execute(text("SELECT version_num FROM alembic_version")).scalar_one_or_none()
        except Exception:
            schema_head = None
        if schema_head != EXPECTED_ALEMBIC_HEAD:
            checks["schema_head"] = {"status":"failed","current":schema_head,"expected":EXPECTED_ALEMBIC_HEAD}
            ok = False
        else:
            checks["schema_head"] = {"status":"ok","current":schema_head,"expected":EXPECTED_ALEMBIC_HEAD}
    else:
        checks["schema_head"] = {"status":"not_required","expected":EXPECTED_ALEMBIC_HEAD}
    if READY_REQUIRE_OIDC and AUTH_MODE != "oidc":
        checks["oidc_required"] = "failed"
        ok = False
    else:
        checks["oidc_required"] = "ok" if READY_REQUIRE_OIDC else "not_required"
    if AUTH_MODE == "oidc":
        configured = bool(OIDC_DISCOVERY_URL and OIDC_CLIENT_ID and OIDC_CLIENT_SECRET and OIDC_STATE_SECRET and OIDC_STATE_SECRET != "dev-only-change-me")
        checks["oidc_config"] = "ok" if configured else "failed"
        ok = ok and configured
    else:
        checks["oidc_config"] = "not_required"
    if APP_ENV in {"pilot", "production"} and OIDC_STATE_SECRET == "dev-only-change-me":
        checks["state_secret"] = "failed"
        ok = False
    else:
        checks["state_secret"] = "ok"
    if READY_REQUIRE_REGISTRATION_DISABLED and REGISTRATION_ENABLED:
        checks["self_registration"] = "failed"
        ok = False
    else:
        checks["self_registration"] = "disabled" if not REGISTRATION_ENABLED else "allowed"
    if READY_REQUIRE_METRICS_TOKEN and (not METRICS_ENABLED or not METRICS_TOKEN or METRICS_TOKEN.upper().startswith("CHANGE_ME")):
        checks["metrics_protection"] = "failed"
        ok = False
    else:
        checks["metrics_protection"] = "token" if METRICS_ENABLED and METRICS_TOKEN else ("open" if METRICS_ENABLED else "disabled")
    if READY_REQUIRE_SECURE_COOKIE and not COOKIE_SECURE:
        checks["secure_cookie"] = "failed"
        ok = False
    else:
        checks["secure_cookie"] = "enabled" if COOKIE_SECURE else "not_required"
    if COOKIE_SAMESITE == "none" and not COOKIE_SECURE:
        checks["cookie_policy"] = "failed"
        ok = False
    else:
        checks["cookie_policy"] = "ok"
    if APP_ENV in {"pilot", "production"} and (not TRUSTED_HOSTS or "*" in TRUSTED_HOSTS):
        checks["trusted_hosts"] = "failed"
        ok = False
    else:
        checks["trusted_hosts"] = "ok"
    if APP_ENV in {"pilot", "production"} and "*" in CORS_ORIGINS:
        checks["cors_policy"] = "failed"
        ok = False
    else:
        checks["cors_policy"] = "ok"
    if APP_ENV in {"pilot", "production"} and TTS_LEGACY_GET_ENABLED:
        checks["tts_legacy_get"] = "failed"
        ok = False
    else:
        checks["tts_legacy_get"] = "disabled" if not TTS_LEGACY_GET_ENABLED else "dev_compatibility"
    if READY_REQUIRE_RLS and not RLS_ENABLED:
        checks["row_level_security"] = "failed"; ok = False
    else:
        checks["row_level_security"] = "enabled" if RLS_ENABLED else "not_required"
    if READY_REQUIRE_TERM_APPROVAL and not TERM_APPROVAL_REQUIRED:
        checks["term_approval"] = "failed"; ok = False
    else:
        checks["term_approval"] = "required" if TERM_APPROVAL_REQUIRED else "not_required"
    tts_health = _tts_health()
    if TTS_DISK_CACHE_ENABLED and not tts_health.get("cache_writable", False):
        checks["tts_cache"] = "failed"
        ok = False
    else:
        checks["tts_cache"] = "ok"
    if APP_ENV in {"pilot", "production"} and TTS_CACHE_PERSISTENCE == "ephemeral" and not str(TTS_CACHE_DIR).startswith("/tmp/"):
        checks["tts_cache_persistence"] = "failed"
        ok = False
    else:
        checks["tts_cache_persistence"] = TTS_CACHE_PERSISTENCE
    # Pronunciation is resilient by design: missing server TTS does not take the service down,
    # because the browser voice is an explicit fallback.
    checks["pronunciation"] = "offline_server" if tts_health["server_available"] else "browser_fallback"
    return ok, checks


@app.get("/health")
@app.get("/health/live")
def health():
    return {"status":"ok", "service":"mgc-languages", "version":APP_VERSION, "recovery":_recovery_view()}


@app.get("/ready")
@app.get("/health/ready")
def ready(db: Session = Depends(db_session)):
    ok, checks = readiness_checks(db)
    recovery = update_recovery_state(ready_ok=ok, checks=checks, tts_health=_tts_health())
    payload = {"status":"ready" if ok else "not_ready", "checks":checks, "recovery":recovery}
    if not ok:
        return JSONResponse(payload, status_code=503)
    return payload


@app.get("/metrics", response_class=PlainTextResponse)
def metrics(request: Request):
    if not METRICS_ENABLED:
        raise HTTPException(404, "Metrics disabled")
    if METRICS_TOKEN:
        supplied = request.headers.get("authorization", "")
        expected = f"Bearer {METRICS_TOKEN}"
        if not hmac.compare_digest(supplied, expected):
            raise HTTPException(401, "Metrics authentication required", headers={"WWW-Authenticate":"Bearer"})
    lines = ["# HELP mgc_http_requests_total HTTP requests", "# TYPE mgc_http_requests_total counter"]
    with _METRIC_LOCK:
        for (method, path, status), value in sorted(_METRIC_REQUESTS.items()):
            lines.append(f'mgc_http_requests_total{{method="{method}",path="{path}",status="{status}"}} {value}')
        lines += ["# HELP mgc_http_request_duration_seconds_sum HTTP request duration sum", "# TYPE mgc_http_request_duration_seconds_sum counter"]
        for (method, path), value in sorted(_METRIC_LATENCY.items()):
            lines.append(f'mgc_http_request_duration_seconds_sum{{method="{method}",path="{path}"}} {value:.6f}')
    with _RELIABILITY_LOCK:
        reliability_runtime = dict(_RELIABILITY_RUNTIME)
    lines += [
        "# HELP mgc_database_failures_total Database failures observed since process start",
        "# TYPE mgc_database_failures_total counter",
        f"mgc_database_failures_total {reliability_runtime.get('database_failures_total',0)}",
        "# HELP mgc_http_5xx_runtime_total HTTP 5xx responses observed since process start",
        "# TYPE mgc_http_5xx_runtime_total counter",
        f"mgc_http_5xx_runtime_total {reliability_runtime.get('http_5xx_total',0)}",
    ]
    with _TTS_RUNTIME_LOCK:
        tts_runtime = dict(_TTS_RUNTIME)
    try:
        now = datetime.now(timezone.utc)
        with SessionLocal() as metric_db:
            op_24h = int(metric_db.scalar(select(func.count()).select_from(OperationalEvent).where(OperationalEvent.created_at >= now - timedelta(hours=24))) or 0)
            op_err_24h = int(metric_db.scalar(select(func.count()).select_from(OperationalEvent).where(OperationalEvent.created_at >= now - timedelta(hours=24), OperationalEvent.severity == "error")) or 0)
            expired_sessions = int(metric_db.scalar(select(func.count()).select_from(LoginSession).where(LoginSession.expires_at <= now)) or 0)
            open_alerts = int(metric_db.scalar(select(func.count()).select_from(PilotAlert).where(PilotAlert.status.in_(["open","acknowledged"]))) or 0)
            apply_rls_context(metric_db, system_admin=True)
            srs_due = int(metric_db.scalar(select(func.count()).select_from(SRSCard).where(SRSCard.due_at <= now)) or 0)
            q_total = int(metric_db.scalar(select(func.count()).select_from(QuestionAttempt)) or 0)
        lines += [
            "# HELP mgc_operational_events_24h Operational events recorded in last 24 hours",
            "# TYPE mgc_operational_events_24h gauge",
            f"mgc_operational_events_24h {op_24h}",
            "# HELP mgc_operational_errors_24h Operational error events recorded in last 24 hours",
            "# TYPE mgc_operational_errors_24h gauge",
            f"mgc_operational_errors_24h {op_err_24h}",
            "# HELP mgc_expired_sessions_pending_cleanup Expired sessions awaiting maintenance cleanup",
            "# TYPE mgc_expired_sessions_pending_cleanup gauge",
            f"mgc_expired_sessions_pending_cleanup {expired_sessions}",
            "# HELP mgc_pilot_alerts_open Active pilot operations alerts",
            "# TYPE mgc_pilot_alerts_open gauge",
            f"mgc_pilot_alerts_open {open_alerts}",
            "# HELP mgc_srs_cards_due Adaptive SRS cards currently due across the pilot",
            "# TYPE mgc_srs_cards_due gauge",
            f"mgc_srs_cards_due {srs_due}",
            "# HELP mgc_question_attempts_total Learning question attempts recorded",
            "# TYPE mgc_question_attempts_total gauge",
            f"mgc_question_attempts_total {q_total}",
        ]
    except Exception:
        lines += [
            "# HELP mgc_database_metrics_available Whether DB-backed metrics could be collected",
            "# TYPE mgc_database_metrics_available gauge",
            "mgc_database_metrics_available 0",
        ]
    else:
        lines += [
            "# HELP mgc_database_metrics_available Whether DB-backed metrics could be collected",
            "# TYPE mgc_database_metrics_available gauge",
            "mgc_database_metrics_available 1",
        ]
    db_queries = db_query_telemetry_snapshot()
    db_pool = db_pool_snapshot()
    recovery_evidence = recovery_evidence_snapshot()
    backup_age = recovery_evidence["backup"].get("age_minutes")
    restore_age = recovery_evidence["restore_rehearsal"].get("age_days")
    lines += [
        "# HELP mgc_db_query_p95_ms Database query p95 latency in the in-process telemetry window",
        "# TYPE mgc_db_query_p95_ms gauge",
        f"mgc_db_query_p95_ms {db_queries['p95_ms']}",
        "# HELP mgc_db_slow_queries_window Slow queries in the current telemetry window",
        "# TYPE mgc_db_slow_queries_window gauge",
        f"mgc_db_slow_queries_window {db_queries['slow_queries']}",
        "# HELP mgc_db_slow_queries_total Slow queries since process start",
        "# TYPE mgc_db_slow_queries_total counter",
        f"mgc_db_slow_queries_total {db_queries['lifetime_slow_queries']}",
        "# HELP mgc_db_pool_checked_out Checked-out PostgreSQL pool connections",
        "# TYPE mgc_db_pool_checked_out gauge",
        f"mgc_db_pool_checked_out {float(db_pool.get('checked_out') or 0)}",
        "# HELP mgc_db_pool_capacity Configured PostgreSQL pool plus overflow capacity",
        "# TYPE mgc_db_pool_capacity gauge",
        f"mgc_db_pool_capacity {float(db_pool.get('capacity') or 0)}",
        "# HELP mgc_db_pool_saturation_percent PostgreSQL connection-pool saturation",
        "# TYPE mgc_db_pool_saturation_percent gauge",
        f"mgc_db_pool_saturation_percent {float(db_pool.get('saturation_percent') or 0)}",
        "# HELP mgc_recovery_backup_evidence_available Whether a backup artifact is visible to the app",
        "# TYPE mgc_recovery_backup_evidence_available gauge",
        f"mgc_recovery_backup_evidence_available {1 if backup_age is not None else 0}",
        "# HELP mgc_recovery_backup_age_minutes Age of newest observed backup artifact",
        "# TYPE mgc_recovery_backup_age_minutes gauge",
        f"mgc_recovery_backup_age_minutes {float(backup_age) if backup_age is not None else -1}",
        "# HELP mgc_recovery_restore_evidence_available Whether restore rehearsal evidence is visible to the app",
        "# TYPE mgc_recovery_restore_evidence_available gauge",
        f"mgc_recovery_restore_evidence_available {1 if restore_age is not None else 0}",
        "# HELP mgc_recovery_restore_evidence_age_days Age of newest restore rehearsal evidence",
        "# TYPE mgc_recovery_restore_evidence_age_days gauge",
        f"mgc_recovery_restore_evidence_age_days {float(restore_age) if restore_age is not None else -1}",
        "# HELP mgc_recovery_rpo_target_minutes Configured pilot RPO target",
        "# TYPE mgc_recovery_rpo_target_minutes gauge",
        f"mgc_recovery_rpo_target_minutes {RPO_TARGET_MINUTES}",
        "# HELP mgc_recovery_rto_target_minutes Configured pilot RTO target",
        "# TYPE mgc_recovery_rto_target_minutes gauge",
        f"mgc_recovery_rto_target_minutes {RTO_TARGET_MINUTES}",
    ]
    slo = slo_snapshot()
    recovery = _recovery_view()
    state_map = {"healthy":0,"degraded":1,"recovering":2,"unavailable":3}
    lines += [
        "# HELP mgc_slo_availability_percent Rolling pilot availability SLI",
        "# TYPE mgc_slo_availability_percent gauge",
        f"mgc_slo_availability_percent {slo['availability_percent']}",
        "# HELP mgc_slo_error_rate_percent Rolling pilot HTTP 5xx rate",
        "# TYPE mgc_slo_error_rate_percent gauge",
        f"mgc_slo_error_rate_percent {slo['error_rate_percent']}",
        "# HELP mgc_slo_p95_ms Rolling pilot HTTP p95 latency in milliseconds",
        "# TYPE mgc_slo_p95_ms gauge",
        f"mgc_slo_p95_ms {slo['p95_ms']}",
        "# HELP mgc_slo_samples Rolling pilot SLO sample count",
        "# TYPE mgc_slo_samples gauge",
        f"mgc_slo_samples {slo['samples']}",
        "# HELP mgc_recovery_state Pilot recovery state (0 healthy, 1 degraded, 2 recovering, 3 unavailable)",
        "# TYPE mgc_recovery_state gauge",
        f"mgc_recovery_state {state_map.get(recovery.get('state'),3)}",
    ]
    lines += [
        "# HELP mgc_tts_success_total Successful server-side TTS syntheses",
        "# TYPE mgc_tts_success_total counter",
        f"mgc_tts_success_total {tts_runtime['success_total']}",
        "# HELP mgc_tts_failure_total Failed server-side TTS syntheses",
        "# TYPE mgc_tts_failure_total counter",
        f"mgc_tts_failure_total {tts_runtime['failure_total']}",
        "# HELP mgc_tts_disk_cache_hits_total TTS disk cache hits",
        "# TYPE mgc_tts_disk_cache_hits_total counter",
        f"mgc_tts_disk_cache_hits_total {tts_runtime['disk_cache_hits']}",
        "# HELP mgc_tts_circuit_open TTS circuit breaker state",
        "# TYPE mgc_tts_circuit_open gauge",
        f"mgc_tts_circuit_open {1 if _tts_circuit_open() else 0}",
        "# HELP mgc_tts_cache_writable TTS cache writability",
        "# TYPE mgc_tts_cache_writable gauge",
        f"mgc_tts_cache_writable {1 if _tts_health().get('cache_writable') else 0}",
    ]
    return "\n".join(lines) + "\n"


@app.get("/api/meta")
def meta():
    return {"title":"MGC Languages","version":APP_VERSION,"instance_id":INSTANCE_ID,"environment":APP_ENV,"auth_mode":AUTH_MODE,"local_auth_enabled":AUTH_MODE == "local","registration_enabled":REGISTRATION_ENABLED and AUTH_MODE == "local","languages":[{"id":"english","label":"Английский","native":"English"},{"id":"chinese","label":"Китайский","native":"中文"}],"server_tts_available":_tts_health()["server_available"],"voice_recording_enabled":False,"pronunciation_transport":"POST","legacy_tts_get_enabled":TTS_LEGACY_GET_ENABLED,"tts_cache_persistence":TTS_CACHE_PERSISTENCE,"expected_schema_head":EXPECTED_ALEMBIC_HEAD,"reliability_profile":"pilot-v5.7.1","slo_profile":"pilot-operations-v5.5","governance_profile":"pilot-governance-v5.6","security_profile":"rls-content-integrity-v5.7","adaptive_learning":"srs-v5.7","recovery_profile":"observability-recovery-v5.7.1","otel_enabled":OTEL_ENABLED,"chinese_learning_standard":"Путунхуа (普通话) — стандартный китайский"}


@app.post("/api/register")
def register(payload: AuthPayload, request: Request, response: Response, db: Session = Depends(db_session)):
    if AUTH_MODE != "local" or not REGISTRATION_ENABLED:
        raise HTTPException(403, "Самостоятельная регистрация отключена")
    rate_limit(request, "register", 5, 600)
    username = payload.username.strip().lower()
    if not re.fullmatch(r"[a-zA-Z0-9_.-]{3,80}", username):
        raise HTTPException(400, "Логин: латинские буквы, цифры, точка, дефис или подчёркивание")
    if db.scalar(select(User).where(User.username == username)):
        raise HTTPException(409, "Такой логин уже зарегистрирован")
    user = User(username=username, display_name=(payload.display_name or username).strip() or username, password_hash=make_password_hash(payload.password))
    db.add(user); db.flush()
    audit_event(db, "auth.register", actor_user_id=user.id, target_type="user", target_id=str(user.id), request=request)
    db.commit(); db.refresh(user)
    return create_login_session(db, user, response)


@app.post("/api/login")
def login(payload: AuthPayload, request: Request, response: Response, db: Session = Depends(db_session)):
    if AUTH_MODE != "local":
        raise HTTPException(403, "Local auth отключён; используйте корпоративный вход")
    rate_limit(request, "login", 10, 300)
    user = db.scalar(select(User).where(User.username == payload.username.strip().lower()))
    if not user or not verify_password(payload.password, user.password_hash):
        raise HTTPException(401, "Неверный логин или пароль")
    audit_event(db, "auth.login", actor_user_id=user.id, target_type="user", target_id=str(user.id), request=request)
    db.commit()
    return create_login_session(db, user, response)


@app.post("/api/logout")
def logout(request: Request, response: Response, user: User = Depends(current_user), mgc_session: str | None = Cookie(default=None), db: Session = Depends(db_session)):
    if mgc_session:
        db.execute(delete(LoginSession).where(LoginSession.token_hash == token_digest(mgc_session)))
    audit_event(db, "auth.logout", actor_user_id=user.id, target_type="user", target_id=str(user.id), request=request)
    db.commit()
    response.delete_cookie("mgc_session", path="/")
    response.delete_cookie("mgc_csrf", path="/")
    return {"ok":True}


def _oidc_role(groups: list[str]) -> str:
    group_set = set(groups)
    if OIDC_ADMIN_GROUP and OIDC_ADMIN_GROUP in group_set: return "admin"
    if OIDC_EDITOR_GROUP and OIDC_EDITOR_GROUP in group_set: return "editor"
    if OIDC_MANAGER_GROUP and OIDC_MANAGER_GROUP in group_set: return "manager"
    return "user"


@app.get("/api/auth/oidc/login")
async def oidc_login(request: Request):
    if AUTH_MODE != "oidc":
        raise HTTPException(404, "OIDC не включён")
    try:
        from authlib.integrations.starlette_client import OAuth
    except ImportError:
        raise HTTPException(503, "OIDC dependency is not installed")
    oauth = OAuth()
    oauth.register(name="corporate", server_metadata_url=OIDC_DISCOVERY_URL, client_id=OIDC_CLIENT_ID, client_secret=OIDC_CLIENT_SECRET, client_kwargs={"scope":OIDC_SCOPE})
    redirect_uri = str(request.url_for("oidc_callback"))
    return await oauth.corporate.authorize_redirect(request, redirect_uri)


@app.get("/api/auth/oidc/callback", name="oidc_callback")
async def oidc_callback(request: Request, db: Session = Depends(db_session)):
    if AUTH_MODE != "oidc":
        raise HTTPException(404, "OIDC не включён")
    from authlib.integrations.starlette_client import OAuth
    oauth = OAuth()
    oauth.register(name="corporate", server_metadata_url=OIDC_DISCOVERY_URL, client_id=OIDC_CLIENT_ID, client_secret=OIDC_CLIENT_SECRET, client_kwargs={"scope":OIDC_SCOPE})
    token = await oauth.corporate.authorize_access_token(request)
    claims = token.get("userinfo") or await oauth.corporate.userinfo(token=token)
    username = str(claims.get(OIDC_USERNAME_CLAIM) or claims.get("email") or claims.get("sub") or "").strip().lower()
    if not username:
        raise HTTPException(403, "OIDC не вернул идентификатор пользователя")
    display_name = str(claims.get(OIDC_DISPLAY_NAME_CLAIM) or username).strip()[:120]
    groups_raw = claims.get(OIDC_GROUPS_CLAIM) or []
    groups = [str(x) for x in groups_raw] if isinstance(groups_raw, (list, tuple, set)) else [str(groups_raw)]
    role = _oidc_role(groups)
    department = str(claims.get(OIDC_DEPARTMENT_CLAIM) or "General").strip()[:160] or "General"
    user = db.scalar(select(User).where(User.username == username))
    if not user:
        user = User(username=username[:80], display_name=display_name or username[:120], password_hash="OIDC", role=role, department=department)
        db.add(user); db.flush()
    else:
        user.display_name = display_name or user.display_name
        user.role = role
        user.department = department
    audit_event(db, "auth.oidc.login", actor_user_id=user.id, target_type="user", target_id=str(user.id), request=request, metadata={"role":role,"department":department})
    db.commit(); db.refresh(user)
    response = RedirectResponse(url="/")
    create_login_session(db, user, response)
    return response


@app.get("/api/me")
def me(user: User = Depends(current_user)):
    return user_view(user)


@app.post("/api/me/language")
def set_language(payload: PreferencePayload, user: User = Depends(current_user), db: Session = Depends(db_session)):
    language = validate_language(payload.language)
    user.preferred_language = language
    db.add(user)
    db.commit()
    return {"ok": True, "language": language}


def learning_pref(db: Session, user_id: int) -> LearningPreference:
    pref = db.get(LearningPreference, user_id)
    if not pref:
        pref = LearningPreference(user_id=user_id)
        db.add(pref)
        db.flush()
    return pref


@app.get("/api/learning/preferences")
def get_learning_preferences(user: User = Depends(current_user), db: Session = Depends(db_session)):
    pref = learning_pref(db, user.id)
    db.commit()
    return {"show_pinyin": pref.show_pinyin, "show_reading": pref.show_reading, "server_audio_enabled": pref.server_audio_enabled}


@app.put("/api/learning/preferences")
def set_learning_preferences(payload: LearningSettingsPayload, user: User = Depends(current_user), db: Session = Depends(db_session)):
    pref = learning_pref(db, user.id)
    pref.show_pinyin = payload.show_pinyin
    pref.show_reading = payload.show_reading
    pref.server_audio_enabled = payload.server_audio_enabled
    pref.updated_at = datetime.now(timezone.utc)
    db.commit()
    return {"show_pinyin": pref.show_pinyin, "show_reading": pref.show_reading, "server_audio_enabled": pref.server_audio_enabled}


@app.get("/api/chinese/foundations")
def chinese_foundations(user: User = Depends(current_user), db: Session = Depends(db_session)):
    require_feature(db, user.id, "chinese_reference")
    return CHINESE_FOUNDATIONS


@app.get("/api/pilot/me")
def pilot_me(user: User = Depends(current_user), db: Session = Depends(db_session)):
    return pilot_me_view(db, user)


_TTS_SEMAPHORE = threading.BoundedSemaphore(TTS_CONCURRENCY)
_TTS_RUNTIME_LOCK = threading.Lock()
_TTS_RUNTIME: dict[str, Any] = {
    "success_total": 0, "failure_total": 0, "disk_cache_hits": 0, "disk_cache_writes": 0,
    "consecutive_failures": 0, "circuit_open_until": 0.0,
}
_TTS_HEALTH_LOCK = threading.Lock()
_TTS_HEALTH_CACHE: dict[str, Any] = {"expires_at": 0.0, "value": None}


def _tts_cache_path(language: str, text_value: str, rate_key: int) -> Path:
    digest = hashlib.sha256(f"v5.3.2|{language}|{rate_key}|{text_value}".encode("utf-8")).hexdigest()
    return TTS_CACHE_DIR / f"{digest}.wav"


def _tts_cache_read(language: str, text_value: str, rate_key: int) -> bytes | None:
    if not TTS_DISK_CACHE_ENABLED:
        return None
    path = _tts_cache_path(language, text_value, rate_key)
    try:
        if not path.is_file():
            return None
        age = time.time() - path.stat().st_mtime
        if age > TTS_DISK_CACHE_TTL_HOURS * 3600:
            path.unlink(missing_ok=True)
            return None
        data = path.read_bytes()
        if not data.startswith(b"RIFF") or len(data) < 512:
            path.unlink(missing_ok=True)
            return None
        with _TTS_RUNTIME_LOCK:
            _TTS_RUNTIME["disk_cache_hits"] += 1
        return data
    except OSError:
        return None


def _prune_tts_cache() -> None:
    if not TTS_DISK_CACHE_ENABLED:
        return
    try:
        TTS_CACHE_DIR.mkdir(parents=True, exist_ok=True, mode=0o700)
        files = [x for x in TTS_CACHE_DIR.glob("*.wav") if x.is_file()]
        max_bytes = TTS_DISK_CACHE_MAX_MB * 1024 * 1024
        total = sum(x.stat().st_size for x in files)
        if total <= max_bytes:
            return
        target = int(max_bytes * 0.8)
        for item in sorted(files, key=lambda x: x.stat().st_mtime):
            size = item.stat().st_size
            item.unlink(missing_ok=True)
            total -= size
            if total <= target:
                break
    except OSError:
        pass


def _tts_cache_write(language: str, text_value: str, rate_key: int, data: bytes) -> None:
    if not TTS_DISK_CACHE_ENABLED or not data.startswith(b"RIFF"):
        return
    try:
        TTS_CACHE_DIR.mkdir(parents=True, exist_ok=True, mode=0o700)
        path = _tts_cache_path(language, text_value, rate_key)
        tmp = path.with_suffix(".tmp")
        tmp.write_bytes(data)
        os.chmod(tmp, 0o600)
        os.replace(tmp, path)
        with _TTS_RUNTIME_LOCK:
            _TTS_RUNTIME["disk_cache_writes"] += 1
        _prune_tts_cache()
    except OSError:
        pass


def _tts_circuit_open() -> bool:
    with _TTS_RUNTIME_LOCK:
        return time.monotonic() < float(_TTS_RUNTIME["circuit_open_until"])


def _tts_mark_success() -> None:
    with _TTS_RUNTIME_LOCK:
        _TTS_RUNTIME["success_total"] += 1
        _TTS_RUNTIME["consecutive_failures"] = 0
        _TTS_RUNTIME["circuit_open_until"] = 0.0


def _tts_mark_failure() -> None:
    with _TTS_RUNTIME_LOCK:
        _TTS_RUNTIME["failure_total"] += 1
        _TTS_RUNTIME["consecutive_failures"] += 1
        if _TTS_RUNTIME["consecutive_failures"] >= TTS_FAILURE_THRESHOLD:
            _TTS_RUNTIME["circuit_open_until"] = time.monotonic() + TTS_CIRCUIT_COOLDOWN_SECONDS


@lru_cache(maxsize=256)
def _synthesize_wav(language: str, text_value: str, rate_key: int) -> bytes:
    if not TTS_ENABLED or not TTS_BINARY:
        raise RuntimeError("server TTS unavailable")
    cached = _tts_cache_read(language, text_value, rate_key)
    if cached:
        return cached
    if _tts_circuit_open():
        raise RuntimeError("TTS circuit breaker open")
    acquired = _TTS_SEMAPHORE.acquire(timeout=TTS_TIMEOUT_SECONDS)
    if not acquired:
        _tts_mark_failure()
        raise RuntimeError("TTS concurrency limit reached")
    try:
        voice = TTS_VOICE_CHINESE if language == "chinese" else TTS_VOICE_ENGLISH
        speed = max(80, min(240, int(175 * (rate_key / 100))))
        result = subprocess.run(
            [TTS_BINARY, "-v", voice, "-s", str(speed), "--stdout"],
            input=text_value.encode("utf-8"),
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=TTS_TIMEOUT_SECONDS,
            check=False,
        )
        if result.returncode != 0 or not result.stdout.startswith(b"RIFF"):
            error = (result.stderr or b"TTS failed").decode("utf-8", "ignore")[:500]
            _tts_mark_failure()
            raise RuntimeError(error)
        _tts_mark_success()
        _tts_cache_write(language, text_value, rate_key, result.stdout)
        return result.stdout
    except subprocess.SubprocessError:
        _tts_mark_failure()
        raise
    finally:
        _TTS_SEMAPHORE.release()


def _tts_health(force: bool = False) -> dict[str, Any]:
    now = time.monotonic()
    with _TTS_HEALTH_LOCK:
        cached = _TTS_HEALTH_CACHE.get("value")
        if not force and cached is not None and now < float(_TTS_HEALTH_CACHE.get("expires_at", 0)):
            return dict(cached)
    engine = Path(TTS_BINARY).name if TTS_BINARY else "browser-fallback"
    result: dict[str, Any] = {
        "server_available": False, "engine": engine, "offline": False,
        "voices": {"chinese": TTS_VOICE_CHINESE, "english": TTS_VOICE_ENGLISH},
        "language_checks": {"chinese": False, "english": False},
        "circuit_open": _tts_circuit_open(),
        "timeout_seconds": TTS_TIMEOUT_SECONDS,
        "concurrency": TTS_CONCURRENCY,
        "disk_cache_enabled": TTS_DISK_CACHE_ENABLED,
        "cache_persistence": TTS_CACHE_PERSISTENCE,
        "cache_writable": False,
    }
    if TTS_DISK_CACHE_ENABLED:
        try:
            TTS_CACHE_DIR.mkdir(parents=True, exist_ok=True, mode=0o700)
            probe = TTS_CACHE_DIR / ".write-probe"
            probe.write_bytes(b"ok")
            os.chmod(probe, 0o600)
            probe.unlink(missing_ok=True)
            result["cache_writable"] = True
        except OSError:
            result["cache_writable"] = False
    else:
        result["cache_writable"] = True
    if TTS_ENABLED and TTS_BINARY and not result["circuit_open"]:
        samples = {"chinese": "你好", "english": "hello"}
        for language, sample in samples.items():
            try:
                wav = _synthesize_wav(language, sample, 90)
                result["language_checks"][language] = wav.startswith(b"RIFF") and len(wav) > 512
            except Exception as exc:
                logger.warning(json.dumps({"event":"tts_probe_failed","language":language,"error":str(exc)[:200]}, ensure_ascii=False))
        result["server_available"] = all(result["language_checks"].values())
        result["offline"] = result["server_available"]
    with _TTS_RUNTIME_LOCK:
        result["runtime"] = dict(_TTS_RUNTIME)
    with _TTS_HEALTH_LOCK:
        _TTS_HEALTH_CACHE["value"] = dict(result)
        _TTS_HEALTH_CACHE["expires_at"] = now + TTS_HEALTH_TTL_SECONDS
    return result


@app.get("/api/pronunciation/status")
def pronunciation_status(user: User = Depends(current_user)):
    return _tts_health()


def _pronunciation_response(request: Request, language: str, text_value: str, rate: float) -> StreamingResponse:
    language = validate_language(language)
    rate_limit(request, "tts", 90, 60)
    cleaned = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", " ", text_value).strip()
    if not cleaned:
        raise HTTPException(400, "Нет текста для произношения")
    health = _tts_health()
    if not health.get("server_available") or not health.get("language_checks", {}).get(language):
        operational_event("tts.fallback", severity="warning", component="tts", status_code=503, request=request, metadata={"language":language,"circuit_open":_tts_circuit_open()}, throttle_seconds=60)
        raise HTTPException(503, "Серверное произношение недоступно; используйте browser fallback")
    try:
        wav = _synthesize_wav(language, cleaned, int(round(rate * 100)))
    except (subprocess.SubprocessError, RuntimeError) as exc:
        logger.warning(json.dumps({"event":"tts_failed","language":language,"error":str(exc)[:200]}, ensure_ascii=False))
        operational_event("tts.failure", severity="error", component="tts", status_code=503, request=request, detail=type(exc).__name__, metadata={"language":language,"circuit_open":_tts_circuit_open()}, throttle_seconds=10)
        raise HTTPException(503, "Серверное произношение временно недоступно")
    return StreamingResponse(
        io.BytesIO(wav), media_type="audio/wav",
        headers={
            "Content-Disposition":"inline; filename=pronunciation.wav",
            "X-TTS-Engine":Path(TTS_BINARY).name,
            "X-TTS-Voice": TTS_VOICE_CHINESE if language == "chinese" else TTS_VOICE_ENGLISH,
            "Cache-Control":"no-store",
            "X-Content-Type-Options":"nosniff",
        },
    )


@app.post("/api/pronunciation/audio")
def pronunciation_audio_post(payload: PronunciationPayload, request: Request, user: User = Depends(current_user), db: Session = Depends(db_session)):
    # POST keeps training text out of URLs and reverse-proxy access logs.
    require_feature(db, user.id, "server_audio")
    _consume_daily_quota(db, user.id, "tts_requests", PILOT_DAILY_TTS_CAP)
    db.commit()
    return _pronunciation_response(request, payload.language, payload.text, payload.rate)


if TTS_LEGACY_GET_ENABLED:
    @app.get("/api/pronunciation/audio", deprecated=True)
    def pronunciation_audio_legacy(
        request: Request,
        language: str = Query(...),
        text_value: str = Query(..., alias="text", min_length=1, max_length=TTS_MAX_CHARS),
        rate: float = Query(default=0.9, ge=0.55, le=1.25),
        user: User = Depends(current_user),
    ):
        response = _pronunciation_response(request, language, text_value, rate)
        response.headers["Deprecation"] = "true"
        response.headers["Sunset"] = "v5.4"
        return response


@app.get("/api/language/{language}/summary")
def language_summary(language: str, user: User = Depends(current_user), db: Session = Depends(db_session)):
    language = validate_language(language)
    effective_terms = terms_for(language, db)
    rows = db.scalars(select(TermProgress).where(
        TermProgress.user_id == user.id, TermProgress.language == language
    )).all()
    known = sum(1 for row in rows if row.status == "known")
    completed_days = len(db.scalars(select(CourseDayResult).where(
        CourseDayResult.user_id == user.id, CourseDayResult.language == language,
    )).all())
    exam_rows = db.scalars(select(ExamResult).where(
        ExamResult.user_id == user.id, ExamResult.language == language,
    )).all()
    best_exam = max((row.score for row in exam_rows), default=0)
    name = "английский" if language == "english" else "китайский"
    return {
        "language": language,
        "label": "Английский" if language == "english" else "Китайский",
        "learning_standard": None if language == "english" else {"name":"Путунхуа (普通话)","label":"стандартный китайский","dialects_reference_only":True},
        "level_labels": {
            "A1": f"Базовый {name}",
            "A2": f"Начальный {name}",
            "B1": f"Рабочий {name}",
            "B2": f"Продвинутый {name}",
            "C1": f"Экспертный {name}",
        },
        "progress": {
            "known": known, "touched": len(rows),
            "percent": round(known * 100 / max(1, len(effective_terms))),
            "completed_days": completed_days,
            "course_percent": round(completed_days * 100 / 30),
            "best_exam": best_exam,
            "exam_passed": best_exam >= 35,
        },
    }


@app.get("/api/language/{language}/topics")
def language_topics(language: str, user: User = Depends(current_user), db: Session = Depends(db_session)):
    language = validate_language(language)
    effective_terms = terms_for(language, db)
    counts = Counter(item["topic"] for item in effective_terms)
    descriptions = {row["label"]: row["description"] for row in TOPIC_ROWS}
    labels = [row["label"] for row in TOPIC_ROWS if counts.get(row["label"], 0)]
    labels.extend(sorted(label for label in counts if label not in descriptions))
    return [
        {
            "id": str(index + 1), "label": label,
            "description": descriptions.get(label, "Корпоративная тема, добавленная администратором"),
            "count": counts.get(label, 0),
            "examples": [item["term"] for item in effective_terms if item["topic"] == label][:3],
        }
        for index, label in enumerate(labels)
    ]


@app.get("/api/language/{language}/terms")
def language_terms(
    language: str, topic: str | None = None, level: str | None = None,
    offset: int = Query(default=0, ge=0), limit: int = Query(default=60, ge=1, le=500),
    user: User = Depends(current_user), db: Session = Depends(db_session),
):
    language = validate_language(language)
    items = terms_for(language, db)
    if topic:
        items = [item for item in items if item["topic"] == topic]
    if level in LEVELS:
        items = [item for item in items if item["level"] == level]
    return {"total": len(items), "offset": offset, "items": items[offset:offset + limit]}


def make_question(item: dict[str, Any], pool: list[dict[str, Any]], rng: random.Random, language: str):
    level = item["level"]
    if level == "A1":
        prompt = f"Что означает «{item['term']}»?"
        option_value = lambda row: row["translation"]
        pronunciation = item["pronunciation"]
    elif level == "A2":
        prompt = f"Как сказать «{item['translation']}»?"
        option_value = lambda row: row["term"]
        pronunciation = item["pronunciation"]
    elif level == "B1":
        source = item.get("example") or item["term"]
        prompt = f"Выберите точный перевод рабочей фразы: «{source}»"
        option_value = lambda row: row.get("example_translation") or row["translation"]
        pronunciation = item.get("example_pronunciation") or item["pronunciation"]
    else:
        situation = item.get("example_translation") or item["translation"]
        prompt = f"Какая формулировка точнее всего выражает: «{situation}»?"
        option_value = lambda row: row.get("example") or row["term"]
        pronunciation = item.get("example_pronunciation") or item["pronunciation"]

    correct = option_value(item)
    candidates = [
        candidate for candidate in pool
        if candidate["id"] != item["id"] and option_value(candidate) != correct
    ]
    same_topic = [candidate for candidate in candidates if candidate["topic"] == item["topic"]]
    rng.shuffle(same_topic)
    rng.shuffle(candidates)
    wrong = []
    seen = {correct}
    for candidate in [*same_topic, *candidates]:
        value = option_value(candidate)
        if value in seen:
            continue
        seen.add(value)
        wrong.append(value)
        if len(wrong) == 3:
            break
    if len(wrong) < 3:
        raise HTTPException(400, "Недостаточно уникальных вариантов для выбранной темы")
    options = [correct, *wrong]
    rng.shuffle(options)
    return {
        "id": item["id"], "level": item["level"], "topic": item["topic"],
        "prompt": prompt,
        "pronunciation": pronunciation if language == "chinese" else item.get("pronunciation", ""),
        "reading": item.get("reading", ""),
        "audio_text": (item.get("example") or item["term"]) if level in {"B1", "B2", "C1"} else item["term"],
        "options": options, "correct_index": options.index(correct),
        "explanation": (
            f"{item['term']} — {item['translation']}. "
            f"{item.get('example') or ''} — {item.get('example_translation') or ''}"
        ).strip(" —."),
    }


@app.get("/api/language/{language}/quiz")
def language_quiz(
    language: str, topic: str | None = None, level: str | None = None,
    count: int = Query(default=10, ge=5, le=30), user: User = Depends(current_user), db: Session = Depends(db_session),
):
    language = validate_language(language)
    pool = terms_for(language, db)
    if topic:
        pool = [item for item in pool if item["topic"] == topic]
    topic_pool = pool
    if level in LEVELS:
        pool = [item for item in pool if item["level"] == level]
    if len(pool) < 4 and len(topic_pool) >= 4:
        pool = topic_pool
    if len(pool) < 4:
        raise HTTPException(400, "В выбранной теме недостаточно терминов")
    rng = random.Random(secrets.randbits(64))
    targets = rng.sample(pool, min(count, len(pool)))
    return {"questions": [make_question(item, pool, rng, language) for item in targets]}


@app.get("/api/language/{language}/situations")
def situations(language: str, user: User = Depends(current_user)):
    language = validate_language(language)
    if language == "chinese":
        return [{
            "id": f"sit-zh-{row['id']}", "topic": topic_for(row.get("topic", "")),
            "source_topic": row.get("topic", ""), "title": row.get("ru", ""),
            "target": row.get("zh", ""), "pronunciation": row.get("pinyin", ""),
            "translation": row.get("ru", ""), "note": row.get("note", ""),
        } for row in CHINESE_RAW.get("situations", [])]
    return [{
        "id": f"sit-en-{row['id']}", "topic": row["topic"], "source_topic": row["topic"],
        "title": row["title"], "target": row["en"], "pronunciation": english_pronunciation(row["en"])["ipa"],
        "reading": english_pronunciation(row["en"])["reading"],
        "translation": row["en_ru"], "note": f"Цель: {row['goal']}",
    } for row in APP_CONTENT["roleplays"]]


@app.get("/api/language/{language}/roleplays")
def roleplays(language: str, user: User = Depends(current_user)):
    language = validate_language(language)
    rng = random.Random(f"mgc-scenario-v4:{user.id}:{language}")
    items = []
    for row in ROLEPLAYS:
        distractor_pool = [candidate for candidate in ROLEPLAYS if candidate["id"] != row["id"]]
        distractors = rng.sample(distractor_pool, 2)
        choices = []
        for candidate in [row, *distractors]:
            target_text = candidate["en"] if language == "english" else candidate["zh"]
            choices.append({
                "text": target_text,
                "pronunciation": english_pronunciation(target_text)["ipa"] if language == "english" else candidate["pinyin"],
                "reading": english_pronunciation(target_text)["reading"] if language == "english" else pinyin_to_ru_approx(candidate["pinyin"]),
                "translation": candidate["en_ru"] if language == "english" else candidate["zh_ru"],
                "correct": candidate["id"] == row["id"],
            })
        rng.shuffle(choices)
        if language == "english":
            question = "A colleague asks you to respond to this issue. Which answer is the most professional?"
            question_pronunciation = english_pronunciation(question)["ipa"]
            question_reading = english_pronunciation(question)["reading"]
        else:
            question = "同事请你处理这个问题。哪一个回答最专业？"
            question_pronunciation = "tóngshì qǐng nǐ chǔlǐ zhè ge wèntí. nǎ yí ge huídá zuì zhuānyè"
            question_reading = "тун-ши цин ни чу-ли чжэ гэ вэнь-ти; на и гэ хуэй-да цзуй чжуань-е"
        items.append({
            "id": row["id"], "topic": topic_for(row["topic"], row["title"]), "title": row["title"], "goal": row["goal"],
            "roles": row["roles"], "question": question,
            "question_pronunciation": question_pronunciation, "question_reading": question_reading,
            "question_translation": f"Ситуация: {row['title']}. Выберите наиболее профессиональный ответ.",
            "options": choices, "prompts": row["prompts"],
        })
    return items


@app.get("/api/language/{language}/mgc-scenarios")
def mgc_scenarios(language: str, user: User = Depends(current_user)):
    language = validate_language(language)
    return [{
        "id": row["id"], "topic": row["topic"], "title": row["title"],
        "target": row["en"] if language == "english" else row["zh"],
        "pronunciation": english_pronunciation(row["en"])["ipa"] if language == "english" else row["pinyin"],
        "reading": english_pronunciation(row["en"])["reading"] if language == "english" else pinyin_to_ru_approx(row["pinyin"]),
        "translation": row["en_ru"] if language == "english" else row["zh_ru"],
        "usage": row["usage"],
    } for row in APP_CONTENT["mgc_scenarios"]]


@app.get("/api/language/{language}/course30")
def course30(language: str, user: User = Depends(current_user), db: Session = Depends(db_session)):
    language = validate_language(language)
    effective_terms = terms_for(language, db)
    topic_labels = [row["label"] for row in TOPIC_ROWS] + sorted({row["topic"] for row in effective_terms if row["topic"] not in {t["label"] for t in TOPIC_ROWS}})
    topics = [
        label for label in topic_labels
        if sum(1 for term in effective_terms if term["topic"] == label) >= 5
    ]
    level_plan = ["A1"] * 5 + ["A2"] * 6 + ["B1"] * 7 + ["B2"] * 6 + ["C1"] * 6
    saved_results = {
        row.day: row for row in db.scalars(select(CourseDayResult).where(
            CourseDayResult.user_id == user.id, CourseDayResult.language == language,
        )).all()
    }
    days = []
    for index in range(30):
        topic = topics[index % len(topics)]
        level = level_plan[index]
        pool = [row for row in effective_terms if row["topic"] == topic and row["level"] == level]
        if len(pool) < 5:
            pool = [row for row in effective_terms if row["topic"] == topic] or effective_terms
        start = (index * 5) % len(pool)
        selection = (pool[start:] + pool[:start])[:5]
        question_pool = [row for row in effective_terms if row["topic"] == topic]
        if len(question_pool) < 4:
            question_pool = effective_terms
        rng = random.Random(f"mgc-course-v4:{user.id}:{language}:{index + 1}")
        quiz = [make_question(row, question_pool, rng, language) for row in selection]
        language_name = "английском" if language == "english" else "китайском"
        result = saved_results.get(index + 1)
        days.append({
            "day": index + 1, "level": level, "topic": topic, "title": f"День {index + 1}: {topic}",
            "task": f"Разберите 5 терминов, соберите пары, решите мини-кейс и пройдите тест на {language_name}.",
            "terms": selection, "quiz": quiz, "completed": result is not None,
            "best_score": result.score if result else 0, "total": 5,
            "activities": ["Разбор терминов", "Собрать пары", "Мини-кейс", "Тест из 5 вопросов"],
        })
    return {"language": language, "completed_days": len(saved_results), "days": days}


@app.post("/api/course-day/result")
def save_course_day(payload: CourseDayPayload, user: User = Depends(current_user), db: Session = Depends(db_session)):
    language = validate_language(payload.language)
    effective_terms = terms_for(language, db)
    if payload.score < 4:
        return {"ok": True, "day": payload.day, "score": payload.score, "total": payload.total, "completed": False}
    row = db.scalar(select(CourseDayResult).where(
        CourseDayResult.user_id == user.id,
        CourseDayResult.language == language,
        CourseDayResult.day == payload.day,
    ))
    if row:
        row.score = max(row.score, payload.score)
        row.total = payload.total
        row.completed_at = datetime.now(timezone.utc)
    else:
        row = CourseDayResult(
            user_id=user.id, language=language, day=payload.day,
            score=payload.score, total=payload.total,
        )
        db.add(row)
    topics = [
        topic_row["label"] for topic_row in TOPIC_ROWS
        if sum(1 for term in effective_terms if term["topic"] == topic_row["label"]) >= 5
    ]
    level_plan = ["A1"] * 5 + ["A2"] * 6 + ["B1"] * 7 + ["B2"] * 6 + ["C1"] * 6
    topic = topics[(payload.day - 1) % len(topics)]
    level = level_plan[payload.day - 1]
    pool = [term for term in effective_terms if term["topic"] == topic and term["level"] == level]
    if len(pool) < 5:
        pool = [term for term in effective_terms if term["topic"] == topic]
    start = ((payload.day - 1) * 5) % len(pool)
    selection = (pool[start:] + pool[:start])[:5]
    for term in selection:
        progress = db.scalar(select(TermProgress).where(
            TermProgress.user_id == user.id,
            TermProgress.language == language,
            TermProgress.term_id == term["id"],
        ))
        if progress:
            progress.status = "known"
            progress.updated_at = datetime.now(timezone.utc)
        else:
            db.add(TermProgress(
                user_id=user.id, language=language, term_id=term["id"], status="known",
            ))
    db.commit()
    return {"ok": True, "day": payload.day, "score": row.score, "total": row.total, "completed": True}


@app.get("/api/language/{language}/final-exam")
def final_exam(language: str, user: User = Depends(current_user), db: Session = Depends(db_session)):
    language = validate_language(language)
    effective_terms = terms_for(language, db)
    rng = random.Random(f"mgc-final-v3:{user.id}:{language}")
    questions = []
    for level in LEVELS:
        pool = [row for row in effective_terms if row["level"] == level]
        for row in rng.sample(pool, 10):
            questions.append(make_question(row, pool, rng, language))
    rng.shuffle(questions)
    return {"language": language, "total": 50, "levels": {level: 10 for level in LEVELS}, "questions": questions}


@app.post("/api/final-exam/result")
def save_exam(payload: ExamPayload, user: User = Depends(current_user), db: Session = Depends(db_session)):
    language = validate_language(payload.language)
    db.add(ExamResult(
        user_id=user.id, language=language, score=payload.score, total=payload.total,
        details=json.dumps(payload.answers, ensure_ascii=False),
    ))
    db.commit()
    return {"ok": True, "score": payload.score, "total": payload.total, "passed": payload.score >= 35}


@app.get("/api/language/{language}/progress")
def get_progress(language: str, user: User = Depends(current_user), db: Session = Depends(db_session)):
    language = validate_language(language)
    rows = db.scalars(select(TermProgress).where(
        TermProgress.user_id == user.id, TermProgress.language == language
    )).all()
    return {row.term_id: row.status for row in rows}


@app.post("/api/language/{language}/progress")
def set_progress(
    language: str, payload: ProgressPayload, user: User = Depends(current_user),
    db: Session = Depends(db_session),
):
    language = validate_language(language)
    row = db.scalar(select(TermProgress).where(
        TermProgress.user_id == user.id, TermProgress.language == language,
        TermProgress.term_id == payload.term_id,
    ))
    if row:
        row.status = payload.status
        row.updated_at = datetime.now(timezone.utc)
    else:
        db.add(TermProgress(user_id=user.id, language=language, term_id=payload.term_id, status=payload.status))
    term = term_by_id(language, payload.term_id, db)
    if term:
        card = get_or_create_srs_card(db, user.id, language, payload.term_id, term.get("topic", ""))
        if payload.status == "known" and card.repetitions == 0:
            schedule_srs(card, 4)
    db.commit()
    return {"ok": True}


def search_score(query: str, item: dict[str, Any]) -> int:
    haystack = " ".join(str(item.get(key, "")) for key in (
        "term", "translation", "topic", "category", "example", "example_translation"
    )).lower()
    query = query.lower().strip()
    if query in haystack:
        return 20 + len(query)
    tokens = [token for token in re.findall(r"[\w\u4e00-\u9fff]+", query) if len(token) > 1]
    return sum(3 for token in tokens if token in haystack)


@app.get("/api/review/queue")
def review_queue(language: str = Query(default="chinese"), limit: int = Query(default=20, ge=1, le=100), user: User = Depends(current_user), db: Session = Depends(db_session)):
    language=validate_language(language); now=datetime.now(timezone.utc)
    cards=db.scalars(select(SRSCard).where(SRSCard.user_id==user.id,SRSCard.language==language,SRSCard.due_at<=now).order_by(SRSCard.due_at.asc()).limit(min(limit,SRS_DAILY_REVIEW_LIMIT))).all()
    items=[]
    for card in cards:
        term=term_by_id(language,card.term_id,db)
        if term: items.append({"card_id":card.id,"due_at":card.due_at.isoformat(),"interval_days":card.interval_days,"lapses":card.lapses,"ease_factor_pct":card.ease_factor_pct,"term":term})
    return {"language":language,"due":len(items),"items":items,"algorithm":"SM2-inspired adaptive SRS"}


@app.post("/api/review/result")
def review_result(payload: SRSReviewPayload, user: User = Depends(current_user), db: Session = Depends(db_session)):
    language=validate_language(payload.language); term=term_by_id(language,payload.term_id,db)
    if not term: raise HTTPException(404,"Термин не найден")
    card=get_or_create_srs_card(db,user.id,language,payload.term_id,term.get("topic", "")); schedule_srs(card,payload.quality)
    progress=db.scalar(select(TermProgress).where(TermProgress.user_id==user.id,TermProgress.language==language,TermProgress.term_id==payload.term_id))
    status="known" if payload.quality>=4 else "learning"
    if progress: progress.status=status; progress.updated_at=datetime.now(timezone.utc)
    else: db.add(TermProgress(user_id=user.id,language=language,term_id=payload.term_id,status=status))
    xp=award_xp(db,user.id,"srs_review",payload.term_id,5 if payload.quality>=4 else 2,language=language,topic=term.get("topic", ""),anti_farm=True)
    db.commit(); return {"ok":True,"next_due_at":card.due_at.isoformat(),"interval_days":card.interval_days,"ease_factor_pct":card.ease_factor_pct,"xp":xp.get("awarded",0)}


@app.post("/api/learning/question-attempt")
def question_attempt(payload: QuestionAttemptPayload, user: User = Depends(current_user), db: Session = Depends(db_session)):
    language=validate_language(payload.language)
    if payload.session_id:
        existing=db.scalar(select(QuestionAttempt).where(QuestionAttempt.user_id==user.id,QuestionAttempt.session_id==payload.session_id,QuestionAttempt.question_id==payload.question_id))
        if existing: return {"ok":True,"duplicate":True}
    _consume_daily_quota(db,user.id,"practice_submissions",PILOT_DAILY_PRACTICE_CAP)
    selected_hash=hashlib.sha256(payload.selected.encode("utf-8")).hexdigest() if payload.selected else ""
    row=QuestionAttempt(user_id=user.id,session_id=payload.session_id,question_id=payload.question_id,term_id=payload.term_id,language=language,topic=payload.topic,kind=payload.kind,correct=payload.correct,response_ms=payload.response_ms,selected_hash=selected_hash)
    db.add(row)
    if payload.term_id:
        term=term_by_id(language,payload.term_id,db)
        if term:
            card=get_or_create_srs_card(db,user.id,language,payload.term_id,payload.topic or term.get("topic", ""))
            schedule_srs(card,5 if payload.correct else 1)
    db.commit(); return {"ok":True}


@app.get("/api/admin/learning/question-quality")
def admin_question_quality(language: str | None = Query(default=None), user: User = Depends(require_roles("admin","editor")), db: Session = Depends(db_session)):
    # Explicit privileged aggregate: Editor receives content-level statistics only, never user-level rows in the response.
    apply_rls_context(db, system_admin=True)
    query=select(QuestionAttempt.question_id,QuestionAttempt.language,QuestionAttempt.topic,QuestionAttempt.kind,func.count(QuestionAttempt.id),func.sum(func.cast(QuestionAttempt.correct,Integer)),func.avg(QuestionAttempt.response_ms)).group_by(QuestionAttempt.question_id,QuestionAttempt.language,QuestionAttempt.topic,QuestionAttempt.kind)
    if language in LANGUAGES: query=query.where(QuestionAttempt.language==language)
    rows=db.execute(query).all(); items=[]
    for qid,lang,topic,kind,total,correct_sum,avg_ms in rows:
        total=int(total or 0); correct=int(correct_sum or 0); accuracy=correct/max(1,total)
        flag=""
        if total>=QUESTION_QUALITY_MIN_ATTEMPTS:
            if accuracy<=QUESTION_QUALITY_LOW_ACCURACY: flag="review_low_accuracy"
            elif accuracy>=QUESTION_QUALITY_HIGH_ACCURACY: flag="review_too_easy"
        items.append({"question_id":qid,"language":lang,"topic":topic,"kind":kind,"attempts":total,"accuracy_percent":round(accuracy*100,1),"avg_response_ms":round(float(avg_ms or 0),1),"flag":flag})
    items.sort(key=lambda x:(x["flag"]=="", -x["attempts"]))
    return {"min_attempts":QUESTION_QUALITY_MIN_ATTEMPTS,"items":items[:500],"note":"Quality signal for content review; not an employee HR rating."}


@app.get("/api/admin/audit/verify-chain")
def admin_verify_audit_chain(user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    rows=db.scalars(select(AuditLog).order_by(AuditLog.id.asc())).all(); anchor=db.get(AuditAnchor,1); previous=(anchor.last_deleted_hash if anchor else ""); broken=[]
    for row in rows:
        canonical="|".join([previous,row.event_type,str(row.actor_user_id or ""),row.target_type,row.target_id,row.request_id,row.source_ip_hash,row.metadata_json])
        expected=hashlib.sha256(canonical.encode("utf-8")).hexdigest()
        if (row.previous_hash or "")!=previous or (row.event_hash or "")!=expected: broken.append(row.id)
        previous=row.event_hash or expected
    return {"ok":not broken,"events":len(rows),"broken_ids":broken[:50],"algorithm":"sha256 hash chain"}


@app.post("/api/assistant")
def assistant(payload: AssistantPayload, user: User = Depends(current_user), db: Session = Depends(db_session)):
    require_feature(db, user.id, "ai_assistant")
    language = validate_language(payload.language)
    effective_terms = terms_for(language, db)
    ranked = sorted(
        ((search_score(payload.question, row), row) for row in effective_terms),
        key=lambda pair: pair[0], reverse=True,
    )
    evidence = [row for score, row in ranked[:6] if score > 0]
    if evidence:
        answer = "Нашёл наиболее близкие термины и рабочие формулировки:"
    else:
        evidence = [row for row in effective_terms if row["level"] in {"A1", "A2"}][:5]
        answer = "Прямого совпадения нет. Начните с этих базовых формулировок:"
    return {
        "answer": answer, "evidence": evidence,
        "note": "Ответ формируется по локальной базе платформы и не отправляет данные во внешние сервисы.",
    }


@app.get("/api/knowledge")
def knowledge(language: str = Query(default="chinese"), user: User = Depends(current_user)):
    language = validate_language(language)
    return [{
        "id": row["id"], "topic": row["topic"], "title": row["title"],
        "situation": row["situation"], "steps": row["steps"], "avoid": row["avoid"],
        "target": row["en"] if language == "english" else row["zh"],
        "pronunciation": "" if language == "english" else row["pinyin"],
        "translation": row["en_ru"] if language == "english" else row["zh_ru"],
    } for row in EXPERIENCE["knowledge"]]


@app.get("/api/gamification/me")
def gamification_me(user: User = Depends(current_user), db: Session = Depends(db_session)):
    result = gamification_view(db, user.id)
    db.commit()
    return result


@app.get("/api/gamification/levels")
def gamification_levels(user: User = Depends(current_user)):
    return [{**level_info((level - 1) * 500), "required_xp": (level - 1) * 500} for level in range(1, 101)]


@app.get("/api/gamification/rewards")
def gamification_rewards(user: User = Depends(current_user), db: Session = Depends(db_session)):
    profile = gamification_view(db, user.id)
    db.commit()
    return {"balance": profile["spendable_xp"], "items": [{"id": key, **value} for key, value in REWARD_CATALOG.items()]}


def reward_result(reward_id: str, language: str, context: dict[str, Any], db: Session, user: User) -> dict[str, Any]:
    term_id = str(context.get("term_id") or "")
    term = term_by_id(language, term_id, db) if term_id else None
    if reward_id in {"hint_small", "show_pinyin", "slow_audio", "eliminate_option", "sentence_start", "explain_word", "pronunciation_breakdown", "work_example", "mistake_explain"} and not term:
        raise HTTPException(400, "Для этой помощи нужен термин из текущего задания")
    if reward_id == "hint_small":
        clue_source = term["translation"] if term else ""
        clue = (clue_source[:2] + "…") if len(clue_source) > 2 else clue_source
        return {"type": "hint", "text": f"Тема: {term['topic']}. Подсказка к значению: «{clue}»"}
    if reward_id == "show_pinyin":
        return {"type": "pronunciation", "text": term.get("pronunciation") or "Для этого термина транскрипция не требуется."}
    if reward_id == "slow_audio":
        return {"type": "audio", "text": term.get("example") or term["term"], "rate": 0.68, "language": language}
    if reward_id == "eliminate_option":
        options = [str(x) for x in context.get("options", [])]
        correct_values = {term.get("term", ""), term.get("translation", ""), term.get("example", ""), term.get("example_translation", "")}
        wrong = [option for option in options if option not in correct_values]
        if not wrong:
            raise HTTPException(400, "Не удалось определить неверный вариант")
        return {"type": "eliminate", "option": wrong[0]}
    if reward_id == "sentence_start":
        phrase = term.get("example") or term["term"]
        size = max(1, len(phrase) // 3)
        return {"type": "sentence_start", "text": phrase[:size] + "…"}
    if reward_id == "explain_word":
        return {"type": "explanation", "text": f"{term['term']} — {term['translation']}. Тема: {term['topic']}."}
    if reward_id == "pronunciation_breakdown":
        value = term.get("pronunciation") or term["term"]
        return {"type": "pronunciation", "text": " · ".join(value.split()) if " " in value else value}
    if reward_id == "work_example":
        return {"type": "example", "target": term.get("example") or term["term"], "translation": term.get("example_translation") or term["translation"]}
    if reward_id == "mistake_explain":
        return {"type": "mistake", "text": f"Правильная связь: {term['term']} → {term['translation']}. В рабочем контексте: {term.get('example') or term['term']} → {term.get('example_translation') or term['translation']}."}
    if reward_id in {"smart_revision", "focus_mode"}:
        limit = 10 if reward_id == "focus_mode" else 7
        progress_rows = db.scalars(select(TermProgress).where(
            TermProgress.user_id == user.id, TermProgress.language == language
        ).order_by(TermProgress.updated_at.asc())).all()
        ordered_ids = [row.term_id for row in progress_rows if row.status == "learning"] + [row.term_id for row in progress_rows if row.status == "known"]
        pool = terms_for(language, db)
        by_id = {row["id"]: row for row in pool}
        selected = [by_id[item_id] for item_id in ordered_ids if item_id in by_id][:limit]
        if len(selected) < limit:
            used = {row["id"] for row in selected}
            selected.extend([row for row in pool if row["id"] not in used][:limit-len(selected)])
        return {"type": "training_pack", "title": "Focus Mode" if reward_id == "focus_mode" else "Smart Revision", "items": selected}
    if reward_id == "workshop_pack":
        requested_topic = str(context.get("topic") or "")
        requested_shop = str(context.get("shop") or "")
        pool = terms_for(language, db)
        if requested_topic:
            pool = [row for row in pool if row["topic"] == requested_topic]
        custom_shop_ids = {row.public_id for row in db.scalars(select(CustomTerm).where(
            CustomTerm.language == language, CustomTerm.status == "published", CustomTerm.shop == requested_shop
        )).all()} if requested_shop else set()
        if requested_shop and custom_shop_ids:
            pool = [row for row in pool if row["id"] in custom_shop_ids]
        if not pool:
            raise HTTPException(404, "В выбранной теме/цехе пока нет терминов")
        rng = random.Random(f"pack:{user.id}:{language}:{requested_topic}:{requested_shop}:{date.today().isoformat()}")
        selected = rng.sample(pool, min(10, len(pool)))
        return {"type": "training_pack", "title": requested_topic or requested_shop or "Тренировка по цеху", "items": selected}
    if reward_id == "roleplay":
        items = roleplays(language, user)
        rng = random.Random(f"paid-roleplay:{user.id}:{date.today().isoformat()}")
        return {"type": "roleplay_pack", "items": rng.sample(items, min(3, len(items)))}
    if reward_id == "meeting_simulator":
        items = roleplays(language, user)
        meeting_items = [row for row in items if any(key in (row.get("title", "") + row.get("goal", "")).lower() for key in ("meeting", "совещ", "решен", "эскал"))] or items
        rng = random.Random(f"meeting:{user.id}:{date.today().isoformat()}")
        return {"type": "meeting_pack", "items": rng.sample(meeting_items, min(5, len(meeting_items)))}
    raise HTTPException(400, "Эта награда пока не поддерживается")


@app.post("/api/gamification/spend")
def gamification_spend(payload: SpendPayload, request: Request, user: User = Depends(current_user), db: Session = Depends(db_session)):
    rate_limit(request, "xp_spend", 60, 60)
    require_feature(db, user.id, "xp_economy")
    language = validate_language(payload.language)
    # Сначала строим результат, чтобы не списывать XP за некорректный контекст.
    result = reward_result(payload.reward_id, language, payload.context, db, user)
    spent = spend_xp(db, user.id, payload.reward_id, language=language, metadata={"context": payload.context})
    db.commit()
    return {**spent, "result": result}


@app.post("/api/practice/result")
def save_practice_result(payload: PracticePayload, user: User = Depends(current_user), db: Session = Depends(db_session)):
    language = validate_language(payload.language)
    existing = db.scalar(select(PracticeResult).where(PracticeResult.session_id == payload.session_id))
    if existing:
        return {"ok": True, "duplicate": True, "profile": gamification_view(db, user.id)}
    _consume_daily_quota(db, user.id, "practice_submissions", PILOT_DAILY_PRACTICE_CAP)
    if payload.score > payload.total:
        raise HTTPException(400, "Некорректный результат")
    db.add(PracticeResult(
        user_id=user.id, session_id=payload.session_id, kind=payload.kind, language=language,
        topic=payload.topic, score=payload.score, total=payload.total,
    ))
    ratio = payload.score / max(1, payload.total)
    if payload.kind == "quiz":
        raw = 20 + (10 if ratio >= .9 else 0) + (10 if ratio == 1 else 0)
    elif payload.kind == "scenario":
        raw = max(2, payload.score * 5)
    elif payload.kind == "pair":
        raw = 5 if payload.score else 1
    elif payload.kind == "course_day":
        raw = 25 + (10 if ratio == 1 else 0)
    elif payload.kind == "tone_lab":
        raw = 8 + payload.score * 2 + (7 if ratio == 1 else 0)
    else:
        raw = (100 if ratio >= .7 else 25) + (50 if ratio >= .9 else 0)
    award = award_xp(
        db, user.id, payload.kind, payload.topic or payload.session_id, raw,
        language=language, topic=payload.topic,
        idempotency_key=f"practice:{user.id}:{payload.session_id}",
        metadata={"score": payload.score, "total": payload.total}, anti_farm=True,
    )
    db.commit()
    return {"ok": True, "duplicate": False, **award}


def game_pool(language: str, topic: str, db: Session) -> list[dict[str, Any]]:
    pool = terms_for(language, db)
    if topic:
        filtered = [row for row in pool if row["topic"] == topic]
        if len(filtered) >= 6:
            pool = filtered
    return pool


def chunk_chinese(text_value: str) -> list[str]:
    text_value = text_value.strip()
    if " " in text_value:
        return [x for x in text_value.split() if x]
    return [text_value[i:i+2] for i in range(0, len(text_value), 2)]


@app.post("/api/games/{game_type}/start")
def start_game(
    game_type: str, language: str = Query(default="chinese"), topic: str = Query(default=""),
    user: User = Depends(current_user), db: Session = Depends(db_session),
):
    language = validate_language(language)
    require_feature(db, user.id, "games")
    if game_type not in {"match", "listening", "mistake", "phrase"}:
        raise HTTPException(404, "Неизвестная игра")
    _consume_daily_quota(db, user.id, "game_starts", PILOT_DAILY_GAME_START_CAP)
    pool = game_pool(language, topic, db)
    if len(pool) < 6:
        raise HTTPException(400, "Недостаточно терминов для игры")
    public_id = secrets.token_urlsafe(18)
    rng = random.Random(secrets.randbits(64))
    selected = rng.sample(pool, min(8, len(pool)))
    items: list[dict[str, Any]] = []
    answers: list[Any] = []
    if game_type == "match":
        items = [{"id": row["id"], "term": row["term"], "pronunciation": row.get("pronunciation", ""), "reading": row.get("reading", ""), "translation": row["translation"]} for row in selected]
        answers = [row["id"] for row in selected]
    elif game_type == "listening":
        for row in selected:
            distractors = rng.sample([x for x in pool if x["id"] != row["id"]], 3)
            options = [row["translation"], *[x["translation"] for x in distractors]]
            rng.shuffle(options)
            items.append({"id": row["id"], "audio_text": row["term"], "pronunciation": row.get("pronunciation", ""), "reading": row.get("reading", ""), "options": options})
            answers.append(options.index(row["translation"]))
    elif game_type == "mistake":
        for row in selected:
            distractors = rng.sample([x for x in pool if x["id"] != row["id"]], 3)
            options = [row["translation"], *[x["translation"] for x in distractors]]
            rng.shuffle(options)
            items.append({"id": row["id"], "term": row["term"], "pronunciation": row.get("pronunciation", ""), "reading": row.get("reading", ""), "options": options})
            answers.append(options.index(row["translation"]))
    else:
        phrase_rows = [row for row in pool if len((row.get("example") or "").strip()) >= 8]
        selected = rng.sample(phrase_rows or pool, min(6, len(phrase_rows or pool)))
        for row in selected:
            phrase = row.get("example") or row["term"]
            tokens = phrase.split() if language == "english" else chunk_chinese(phrase)
            shuffled = list(tokens)
            rng.shuffle(shuffled)
            items.append({"id": row["id"], "translation": row.get("example_translation") or row["translation"], "tokens": shuffled})
            answers.append(tokens)
    db.add(GameSession(
        public_id=public_id, user_id=user.id, game_type=game_type, language=language, topic=topic,
        payload_json=json.dumps({"items": items, "answers": answers}, ensure_ascii=False),
        total=len(items),
    ))
    db.commit()
    return {"session_id": public_id, "game_type": game_type, "language": language, "topic": topic, "items": items, "total": len(items)}


@app.post("/api/games/{session_id}/finish")
def finish_game(session_id: str, payload: GameFinishPayload, request: Request, user: User = Depends(current_user), db: Session = Depends(db_session)):
    rate_limit(request, "game_finish", 120, 60)
    session = db.scalar(select(GameSession).where(GameSession.public_id == session_id, GameSession.user_id == user.id))
    if not session:
        raise HTTPException(404, "Игровая сессия не найдена")
    if session.status == "completed":
        return {"ok": True, "duplicate": True, "score": session.score, "total": session.total, "profile": gamification_view(db, user.id)}
    data = json.loads(session.payload_json)
    correct_answers = data.get("answers", [])
    score = 0
    for index, correct in enumerate(correct_answers):
        if index >= len(payload.answers):
            break
        answer = payload.answers[index]
        if isinstance(correct, list):
            if list(answer) == correct:
                score += 1
        elif answer == correct:
            score += 1
    session.status = "completed"
    session.score = score
    session.completed_at = datetime.now(timezone.utc)
    raw = 10 + score * 3 + (10 if score == session.total else 0)
    award = award_xp(
        db, user.id, f"game:{session.game_type}", session.topic or session.game_type, raw,
        language=session.language, topic=session.topic,
        idempotency_key=f"game:{user.id}:{session.public_id}", metadata={"score": score, "total": session.total}, anti_farm=True,
    )
    db.commit()
    return {"ok": True, "score": score, "total": session.total, **award}


def notification_pref(db: Session, user_id: int) -> NotificationPreference:
    row = db.get(NotificationPreference, user_id)
    if not row:
        row = NotificationPreference(user_id=user_id)
        db.add(row)
        db.flush()
    return row


def consecutive_nudges_without_activity(db: Session, user_id: int, last_activity: datetime | None) -> int:
    query = select(LearningNudge).where(LearningNudge.user_id == user_id)
    if last_activity:
        query = query.where(LearningNudge.created_at > last_activity)
    return len(db.scalars(query).all())


def maybe_queue_nudge(db: Session, user: User) -> LearningNudge | None:
    if not user_feature_flags(db, user.id).get("learning_nudges", True):
        return None
    pref = notification_pref(db, user.id)
    if pref.mode == "off":
        return None
    profile = get_profile(db, user.id)
    last_activity = profile.last_activity_at or user.created_at
    if last_activity and last_activity.tzinfo is None:
        last_activity = last_activity.replace(tzinfo=timezone.utc)
    count = consecutive_nudges_without_activity(db, user.id, last_activity)
    if count >= 3:
        return None
    intervals = {"minimal": [7, 14, 30], "normal": [3, 7, 14], "active": [1, 3, 7]}
    required_days = intervals[pref.mode][count]
    now = datetime.now(timezone.utc)
    last_nudge = db.scalar(select(LearningNudge).where(LearningNudge.user_id == user.id).order_by(LearningNudge.created_at.desc()).limit(1))
    anchor = last_nudge.created_at if last_nudge and count else last_activity
    if anchor and anchor.tzinfo is None:
        anchor = anchor.replace(tzinfo=timezone.utc)
    if anchor and now - anchor < timedelta(days=required_days):
        return None
    progress = db.scalars(select(TermProgress).where(TermProgress.user_id == user.id)).all()
    due = []
    for row in progress:
        updated = row.updated_at.replace(tzinfo=timezone.utc) if row.updated_at.tzinfo is None else row.updated_at
        threshold = timedelta(days=2 if row.status == "learning" else 7)
        if now - updated >= threshold:
            due.append(row)
    if due:
        title = "Короткое повторение готово"
        body = f"{min(len(due), 12)} терминов пора повторить. Обычно это занимает 3–5 минут."
        target = "games" if user_feature_flags(db, user.id).get("games", True) else "topics"
        reason = "review_due"
    else:
        title = "Можно продолжить с того же места"
        body = "Есть несколько минут? Откройте одну короткую тренировку — без обязательной серии дней."
        target = "home"
        reason = "gentle_return"
    row = LearningNudge(user_id=user.id, title=title, body=body, target_view=target, reason=reason)
    db.add(row)
    db.flush()
    return row


@app.get("/api/notifications/settings")
def get_notification_settings(user: User = Depends(current_user), db: Session = Depends(db_session)):
    row = notification_pref(db, user.id)
    db.commit()
    return {"mode": row.mode, "window_start": row.window_start, "window_end": row.window_end, "browser_enabled": row.browser_enabled}


@app.put("/api/notifications/settings")
def set_notification_settings(payload: NotificationSettingsPayload, user: User = Depends(current_user), db: Session = Depends(db_session)):
    row = notification_pref(db, user.id)
    row.mode = payload.mode
    row.window_start = payload.window_start
    row.window_end = payload.window_end
    row.browser_enabled = payload.browser_enabled
    row.updated_at = datetime.now(timezone.utc)
    db.commit()
    return {"ok": True, "mode": row.mode, "window_start": row.window_start, "window_end": row.window_end, "browser_enabled": row.browser_enabled}


@app.get("/api/notifications/pending")
def pending_notifications(user: User = Depends(current_user), db: Session = Depends(db_session)):
    enabled = user_feature_flags(db, user.id).get("learning_nudges", True)
    if enabled:
        maybe_queue_nudge(db, user)
        rows = db.scalars(select(LearningNudge).where(
            LearningNudge.user_id == user.id, LearningNudge.read == False  # noqa: E712
        ).order_by(LearningNudge.created_at.desc()).limit(5)).all()
    else:
        rows = []
    pref = notification_pref(db, user.id)
    db.commit()
    return {"feature_enabled":enabled,"settings": {"mode": pref.mode, "window_start": pref.window_start, "window_end": pref.window_end, "browser_enabled": pref.browser_enabled}, "items": [
        {"id": row.id, "title": row.title, "body": row.body, "target_view": row.target_view, "reason": row.reason, "created_at": row.created_at.isoformat()} for row in rows
    ]}


@app.post("/api/notifications/{nudge_id}/read")
def read_notification(nudge_id: int, user: User = Depends(current_user), db: Session = Depends(db_session)):
    row = db.scalar(select(LearningNudge).where(LearningNudge.id == nudge_id, LearningNudge.user_id == user.id))
    if not row:
        raise HTTPException(404, "Уведомление не найдено")
    row.read = True
    db.commit()
    return {"ok": True}


def _term_snapshot(row: CustomTerm) -> dict[str, Any]:
    return {key: getattr(row, key) for key in ("language","shop","topic","subtopic","level","term","pronunciation","reading","translation","example","example_translation","tags","source_type","source_ref","status")}


def record_term_revision(db: Session, row: CustomTerm, actor_user_id: int | None, action: str) -> TermRevision:
    last_no = db.scalar(select(func.max(TermRevision.revision_no)).where(TermRevision.term_id == row.id)) or 0
    revision = TermRevision(term_id=row.id, revision_no=int(last_no)+1, action=action, snapshot_json=json.dumps(_term_snapshot(row), ensure_ascii=False, sort_keys=True), actor_user_id=actor_user_id)
    db.add(revision)
    db.flush()
    return revision


def get_or_create_srs_card(db: Session, user_id: int, language: str, term_id: str, topic: str = "") -> SRSCard:
    card = db.scalar(select(SRSCard).where(SRSCard.user_id==user_id, SRSCard.language==language, SRSCard.term_id==term_id))
    if not card:
        card = SRSCard(user_id=user_id, language=language, term_id=term_id, topic=topic, due_at=datetime.now(timezone.utc))
        db.add(card); db.flush()
    elif topic and not card.topic:
        card.topic = topic
    return card


def schedule_srs(card: SRSCard, quality: int, now: datetime | None = None) -> SRSCard:
    now = now or datetime.now(timezone.utc)
    q = max(0, min(5, int(quality)))
    if q < 3:
        card.repetitions = 0
        card.lapses += 1
        card.interval_days = 1
    else:
        card.repetitions += 1
        if card.repetitions == 1:
            card.interval_days = 1
        elif card.repetitions == 2:
            card.interval_days = 3
        else:
            card.interval_days = max(1, round(card.interval_days * card.ease_factor_pct / 100))
    ef = card.ease_factor_pct / 100.0
    ef = ef + (0.1 - (5-q) * (0.08 + (5-q) * 0.02))
    card.ease_factor_pct = int(max(130, min(300, round(ef*100))))
    card.last_quality = q
    card.due_at = now + timedelta(days=card.interval_days)
    card.updated_at = now
    return card


def admin_term_response(row: CustomTerm) -> dict[str, Any]:
    return {**custom_term_view(row), "db_id": row.id, "created_at": row.created_at.isoformat(), "updated_at": row.updated_at.isoformat()}


@app.get("/api/admin/terms")
def admin_terms(
    language: str | None = Query(default=None), status: str | None = Query(default=None),
    user: User = Depends(require_roles("admin", "editor")), db: Session = Depends(db_session),
):
    query = select(CustomTerm).order_by(CustomTerm.updated_at.desc())
    if language in LANGUAGES:
        query = query.where(CustomTerm.language == language)
    if status:
        query = query.where(CustomTerm.status == status)
    rows = db.scalars(query).all()
    return [admin_term_response(row) for row in rows]


@app.post("/api/admin/terms")
def admin_create_term(payload: CustomTermPayload, request: Request, user: User = Depends(require_roles("admin", "editor")), db: Session = Depends(db_session)):
    language = validate_language(payload.language)
    values = payload.model_dump(exclude={"language"})
    if TERM_APPROVAL_REQUIRED and user.role != "admin" and values.get("status") == "published":
        values["status"] = "review"
    row = CustomTerm(public_id=f"custom-{secrets.token_hex(8)}", language=language, created_by=user.id, **values)
    db.add(row)
    db.flush()
    revision = record_term_revision(db, row, user.id, "create")
    audit_event(db, "term.create", actor_user_id=user.id, target_type="custom_term", target_id=row.public_id, request=request, metadata={"language": language, "topic": row.topic, "status": row.status, "revision": revision.revision_no})
    db.commit()
    db.refresh(row)
    return admin_term_response(row)


@app.patch("/api/admin/terms/{term_id}")
def admin_update_term(term_id: int, payload: CustomTermPayload, request: Request, user: User = Depends(require_roles("admin", "editor")), db: Session = Depends(db_session)):
    row = db.get(CustomTerm, term_id)
    if not row:
        raise HTTPException(404, "Термин не найден")
    record_term_revision(db, row, user.id, "before_update")
    values = payload.model_dump()
    values["language"] = validate_language(values["language"])
    if TERM_APPROVAL_REQUIRED and user.role != "admin" and values.get("status") == "published":
        values["status"] = "review"
    for key, value in values.items():
        setattr(row, key, value)
    row.updated_at = datetime.now(timezone.utc)
    revision = record_term_revision(db, row, user.id, "update")
    audit_event(db, "term.update", actor_user_id=user.id, target_type="custom_term", target_id=row.public_id, request=request, metadata={"topic": row.topic, "status": row.status, "revision": revision.revision_no})
    db.commit()
    return admin_term_response(row)


@app.delete("/api/admin/terms/{term_id}")
def admin_delete_term(term_id: int, request: Request, user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    row = db.get(CustomTerm, term_id)
    if not row:
        raise HTTPException(404, "Термин не найден")
    public_id = row.public_id
    audit_event(db, "term.delete", actor_user_id=user.id, target_type="custom_term", target_id=public_id, request=request, metadata={"topic": row.topic})
    db.delete(row)
    db.commit()
    return {"ok": True}


@app.get("/api/admin/terms/{term_id}/revisions")
def admin_term_revisions(term_id: int, user: User = Depends(require_roles("admin", "editor")), db: Session = Depends(db_session)):
    row = db.get(CustomTerm, term_id)
    if not row: raise HTTPException(404, "Термин не найден")
    revisions = db.scalars(select(TermRevision).where(TermRevision.term_id==term_id).order_by(TermRevision.revision_no.desc())).all()
    reviews = db.scalars(select(TermReview).where(TermReview.term_id==term_id).order_by(TermReview.created_at.desc())).all()
    return {"term":admin_term_response(row), "revisions":[{"id":r.id,"revision_no":r.revision_no,"action":r.action,"snapshot":json.loads(r.snapshot_json),"actor_user_id":r.actor_user_id,"created_at":r.created_at.isoformat()} for r in revisions], "reviews":[{"id":x.id,"revision_no":x.revision_no,"decision":x.decision,"note":x.note,"reviewer_user_id":x.reviewer_user_id,"created_at":x.created_at.isoformat()} for x in reviews]}


@app.post("/api/admin/terms/{term_id}/submit-review")
def admin_term_submit_review(term_id: int, request: Request, user: User = Depends(require_roles("admin","editor")), db: Session = Depends(db_session)):
    row=db.get(CustomTerm,term_id)
    if not row: raise HTTPException(404,"Термин не найден")
    if row.status == "published": raise HTTPException(409,"Опубликованный термин сначала измените")
    row.status="review"; row.updated_at=datetime.now(timezone.utc)
    rev=record_term_revision(db,row,user.id,"submit_review")
    audit_event(db,"term.submit_review",actor_user_id=user.id,target_type="custom_term",target_id=row.public_id,request=request,metadata={"revision":rev.revision_no})
    db.commit(); return admin_term_response(row)


@app.post("/api/admin/terms/{term_id}/approve")
def admin_term_approve(term_id: int, payload: TermReviewPayload, request: Request, user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    row=db.get(CustomTerm,term_id)
    if not row: raise HTTPException(404,"Термин не найден")
    if row.status != "review": raise HTTPException(409,"Термин должен быть на проверке")
    row.status="published"; row.updated_at=datetime.now(timezone.utc)
    rev=record_term_revision(db,row,user.id,"approve")
    db.add(TermReview(term_id=row.id,revision_no=rev.revision_no,reviewer_user_id=user.id,decision="approved",note=payload.note))
    audit_event(db,"term.approve",actor_user_id=user.id,target_type="custom_term",target_id=row.public_id,request=request,metadata={"revision":rev.revision_no,"note":payload.note})
    db.commit(); return admin_term_response(row)


@app.post("/api/admin/terms/{term_id}/reject")
def admin_term_reject(term_id: int, payload: TermReviewPayload, request: Request, user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    row=db.get(CustomTerm,term_id)
    if not row: raise HTTPException(404,"Термин не найден")
    row.status="draft"; row.updated_at=datetime.now(timezone.utc)
    rev=record_term_revision(db,row,user.id,"reject")
    db.add(TermReview(term_id=row.id,revision_no=rev.revision_no,reviewer_user_id=user.id,decision="rejected",note=payload.note))
    audit_event(db,"term.reject",actor_user_id=user.id,target_type="custom_term",target_id=row.public_id,request=request,metadata={"revision":rev.revision_no,"note":payload.note})
    db.commit(); return admin_term_response(row)


@app.post("/api/admin/terms/{term_id}/rollback/{revision_no}")
def admin_term_rollback(term_id: int, revision_no: int, request: Request, user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    row=db.get(CustomTerm,term_id); rev=db.scalar(select(TermRevision).where(TermRevision.term_id==term_id,TermRevision.revision_no==revision_no))
    if not row or not rev: raise HTTPException(404,"Термин или версия не найдены")
    record_term_revision(db,row,user.id,"before_rollback")
    snap=json.loads(rev.snapshot_json)
    for key,value in snap.items(): setattr(row,key,value)
    row.status="review" if TERM_APPROVAL_REQUIRED else row.status; row.updated_at=datetime.now(timezone.utc)
    newrev=record_term_revision(db,row,user.id,"rollback")
    audit_event(db,"term.rollback",actor_user_id=user.id,target_type="custom_term",target_id=row.public_id,request=request,metadata={"from_revision":revision_no,"new_revision":newrev.revision_no})
    db.commit(); return admin_term_response(row)


@app.post("/api/admin/terms/import")
def admin_import_terms(
    request: Request, file: UploadFile = File(...), default_language: str = Query(default="chinese"),
    user: User = Depends(require_roles("admin", "editor")), db: Session = Depends(db_session),
):
    from io import BytesIO, StringIO
    import csv
    from openpyxl import load_workbook
    rate_limit(request, "term_import", 10, 60)
    default_language = validate_language(default_language)
    raw = file.file.read(MAX_IMPORT_BYTES + 1)
    if len(raw) > MAX_IMPORT_BYTES:
        raise HTTPException(413, f"Файл превышает лимит {MAX_IMPORT_BYTES} байт")
    rows: list[dict[str, Any]] = []
    filename = (file.filename or "").lower()
    content_type = (file.content_type or "").lower()
    allowed_mime = {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "text/csv", "application/csv", "text/plain", "application/octet-stream"}
    if content_type and content_type not in allowed_mime:
        raise HTTPException(415, "Недопустимый MIME-тип файла")
    if filename.endswith(".xlsx"):
        import zipfile
        if not raw.startswith(b"PK\x03\x04"):
            raise HTTPException(400, "XLSX signature mismatch")
        try:
            with zipfile.ZipFile(BytesIO(raw)) as zf:
                uncompressed = sum(info.file_size for info in zf.infolist())
                if uncompressed > MAX_XLSX_UNCOMPRESSED_BYTES:
                    raise HTTPException(413, "XLSX слишком велик после распаковки")
                if any(info.file_size > 0 and info.compress_size == 0 for info in zf.infolist()):
                    raise HTTPException(400, "Некорректный XLSX archive")
        except zipfile.BadZipFile:
            raise HTTPException(400, "Некорректный XLSX archive")
        wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        values = []
        for row in ws.iter_rows(values_only=True):
            values.append(row)
            if len(values) > MAX_IMPORT_ROWS + 1:
                raise HTTPException(413, f"Импорт ограничен {MAX_IMPORT_ROWS} строками")
        if not values:
            raise HTTPException(400, "Пустой XLSX")
        headers = [str(x or "").strip().lower() for x in values[0]][:50]
        rows = [dict(zip(headers, row[:50])) for row in values[1:]]
    elif filename.endswith(".csv"):
        try:
            text_value = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise HTTPException(400, "CSV должен быть UTF-8")
        reader = csv.DictReader(StringIO(text_value))
        rows = []
        for item in reader:
            rows.append(item)
            if len(rows) > MAX_IMPORT_ROWS:
                raise HTTPException(413, f"Импорт ограничен {MAX_IMPORT_ROWS} строками")
    else:
        raise HTTPException(400, "Поддерживаются .xlsx и .csv")
    created = 0
    errors: list[str] = []
    for index, item in enumerate(rows, start=2):
        try:
            language = str(item.get("language") or default_language).strip().lower()
            language = validate_language(language)
            term = str(item.get("term") or "").strip()
            translation = str(item.get("translation") or "").strip()
            topic = str(item.get("topic") or item.get("shop") or "Корпоративная терминология").strip()
            if not term or not translation:
                raise ValueError("term/translation обязательны")
            level = str(item.get("level") or "A1").strip().upper()
            if level not in LEVELS:
                level = "A1"
            row = CustomTerm(
                public_id=f"custom-{secrets.token_hex(8)}", language=language,
                shop=str(item.get("shop") or "").strip(), topic=topic,
                subtopic=str(item.get("subtopic") or "").strip(), level=level, term=term,
                pronunciation=str(item.get("pronunciation") or item.get("pinyin") or "").strip(),
                reading=str(item.get("reading") or item.get("ru_read") or "").strip(),
                translation=translation, example=str(item.get("example") or "").strip(),
                example_translation=str(item.get("example_translation") or "").strip(),
                tags=str(item.get("tags") or "").strip(), source_type=str(item.get("source_type") or "import").strip().lower(),
                source_ref=str(item.get("source_ref") or file.filename or "").strip(), status=str(item.get("status") or "review").strip().lower(),
                created_by=user.id,
            )
            if row.status not in {"draft", "review", "published", "archived"}:
                row.status = "review"
            if row.source_type not in {"manual","company_standard","supplier","work_instruction","engineering_document","language_expert","public_dictionary","ai_suggestion","import"}:
                row.source_type = "import"
            if TERM_APPROVAL_REQUIRED and user.role != "admin" and row.status == "published":
                row.status = "review"
            db.add(row)
            db.flush()
            record_term_revision(db, row, user.id, "import")
            created += 1
        except Exception as exc:
            errors.append(f"Строка {index}: {exc}")
    audit_event(db, "term.import", actor_user_id=user.id, target_type="custom_term_batch", target_id=file.filename or "upload", request=request, metadata={"created": created, "error_count": len(errors), "default_language": default_language})
    db.commit()
    return {"ok": True, "created": created, "errors": errors[:50], "error_count": len(errors)}


@app.get("/api/admin/taxonomy")
def admin_taxonomy(user: User = Depends(require_roles("admin", "editor", "manager")), db: Session = Depends(db_session)):
    custom = db.scalars(select(CustomTerm)).all()
    topics = sorted({row["label"] for row in TOPIC_ROWS} | {row.topic for row in custom if row.topic})
    shops = sorted({row.shop for row in custom if row.shop} | {"Сборка", "Сварка", "Окраска", "Штамповка", "Компоненты", "Логистика", "Качество", "R&D / Engineering", "Maintenance", "Tool Shop", "Закупки", "Локализация / ВЭД", "Project Management", "HSE"})
    return {"topics": topics, "shops": shops, "levels": LEVELS, "languages": sorted(LANGUAGES)}


def user_admin_stats(db: Session, row: User) -> dict[str, Any]:
    profile = gamification_view(db, row.id)
    progress = db.scalars(select(TermProgress).where(TermProgress.user_id == row.id)).all()
    known = sum(1 for p in progress if p.status == "known")
    exams = db.scalars(select(ExamResult).where(ExamResult.user_id == row.id)).all()
    last = profile.get("last_activity_at")
    pref = notification_pref(db, row.id)
    return {
        "id": row.id, "username": row.username, "display_name": row.display_name, "role": row.role, "department": row.department,
        "preferred_language": row.preferred_language, "level": profile["level"], "level_title": profile["title"],
        "lifetime_xp": profile["lifetime_xp"], "spendable_xp": profile["spendable_xp"], "weekly_xp": profile["weekly_xp"],
        "terms_known": known, "terms_touched": len(progress), "best_exam": max((e.score for e in exams), default=0),
        "last_activity_at": last, "notification_mode": pref.mode, "created_at": row.created_at.isoformat(),
    }


@app.get("/api/admin/users")
def admin_users(user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    rows = db.scalars(select(User).order_by(User.display_name)).all()
    result = [user_admin_stats(db, row) for row in rows]
    db.commit()
    return result


@app.get("/api/admin/users/{user_id}/learning-stats")
def admin_user_learning_stats(user_id: int, user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    target = db.get(User, user_id)
    if not target:
        raise HTTPException(404, "Пользователь не найден")
    progress = db.scalars(select(TermProgress).where(TermProgress.user_id == target.id)).all()
    by_language: dict[str, Any] = {}
    for language in sorted(LANGUAGES):
        lang_rows = [row for row in progress if row.language == language]
        effective = terms_for(language, db)
        topic_by_id = {row["id"]: row["topic"] for row in effective}
        topic_counts: dict[str, dict[str, int]] = {}
        for row in lang_rows:
            topic = topic_by_id.get(row.term_id, "Другое")
            item = topic_counts.setdefault(topic, {"touched": 0, "known": 0})
            item["touched"] += 1
            if row.status == "known": item["known"] += 1
        by_language[language] = {
            "touched": len(lang_rows), "known": sum(1 for r in lang_rows if r.status == "known"),
            "topics": [{"topic": topic, **counts, "mastery_percent": round(counts["known"] * 100 / max(1, counts["touched"]))} for topic, counts in sorted(topic_counts.items())],
        }
    practices = db.scalars(select(PracticeResult).where(PracticeResult.user_id == target.id).order_by(PracticeResult.created_at.desc()).limit(30)).all()
    result = user_admin_stats(db, target)
    result["languages"] = by_language
    result["recent_practice"] = [{"kind": x.kind, "language": x.language, "topic": x.topic, "score": x.score, "total": x.total, "created_at": x.created_at.isoformat()} for x in practices]
    db.commit()
    return result


@app.patch("/api/admin/users/{user_id}/role")
def admin_set_role(user_id: int, payload: UserRolePayload, request: Request, user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    target = db.get(User, user_id)
    if not target:
        raise HTTPException(404, "Пользователь не найден")
    if target.id == user.id and payload.role != "admin":
        raise HTTPException(400, "Нельзя снять роль admin у самого себя из этой панели")
    previous_role = target.role
    target.role = payload.role
    audit_event(db, "user.role.change", actor_user_id=user.id, target_type="user", target_id=str(target.id), request=request, metadata={"from": previous_role, "to": payload.role})
    db.commit()
    return {"ok": True, "user": user_view(target)}


@app.patch("/api/admin/users/{user_id}/department")
def admin_set_department(user_id: int, payload: UserDepartmentPayload, request: Request, user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    target = db.get(User, user_id)
    if not target:
        raise HTTPException(404, "Пользователь не найден")
    previous = target.department
    target.department = payload.department.strip()[:160]
    audit_event(db, "user.department.change", actor_user_id=user.id, target_type="user", target_id=str(target.id), request=request, metadata={"from":previous,"to":target.department})
    db.commit()
    return {"ok":True,"user":user_view(target)}


def _manager_target_allowed(manager: User, target: User) -> bool:
    return manager.role == "admin" or (manager.role == "manager" and manager.department == target.department)


@app.get("/api/manager/team")
def manager_team(user: User = Depends(require_roles("manager", "admin")), db: Session = Depends(db_session)):
    query = select(User).order_by(User.display_name)
    if user.role == "manager":
        query = query.where(User.department == user.department)
    rows = db.scalars(query).all()
    result = [user_admin_stats(db, row) for row in rows]
    db.commit()
    return {"department": user.department if user.role == "manager" else "ALL", "users": result}


@app.get("/api/manager/team/{user_id}/learning-stats")
def manager_user_learning_stats(user_id: int, user: User = Depends(require_roles("manager", "admin")), db: Session = Depends(db_session)):
    target = db.get(User, user_id)
    if not target:
        raise HTTPException(404, "Пользователь не найден")
    if not _manager_target_allowed(user, target):
        raise HTTPException(403, "Доступ только к сотрудникам своего подразделения")
    progress = db.scalars(select(TermProgress).where(TermProgress.user_id == target.id)).all()
    by_language: dict[str, Any] = {}
    for language in sorted(LANGUAGES):
        lang_rows = [row for row in progress if row.language == language]
        effective = terms_for(language, db)
        topic_by_id = {row["id"]: row["topic"] for row in effective}
        topic_counts: dict[str, dict[str, int]] = {}
        for row in lang_rows:
            topic = topic_by_id.get(row.term_id, "Другое")
            item = topic_counts.setdefault(topic, {"touched":0,"known":0})
            item["touched"] += 1
            if row.status == "known":
                item["known"] += 1
        by_language[language] = {
            "touched":len(lang_rows), "known":sum(1 for r in lang_rows if r.status == "known"),
            "topics":[{"topic":topic, **counts, "mastery_percent":round(counts["known"]*100/max(1,counts["touched"]))} for topic,counts in sorted(topic_counts.items())],
        }
    result = user_admin_stats(db, target)
    result["languages"] = by_language
    return result


def _latest_matching_file(directory: Path, patterns: tuple[str, ...]) -> Path | None:
    try:
        if not directory.exists() or not directory.is_dir():
            return None
        candidates: list[Path] = []
        for pattern in patterns:
            candidates.extend([p for p in directory.glob(pattern) if p.is_file()])
        return max(candidates, key=lambda p: p.stat().st_mtime) if candidates else None
    except OSError:
        return None

def _age_minutes(path: Path | None) -> float | None:
    if path is None:
        return None
    try:
        return round(max(0.0, (time.time() - path.stat().st_mtime) / 60.0), 2)
    except OSError:
        return None

def recovery_evidence_snapshot() -> dict[str, Any]:
    backup = _latest_matching_file(BACKUP_EVIDENCE_DIR, ("mgc_languages_*.dump", "mgc_languages_base_*.tar.gz"))
    restore = _latest_matching_file(BACKUP_EVIDENCE_DIR, ("*.restore-ok.json", "restore_evidence_*.json"))
    backup_age = _age_minutes(backup)
    restore_age_minutes = _age_minutes(restore)
    restore_payload: dict[str, Any] = {}
    if restore is not None:
        try:
            if restore.stat().st_size <= 64 * 1024:
                restore_payload = json.loads(restore.read_text(encoding="utf-8"))
        except Exception:
            restore_payload = {"parse_error": True}
    wal = _latest_matching_file(WAL_ARCHIVE_DIR, ("*",))
    wal_age = _age_minutes(wal)
    backup_status = "unavailable" if not BACKUP_EVIDENCE_DIR.exists() else ("missing" if backup_age is None else ("stale" if backup_age > BACKUP_MAX_AGE_MINUTES else "ok"))
    restore_age_days = round(restore_age_minutes / 1440.0, 2) if restore_age_minutes is not None else None
    restore_status = "unavailable" if not BACKUP_EVIDENCE_DIR.exists() else ("missing" if restore_age_days is None else ("stale" if restore_age_days > RESTORE_EVIDENCE_MAX_AGE_DAYS else "ok"))
    rpo_estimate = backup_age if backup_age is not None else None
    rpo_status = "unknown" if rpo_estimate is None else ("met" if rpo_estimate <= RPO_TARGET_MINUTES else "missed")
    restore_duration_seconds = restore_payload.get("duration_seconds") if isinstance(restore_payload, dict) else None
    try:
        restore_duration_minutes = round(float(restore_duration_seconds) / 60.0, 2) if restore_duration_seconds is not None else None
    except (TypeError, ValueError):
        restore_duration_minutes = None
    rto_status = "unknown" if restore_duration_minutes is None else ("met" if restore_duration_minutes <= RTO_TARGET_MINUTES else "missed")
    return {
        "backup_dir": str(BACKUP_EVIDENCE_DIR),
        "backup": {"status": backup_status, "latest_file": backup.name if backup else None, "age_minutes": backup_age, "max_age_minutes": BACKUP_MAX_AGE_MINUTES},
        "restore_rehearsal": {"status": restore_status, "latest_file": restore.name if restore else None, "age_days": restore_age_days, "max_age_days": RESTORE_EVIDENCE_MAX_AGE_DAYS, "evidence": restore_payload},
        "wal_archive": {"directory": str(WAL_ARCHIVE_DIR), "latest_file": wal.name if wal else None, "age_minutes": wal_age, "note": "Pilot WAL archive on the same host is rehearsal evidence only; production PITR requires encrypted off-host WAL/base backups."},
        "objectives": {"rpo_target_minutes": RPO_TARGET_MINUTES, "rto_target_minutes": RTO_TARGET_MINUTES, "rpo_estimate_minutes_from_latest_backup": rpo_estimate, "rpo_status": rpo_status, "last_restore_duration_minutes": restore_duration_minutes, "rto_status": rto_status, "note": "RTO is proven by timed restore/failover drills, not inferred from application uptime."},
        "required_for_alerting": RECOVERY_EVIDENCE_REQUIRED,
    }


def evaluate_pilot_alerts(db: Session) -> dict[str, Any]:
    now = datetime.now(timezone.utc)
    ready_ok, checks = readiness_checks(db)
    tts_health = _tts_health()
    recovery = update_recovery_state(ready_ok=ready_ok, checks=checks, tts_health=tts_health)
    slo = slo_snapshot()
    maintenance = cleanup_runtime_data(db, dry_run=True)
    backlog = sum(int(v or 0) for v in maintenance.get("counts", {}).values())
    conditions: dict[str, dict[str, Any]] = {}

    def add(key: str, severity: str, component: str, title: str, detail: str, metadata: dict[str, Any] | None = None) -> None:
        conditions[key] = {"severity":severity,"component":component,"title":title,"detail":detail,"metadata":metadata or {}}

    if not ready_ok:
        add("readiness.not_ready", "critical", "platform", "Pilot instance is not ready", "Readiness policy failed; inspect schema, database and security configuration.", {"checks":checks})
    db_latency = checks.get("database_latency_ms")
    if isinstance(db_latency, (int, float)) and db_latency > ALERT_DB_LATENCY_MS:
        add("database.latency", "warning", "database", "Database latency is elevated", f"Readiness query latency is {db_latency} ms; alert threshold is {ALERT_DB_LATENCY_MS} ms.", {"latency_ms":db_latency,"threshold_ms":ALERT_DB_LATENCY_MS})
    db_queries = db_query_telemetry_snapshot()
    db_pool = db_pool_snapshot()
    if db_queries["slow_queries"] >= DB_SLOW_QUERY_ALERT_COUNT:
        add("database.slow_queries", "warning", "database", "Slow-query volume is elevated", f"{db_queries['slow_queries']} queries exceeded {DB_SLOW_QUERY_THRESHOLD_MS} ms in the last {DB_QUERY_WINDOW_MINUTES} minutes.", {"telemetry":db_queries})
    if db_pool.get("available") and float(db_pool.get("saturation_percent") or 0) >= DB_POOL_ALERT_PERCENT:
        add("database.pool_saturation", "warning", "database", "Database connection pool saturation is high", f"{db_pool['saturation_percent']}% of configured DB connection capacity is checked out.", {"pool":db_pool})
    if slo["samples"] >= 20 and slo["error_rate_percent"] > ALERT_ERROR_RATE_PERCENT:
        add("slo.error_rate", "critical", "http", "HTTP error rate exceeds pilot threshold", f"5xx rate is {slo['error_rate_percent']}% in the current SLO window.", {"slo":slo})
    if slo["samples"] >= 20 and slo["p95_ms"] > ALERT_P95_MS:
        add("slo.p95_latency", "warning", "http", "P95 latency exceeds pilot threshold", f"P95 latency is {slo['p95_ms']} ms in the current SLO window.", {"slo":slo})
    if _tts_circuit_open():
        add("tts.circuit_open", "warning", "tts", "Offline TTS circuit breaker is open", "Learning remains available through browser speech fallback; investigate local TTS before closing the alert.", {"tts":tts_health})
    if backlog > ALERT_CLEANUP_BACKLOG:
        add("maintenance.backlog", "warning", "maintenance", "Maintenance cleanup backlog is high", f"{backlog} rows are waiting for retention cleanup.", {"backlog":backlog,"threshold":ALERT_CLEANUP_BACKLOG})
    recovery_evidence = recovery_evidence_snapshot()
    if RECOVERY_EVIDENCE_REQUIRED:
        backup_status = recovery_evidence["backup"]["status"]
        restore_status = recovery_evidence["restore_rehearsal"]["status"]
        if backup_status in {"missing", "stale", "unavailable"}:
            add("recovery.backup_evidence", "critical" if backup_status == "stale" else "warning", "recovery", "Backup evidence is not current", f"Backup evidence status is {backup_status}; target maximum age is {BACKUP_MAX_AGE_MINUTES} minutes.", recovery_evidence["backup"])
        if restore_status in {"missing", "stale", "unavailable"}:
            add("recovery.restore_rehearsal", "warning", "recovery", "Restore rehearsal evidence is not current", f"Restore rehearsal evidence status is {restore_status}; target maximum age is {RESTORE_EVIDENCE_MAX_AGE_DAYS} days.", recovery_evidence["restore_rehearsal"])

    existing = {row.alert_key:row for row in db.scalars(select(PilotAlert)).all()}
    for key, item in conditions.items():
        row = existing.get(key)
        if row is None:
            row = PilotAlert(alert_key=key, severity=item["severity"], status="open", component=item["component"], title=item["title"], detail=item["detail"], metadata_json=json.dumps(item["metadata"], ensure_ascii=False, default=str), first_seen_at=now, last_seen_at=now)
            db.add(row)
        else:
            if row.status == "resolved":
                row.status = "open"
                row.first_seen_at = now
                row.resolved_at = None
                row.acknowledged_at = None
                row.acknowledged_by = None
            row.severity = item["severity"]
            row.component = item["component"]
            row.title = item["title"]
            row.detail = item["detail"]
            row.metadata_json = json.dumps(item["metadata"], ensure_ascii=False, default=str)
            row.last_seen_at = now
    active_keys = set(conditions)
    for key, row in existing.items():
        if key not in active_keys and row.status in {"open","acknowledged"}:
            row.status = "resolved"
            row.resolved_at = now
            row.last_seen_at = now
    db.flush()
    rows = db.scalars(select(PilotAlert).order_by(PilotAlert.status.asc(), PilotAlert.severity.desc(), PilotAlert.last_seen_at.desc())).all()
    return {
        "slo":slo, "recovery":recovery, "active_conditions":sorted(active_keys),
        "alerts":[{
            "id":row.id,"alert_key":row.alert_key,"severity":row.severity,"status":row.status,"component":row.component,
            "title":row.title,"detail":row.detail,"metadata":json.loads(row.metadata_json or "{}"),
            "first_seen_at":row.first_seen_at.isoformat() if row.first_seen_at else None,
            "last_seen_at":row.last_seen_at.isoformat() if row.last_seen_at else None,
            "acknowledged_by":row.acknowledged_by,"acknowledged_at":row.acknowledged_at.isoformat() if row.acknowledged_at else None,
            "resolved_at":row.resolved_at.isoformat() if row.resolved_at else None,
        } for row in rows],
    }


@app.get("/api/admin/slo")
def admin_slo(user: User = Depends(require_roles("admin"))):
    return {"slo":slo_snapshot(),"recovery":_recovery_view(),"sla":{"mode":PILOT_SLA_MODE,"service_window":PILOT_SERVICE_WINDOW,"contractual":False}}


@app.get("/api/admin/alerts")
def admin_alerts(status: str | None = Query(default=None, pattern="^(open|acknowledged|resolved)$"), user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    data = evaluate_pilot_alerts(db)
    db.commit()
    rows = data["alerts"]
    if status:
        rows = [row for row in rows if row["status"] == status]
    return {"alerts":rows,"active_count":sum(1 for row in data["alerts"] if row["status"] in {"open","acknowledged"}),"recovery":data["recovery"],"slo":data["slo"]}


@app.patch("/api/admin/alerts/{alert_id}/ack")
def admin_alert_ack(alert_id: int, payload: AlertAckPayload, request: Request, user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    row = db.get(PilotAlert, alert_id)
    if not row:
        raise HTTPException(404, "Alert not found")
    if row.status == "resolved":
        return {"ok":True,"status":"resolved"}
    row.status = "acknowledged"
    row.acknowledged_by = user.id
    row.acknowledged_at = datetime.now(timezone.utc)
    if payload.note:
        metadata = json.loads(row.metadata_json or "{}")
        metadata["ack_note"] = payload.note
        row.metadata_json = json.dumps(metadata, ensure_ascii=False)
    audit_event(db, "pilot_alert.acknowledged", actor_user_id=user.id, target_type="pilot_alert", target_id=str(row.id), request=request, metadata={"alert_key":row.alert_key})
    db.commit()
    return {"ok":True,"status":row.status}


@app.get("/api/admin/learning-error-telemetry")
def learning_error_telemetry(days: int = Query(default=7, ge=1, le=90), user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    now = datetime.now(timezone.utc)
    since = now - timedelta(days=days)
    previous_since = since - timedelta(days=days)
    rows = db.scalars(select(PracticeResult).where(PracticeResult.created_at >= previous_since)).all()
    def aware(value: datetime) -> datetime:
        return value if value.tzinfo else value.replace(tzinfo=timezone.utc)
    current = [r for r in rows if aware(r.created_at) >= since]
    previous = [r for r in rows if previous_since <= aware(r.created_at) < since]

    def accuracy(items: list[PracticeResult]) -> float:
        total = sum(max(0, int(r.total or 0)) for r in items)
        correct = sum(max(0, int(r.score or 0)) for r in items)
        return round(correct * 100 / total, 1) if total else 0.0

    def aggregate(items: list[PracticeResult], key_fn) -> list[dict[str, Any]]:
        groups: dict[str, dict[str, Any]] = {}
        for row in items:
            key = key_fn(row) or "Без темы"
            item = groups.setdefault(key,{"name":key,"sessions":0,"questions":0,"correct":0,"errors":0})
            item["sessions"] += 1
            item["questions"] += max(0, int(row.total or 0))
            item["correct"] += max(0, int(row.score or 0))
            item["errors"] += max(0, int(row.total or 0) - int(row.score or 0))
        out=[]
        for item in groups.values():
            item["accuracy_percent"] = round(item["correct"] * 100 / max(1,item["questions"]),1)
            item["error_rate_percent"] = round(item["errors"] * 100 / max(1,item["questions"]),1)
            out.append(item)
        return sorted(out,key=lambda x:(-x["errors"],x["accuracy_percent"],x["name"]))

    by_topic = aggregate(current, lambda r:r.topic)
    by_kind = aggregate(current, lambda r:r.kind)
    by_language = aggregate(current, lambda r:r.language)
    return {
        "days":days,"sessions":len(current),"accuracy_percent":accuracy(current),"previous_accuracy_percent":accuracy(previous),
        "errors":sum(max(0,int(r.total or 0)-int(r.score or 0)) for r in current),
        "users_with_errors":len({r.user_id for r in current if int(r.score or 0) < int(r.total or 0)}),
        "weak_topics":[x for x in by_topic if x["questions"] >= 3][:10],
        "by_kind":by_kind[:10],"by_language":by_language,
        "note":"Aggregated learning telemetry is for improving content and pilot UX; it is not an HR performance rating.",
    }


@app.get("/api/admin/pilot-telemetry")
def pilot_telemetry(user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    now = datetime.now(timezone.utc)
    users = db.scalars(select(User)).all()
    profiles = {p.user_id:p for p in db.scalars(select(GamificationProfile)).all()}
    def active(days: int) -> int:
        count = 0
        for row in users:
            profile = profiles.get(row.id)
            if not profile or not profile.last_activity_at:
                continue
            activity = profile.last_activity_at if profile.last_activity_at.tzinfo else profile.last_activity_at.replace(tzinfo=timezone.utc)
            if now - activity <= timedelta(days=days):
                count += 1
        return count
    since = now - timedelta(days=7)
    practices = db.scalars(select(PracticeResult).where(PracticeResult.created_at >= since)).all()
    games = db.scalars(select(GameSession).where(GameSession.completed_at >= since)).all()
    nudges = db.scalars(select(LearningNudge).where(LearningNudge.created_at >= since)).all()
    accuracy = round(sum(row.score for row in practices) * 100 / max(1, sum(row.total for row in practices)), 1)
    by_language = {lang: sum(1 for row in practices if row.language == lang) for lang in LANGUAGES}
    tts_health = _tts_health()
    tts_cache = _synthesize_wav.cache_info()
    return {
        "users_total":len(users), "active_1d":active(1), "active_7d":active(7), "active_30d":active(30),
        "practice_sessions_7d":len(practices), "practice_accuracy_7d":accuracy, "games_completed_7d":len(games),
        "practice_by_language_7d":by_language, "nudges_created_7d":len(nudges), "nudges_unread_7d":sum(1 for row in nudges if not row.read),
        "server_tts_available":tts_health["server_available"], "tts_engine":tts_health["engine"],
        "tts_language_checks":tts_health["language_checks"],
        "tts_cache":{"hits":tts_cache.hits,"misses":tts_cache.misses,"size":tts_cache.currsize,"maxsize":tts_cache.maxsize,
                     "disk_enabled":TTS_DISK_CACHE_ENABLED,"writable":tts_health.get("cache_writable",False),
                     "persistence":TTS_CACHE_PERSISTENCE},
        "legacy_tts_get_enabled":TTS_LEGACY_GET_ENABLED,
        "operational_events_24h": int(db.scalar(select(func.count()).select_from(OperationalEvent).where(OperationalEvent.created_at >= now - timedelta(hours=24))) or 0),
        "operational_errors_24h": int(db.scalar(select(func.count()).select_from(OperationalEvent).where(OperationalEvent.created_at >= now - timedelta(hours=24), OperationalEvent.severity == "error")) or 0),
        "expired_sessions_pending_cleanup": int(db.scalar(select(func.count()).select_from(LoginSession).where(LoginSession.expires_at <= now)) or 0),
        "retention_days":{"audit":AUDIT_RETENTION_DAYS,"operational_events":OPERATIONAL_EVENT_RETENTION_DAYS,"nudges":NUDGE_RETENTION_DAYS,"resolved_alerts":ALERT_RETENTION_DAYS,"pilot_usage":PILOT_USAGE_RETENTION_DAYS},
        "departments":sorted({row.department for row in users}),
        "slo":slo_snapshot(), "recovery":_recovery_view(),
        "active_alerts":int(db.scalar(select(func.count()).select_from(PilotAlert).where(PilotAlert.status.in_(["open","acknowledged"]))) or 0),
    }


def _pilot_group_view(db: Session, group: PilotGroup) -> dict[str, Any]:
    members = db.scalars(select(User).join(PilotGroupMember, PilotGroupMember.user_id == User.id).where(PilotGroupMember.group_id == group.id).order_by(User.display_name)).all()
    features = {row.flag_key:bool(row.enabled) for row in db.scalars(select(PilotGroupFeature).where(PilotGroupFeature.group_id == group.id)).all()}
    assignments = db.scalars(select(PilotTrackAssignment).where(PilotTrackAssignment.group_id == group.id).order_by(PilotTrackAssignment.created_at.desc())).all()
    return {"id":group.id,"name":group.name,"department":group.department,"description":group.description,"status":group.status,"wave":group.wave,"active":_group_active(group),
            "starts_at":group.starts_at.isoformat() if group.starts_at else None,"ends_at":group.ends_at.isoformat() if group.ends_at else None,
            "members":[{"id":u.id,"username":u.username,"display_name":u.display_name,"department":u.department} for u in members],"features":features,
            "assignments":[{"id":a.id,"language":a.language,"track_name":a.track_name,"topic":a.topic,"target_level":a.target_level,"due_date":a.due_date} for a in assignments]}


@app.get("/api/admin/pilot/groups")
def admin_pilot_groups(user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    return [_pilot_group_view(db, g) for g in db.scalars(select(PilotGroup).order_by(PilotGroup.wave, PilotGroup.name)).all()]


@app.post("/api/admin/pilot/groups")
def admin_create_pilot_group(payload: PilotGroupPayload, request: Request, user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    if db.scalar(select(PilotGroup).where(func.lower(PilotGroup.name) == payload.name.strip().lower())):
        raise HTTPException(409, "Группа пилота уже существует")
    if payload.starts_at and payload.ends_at and payload.ends_at <= payload.starts_at:
        raise HTTPException(400, "Дата окончания волны должна быть позже даты начала")
    group = PilotGroup(name=payload.name.strip(), department=payload.department.strip(), description=payload.description.strip(), status=payload.status, wave=payload.wave, starts_at=payload.starts_at, ends_at=payload.ends_at, created_by=user.id)
    db.add(group); db.flush()
    audit_event(db, "pilot.group.create", actor_user_id=user.id, target_type="pilot_group", target_id=str(group.id), request=request, metadata={"name":group.name,"wave":group.wave,"status":group.status})
    db.commit(); db.refresh(group)
    return _pilot_group_view(db, group)


@app.patch("/api/admin/pilot/groups/{group_id}/status")
def admin_pilot_group_status(group_id: int, payload: PilotGroupStatusPayload, request: Request, user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    group = db.get(PilotGroup, group_id)
    if not group: raise HTTPException(404, "Группа не найдена")
    previous = group.status; group.status = payload.status
    audit_event(db, "pilot.group.status", actor_user_id=user.id, target_type="pilot_group", target_id=str(group.id), request=request, metadata={"from":previous,"to":group.status})
    db.commit(); return _pilot_group_view(db, group)


@app.patch("/api/admin/pilot/groups/{group_id}")
def admin_update_pilot_group(group_id: int, payload: PilotGroupUpdatePayload, request: Request, user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    group = db.get(PilotGroup, group_id)
    if not group:
        raise HTTPException(404, "Группа не найдена")
    before = {"name":group.name,"department":group.department,"status":group.status,"wave":group.wave,"starts_at":group.starts_at.isoformat() if group.starts_at else None,"ends_at":group.ends_at.isoformat() if group.ends_at else None}
    values = payload.model_dump(exclude_unset=True)
    if "name" in values and values["name"]:
        duplicate = db.scalar(select(PilotGroup).where(func.lower(PilotGroup.name) == values["name"].strip().lower(), PilotGroup.id != group_id))
        if duplicate:
            raise HTTPException(409, "Группа пилота с таким именем уже существует")
        group.name = values["name"].strip()
    if "department" in values and values["department"] is not None: group.department = values["department"].strip()
    if "description" in values and values["description"] is not None: group.description = values["description"].strip()
    if "status" in values and values["status"] is not None: group.status = values["status"]
    if "wave" in values and values["wave"] is not None: group.wave = values["wave"]
    if "starts_at" in values: group.starts_at = values["starts_at"]
    if "ends_at" in values: group.ends_at = values["ends_at"]
    if group.starts_at and group.ends_at and group.ends_at <= group.starts_at:
        raise HTTPException(400, "Дата окончания волны должна быть позже даты начала")
    after = {"name":group.name,"department":group.department,"status":group.status,"wave":group.wave,"starts_at":group.starts_at.isoformat() if group.starts_at else None,"ends_at":group.ends_at.isoformat() if group.ends_at else None}
    audit_event(db, "pilot.group.update", actor_user_id=user.id, target_type="pilot_group", target_id=str(group.id), request=request, metadata={"before":before,"after":after})
    db.commit()
    return _pilot_group_view(db, group)


@app.post("/api/admin/pilot/groups/{group_id}/members/{user_id}")
def admin_pilot_add_member(group_id: int, user_id: int, request: Request, user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    group = db.get(PilotGroup, group_id); target = db.get(User, user_id)
    if not group or not target: raise HTTPException(404, "Группа или пользователь не найдены")
    existing = db.scalar(select(PilotGroupMember).where(PilotGroupMember.group_id == group_id, PilotGroupMember.user_id == user_id))
    if not existing: db.add(PilotGroupMember(group_id=group_id, user_id=user_id))
    audit_event(db, "pilot.group.member.add", actor_user_id=user.id, target_type="user", target_id=str(user_id), request=request, metadata={"group_id":group_id})
    db.commit(); return _pilot_group_view(db, group)


@app.delete("/api/admin/pilot/groups/{group_id}/members/{user_id}")
def admin_pilot_remove_member(group_id: int, user_id: int, request: Request, user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    group = db.get(PilotGroup, group_id)
    if not group: raise HTTPException(404, "Группа не найдена")
    db.execute(delete(PilotGroupMember).where(PilotGroupMember.group_id == group_id, PilotGroupMember.user_id == user_id))
    audit_event(db, "pilot.group.member.remove", actor_user_id=user.id, target_type="user", target_id=str(user_id), request=request, metadata={"group_id":group_id})
    db.commit(); return {"ok":True}


@app.get("/api/admin/pilot/features")
def admin_pilot_features(user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    catalog = feature_flag_catalog(db)
    return [{"flag_key":k,**v} for k,v in sorted(catalog.items())]


@app.post("/api/admin/pilot/features")
def admin_create_feature(payload: FeatureFlagPayload, request: Request, user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    row = db.scalar(select(FeatureFlag).where(FeatureFlag.flag_key == payload.flag_key))
    if row:
        row.title=payload.title; row.description=payload.description; row.default_enabled=payload.default_enabled; row.updated_at=datetime.now(timezone.utc)
    else:
        row=FeatureFlag(flag_key=payload.flag_key,title=payload.title,description=payload.description,default_enabled=payload.default_enabled); db.add(row)
    audit_event(db,"pilot.feature.upsert",actor_user_id=user.id,target_type="feature_flag",target_id=payload.flag_key,request=request,metadata={"default_enabled":payload.default_enabled})
    db.commit(); return {"ok":True,"flag_key":payload.flag_key}


@app.put("/api/admin/pilot/groups/{group_id}/features/{flag_key}")
def admin_group_feature(group_id: int, flag_key: str, payload: PilotFeaturePayload, request: Request, user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    group=db.get(PilotGroup,group_id)
    if not group: raise HTTPException(404,"Группа не найдена")
    if flag_key not in feature_flag_catalog(db): raise HTTPException(404,"Feature flag не найден")
    row=db.scalar(select(PilotGroupFeature).where(PilotGroupFeature.group_id==group_id,PilotGroupFeature.flag_key==flag_key))
    if not row: row=PilotGroupFeature(group_id=group_id,flag_key=flag_key,enabled=payload.enabled); db.add(row)
    else: row.enabled=payload.enabled; row.updated_at=datetime.now(timezone.utc)
    audit_event(db,"pilot.feature.group_override",actor_user_id=user.id,target_type="pilot_group",target_id=str(group_id),request=request,metadata={"flag":flag_key,"enabled":payload.enabled})
    db.commit(); return {"ok":True,"flag_key":flag_key,"enabled":payload.enabled}


@app.post("/api/admin/pilot/groups/{group_id}/assignments")
def admin_group_assignment(group_id: int, payload: TrackAssignmentPayload, request: Request, user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    group=db.get(PilotGroup,group_id)
    if not group: raise HTTPException(404,"Группа не найдена")
    language=validate_language(payload.language)
    row=PilotTrackAssignment(group_id=group_id,language=language,track_name=payload.track_name.strip(),topic=payload.topic.strip(),target_level=payload.target_level,due_date=payload.due_date,created_by=user.id)
    db.add(row); db.flush()
    audit_event(db,"pilot.assignment.create",actor_user_id=user.id,target_type="pilot_group",target_id=str(group_id),request=request,metadata={"assignment_id":row.id,"language":language,"track_name":row.track_name})
    db.commit(); return {"ok":True,"id":row.id}


@app.delete("/api/admin/pilot/groups/{group_id}/assignments/{assignment_id}")
def admin_delete_group_assignment(group_id: int, assignment_id: int, request: Request, user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    group=db.get(PilotGroup,group_id)
    if not group: raise HTTPException(404,"Группа не найдена")
    row=db.scalar(select(PilotTrackAssignment).where(PilotTrackAssignment.id==assignment_id, PilotTrackAssignment.group_id==group_id))
    if not row: raise HTTPException(404,"Назначение не найдено")
    metadata={"assignment_id":row.id,"language":row.language,"track_name":row.track_name,"topic":row.topic,"target_level":row.target_level,"due_date":row.due_date}
    db.delete(row)
    audit_event(db,"pilot.assignment.delete",actor_user_id=user.id,target_type="pilot_group",target_id=str(group_id),request=request,metadata=metadata)
    db.commit(); return {"ok":True,"id":assignment_id}


@app.get("/api/admin/pilot/export.csv")
def admin_pilot_export_csv(group_id: int | None = Query(default=None), user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    query=select(User).order_by(User.display_name)
    selected_group = None
    if group_id is not None:
        selected_group=db.get(PilotGroup,group_id)
        if not selected_group: raise HTTPException(404,"Группа не найдена")
        query=query.join(PilotGroupMember,PilotGroupMember.user_id==User.id).where(PilotGroupMember.group_id==group_id)
    users=db.scalars(query.limit(PILOT_EXPORT_MAX_USERS)).all()
    out=io.StringIO(); writer=csv.writer(out)
    writer.writerow(["pilot_groups","waves","user_id","username","display_name","department","role","level","lifetime_xp","weekly_xp","english_known","english_touched","chinese_putonghua_known","chinese_putonghua_touched","terms_known_total","terms_touched_total","best_exam","last_activity_at"])
    for target in users:
        x=user_admin_stats(db,target)
        progress=db.scalars(select(TermProgress).where(TermProgress.user_id==target.id)).all()
        counts={}
        for lang in ("english","chinese"):
            rows=[r for r in progress if r.language==lang]
            counts[lang]=(sum(1 for r in rows if r.status=="known"),len(rows))
        memberships=user_pilot_groups(db,target.id)
        if selected_group is not None: memberships=[g for g in memberships if g.id==selected_group.id]
        group_names=" | ".join(g.name for g in memberships) or "UNASSIGNED"
        waves=" | ".join(str(g.wave) for g in memberships) or ""
        writer.writerow([group_names,waves,target.id,target.username,target.display_name,target.department,target.role,x["level"],x["lifetime_xp"],x["weekly_xp"],counts["english"][0],counts["english"][1],counts["chinese"][0],counts["chinese"][1],x["terms_known"],x["terms_touched"],x["best_exam"],x["last_activity_at"] or ""])
    data=out.getvalue().encode("utf-8-sig")
    filename=f"mgc_pilot_results_group_{group_id}.csv" if group_id is not None else "mgc_pilot_results_all.csv"
    return StreamingResponse(io.BytesIO(data),media_type="text/csv; charset=utf-8",headers={"Content-Disposition":f"attachment; filename={filename}","Cache-Control":"no-store"})


@app.get("/api/admin/pilot/governance-summary")
def admin_governance_summary(user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    groups=db.scalars(select(PilotGroup)).all(); members=db.scalars(select(PilotGroupMember)).all(); assignments=db.scalars(select(PilotTrackAssignment)).all()
    usage_today=db.scalars(select(PilotDailyUsage).where(PilotDailyUsage.day_key==_pilot_day_key())).all()
    return {"groups_total":len(groups),"groups_active":sum(1 for g in groups if _group_active(g)),"waves":sorted({g.wave for g in groups}),"memberships":len(members),"assignments":len(assignments),
            "usage_today":{"users":len(usage_today),"xp_awarded":sum(x.xp_awarded for x in usage_today),"game_starts":sum(x.game_starts for x in usage_today),"tts_requests":sum(x.tts_requests for x in usage_today),"practice_submissions":sum(x.practice_submissions for x in usage_today)},
            "quotas":{"daily_xp":PILOT_DAILY_XP_CAP,"daily_games":PILOT_DAILY_GAME_START_CAP,"daily_tts":PILOT_DAILY_TTS_CAP,"daily_practice":PILOT_DAILY_PRACTICE_CAP},
            "chinese_standard":"Путунхуа (普通话) — единственный основной китайский учебный трек; диалекты — справка."}


@app.get("/api/admin/analytics")
def admin_analytics(user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    users = db.scalars(select(User)).all()
    stats = [user_admin_stats(db, row) for row in users]
    now = datetime.now(timezone.utc)
    active_7d = 0
    for row in stats:
        if row["last_activity_at"]:
            dt = datetime.fromisoformat(row["last_activity_at"])
            if dt.tzinfo is None: dt = dt.replace(tzinfo=timezone.utc)
            if now - dt <= timedelta(days=7): active_7d += 1
    db.commit()
    return {
        "users": len(stats), "active_7d": active_7d,
        "lifetime_xp": sum(x["lifetime_xp"] for x in stats), "weekly_xp": sum(x["weekly_xp"] for x in stats),
        "terms_known": sum(x["terms_known"] for x in stats),
    }


@app.get("/api/admin/audit")
def admin_audit(
    limit: int = Query(default=100, ge=1, le=500), event_type: str | None = None,
    user: User = Depends(require_roles("admin")), db: Session = Depends(db_session),
):
    query = select(AuditLog)
    if event_type:
        query = query.where(AuditLog.event_type == event_type)
    rows = db.scalars(query.order_by(AuditLog.created_at.desc()).limit(limit)).all()
    return [{
        "id": row.id, "actor_user_id": row.actor_user_id, "event_type": row.event_type,
        "target_type": row.target_type, "target_id": row.target_id,
        "request_id": row.request_id, "metadata": json.loads(row.metadata_json or "{}"),
        "created_at": row.created_at.isoformat(),
    } for row in rows]


def _pool_value(name: str) -> Any:
    value = getattr(engine.pool, name, None)
    try:
        return value() if callable(value) else value
    except Exception:
        return None


@app.get("/api/admin/database/telemetry")
def admin_database_telemetry(user: User = Depends(require_roles("admin"))):
    return {"queries": db_query_telemetry_snapshot(), "pool": db_pool_snapshot()}


@app.get("/api/admin/recovery/evidence")
def admin_recovery_evidence(user: User = Depends(require_roles("admin"))):
    return recovery_evidence_snapshot()


@app.get("/api/admin/it-dashboard")
def admin_it_dashboard(user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    now = datetime.now(timezone.utc)
    ready_ok, ready_checks = readiness_checks(db)
    since_24h = now - timedelta(hours=24)
    recent = db.scalars(select(OperationalEvent).where(OperationalEvent.created_at >= since_24h).order_by(OperationalEvent.created_at.desc()).limit(100)).all()
    by_component = Counter(row.component for row in recent)
    by_type = Counter(row.event_type for row in recent)
    with _METRIC_LOCK:
        total_requests = sum(_METRIC_REQUESTS.values())
        server_errors = sum(count for (method, path, status), count in _METRIC_REQUESTS.items() if int(status) >= 500)
    with _RELIABILITY_LOCK:
        reliability_runtime = dict(_RELIABILITY_RUNTIME)
    tts = _tts_health()
    maintenance_preview = cleanup_runtime_data(db, dry_run=True)
    operations = evaluate_pilot_alerts(db)
    db.commit()
    active_alerts = [row for row in operations["alerts"] if row["status"] in {"open","acknowledged"}]
    return {
        "status":"healthy" if ready_ok and operations["recovery"].get("state") == "healthy" else "attention", "version":APP_VERSION, "environment":APP_ENV, "instance_id":INSTANCE_ID,
        "readiness":{"ready":ready_ok,"checks":ready_checks},
        "recovery":operations["recovery"], "slo":operations["slo"],
        "sla":{"mode":PILOT_SLA_MODE,"service_window":PILOT_SERVICE_WINDOW,"contractual":False},
        "alerts":{"active_count":len(active_alerts),"items":active_alerts[:20]},
        "database":{"driver":engine.dialect.name,"pool_size":_pool_value("size"),"checked_out":_pool_value("checkedout"),"pool":db_pool_snapshot(),"query_telemetry":db_query_telemetry_snapshot()},
        "recovery_evidence":recovery_evidence_snapshot(),
        "rpo_rto":{"rpo_target_minutes":RPO_TARGET_MINUTES,"rto_target_minutes":RTO_TARGET_MINUTES,"note":"RPO is estimated from current backup evidence; RTO requires a timed restore/failover drill."},
        "tts":{"server_available":tts.get("server_available",False),"engine":tts.get("engine",""),"circuit_open":_tts_circuit_open(),"runtime":dict(_TTS_RUNTIME)},
        "http":{"requests_since_start":total_requests,"server_errors_since_start":server_errors},
        "reliability_runtime":reliability_runtime,
        "events_24h":{"total":len(recent),"errors":sum(1 for x in recent if x.severity == "error"),"by_component":dict(by_component),"by_type":dict(by_type)},
        "maintenance":maintenance_preview,
        "recent_events":[{"event_type":x.event_type,"severity":x.severity,"component":x.component,"status_code":x.status_code,"request_id":x.request_id,"path":x.path,"detail":x.detail,"created_at":x.created_at.isoformat()} for x in recent[:20]],
    }


@app.post("/api/admin/maintenance/cleanup")
def admin_maintenance_cleanup(request: Request, dry_run: bool = Query(default=True), user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    result = cleanup_runtime_data(db, dry_run=dry_run)
    audit_event(db, "maintenance.cleanup", actor_user_id=user.id, target_type="runtime_data", request=request, metadata=result)
    db.commit()
    return result


@app.get("/api/admin/operational-events")
def admin_operational_events(limit: int = Query(default=100, ge=1, le=500), severity: str | None = None, component: str | None = None, user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    query = select(OperationalEvent)
    if severity:
        query = query.where(OperationalEvent.severity == severity)
    if component:
        query = query.where(OperationalEvent.component == component)
    rows = db.scalars(query.order_by(OperationalEvent.created_at.desc()).limit(limit)).all()
    return [{"id":x.id,"event_type":x.event_type,"severity":x.severity,"component":x.component,"status_code":x.status_code,"request_id":x.request_id,"path":x.path,"detail":x.detail,"metadata":json.loads(x.metadata_json or "{}"),"created_at":x.created_at.isoformat()} for x in rows]


@app.get("/api/admin/system/summary")
def admin_system_summary(user: User = Depends(require_roles("admin")), db: Session = Depends(db_session)):
    ok, checks = readiness_checks(db)
    findings = []
    if APP_ENV in {"pilot", "production"} and not COOKIE_SECURE:
        findings.append("COOKIE_SECURE=false: допустимо только за доверенным TLS reverse proxy на sandbox; включить при прямом HTTPS")
    if AUTH_MODE == "local":
        findings.append("AUTH_MODE=local: для корпоративного rollout рекомендуется OIDC")
    if DATABASE_URL.startswith("sqlite"):
        findings.append("SQLite используется только для dev; pilot profile требует PostgreSQL")
    if not TRUSTED_HOSTS:
        findings.append("TRUSTED_HOSTS не ограничен")
    if APP_ENV in {"pilot", "production"} and TTS_LEGACY_GET_ENABLED:
        findings.append("TTS_LEGACY_GET_ENABLED=true: отключите GET-аудио, чтобы учебный текст не попадал в URL/access-log")
    if APP_ENV in {"pilot", "production"} and TTS_CACHE_PERSISTENCE != "ephemeral":
        findings.append("TTS cache persistent: для privacy-first пилота рекомендуется ephemeral cache")
    return {
        "version":APP_VERSION, "environment":APP_ENV, "auth_mode":AUTH_MODE,
        "database_driver":engine.dialect.name, "auto_create_schema":AUTO_CREATE_SCHEMA,
        "csrf_enabled":True, "cookie_secure":COOKIE_SECURE,
        "registration_enabled":AUTH_MODE == "local" and REGISTRATION_ENABLED,
        "trusted_hosts":TRUSTED_HOSTS, "cors_origins_configured":len(CORS_ORIGINS),
        "metrics_enabled":METRICS_ENABLED, "metrics_token_configured":bool(METRICS_TOKEN),
        "rls_enabled":RLS_ENABLED, "otel":OTEL_STATUS, "term_approval_required":TERM_APPROVAL_REQUIRED,
        "database_observability":{"pool_alert_percent":DB_POOL_ALERT_PERCENT,"slow_query_threshold_ms":DB_SLOW_QUERY_THRESHOLD_MS,"query_window_minutes":DB_QUERY_WINDOW_MINUTES},
        "recovery_evidence":recovery_evidence_snapshot(),
        "rpo_rto":{"rpo_target_minutes":RPO_TARGET_MINUTES,"rto_target_minutes":RTO_TARGET_MINUTES,"recovery_evidence_required":RECOVERY_EVIDENCE_REQUIRED},
        "server_tts_available":_tts_health()["server_available"], "tts_engine":_tts_health()["engine"],
        "oidc_department_claim":OIDC_DEPARTMENT_CLAIM, "voice_recording_enabled":False, "pronunciation_transport":"POST",
        "tts_circuit_open":_tts_circuit_open(), "tts_disk_cache_enabled":TTS_DISK_CACHE_ENABLED,
        "tts_cache_persistence":TTS_CACHE_PERSISTENCE, "legacy_tts_get_enabled":TTS_LEGACY_GET_ENABLED,
        "db_connect_timeout_seconds":DB_CONNECT_TIMEOUT_SECONDS, "db_statement_timeout_ms":DB_STATEMENT_TIMEOUT_MS,
        "retention_days":{"audit":AUDIT_RETENTION_DAYS,"operational_events":OPERATIONAL_EVENT_RETENTION_DAYS,"nudges":NUDGE_RETENTION_DAYS},
        "pilot_governance":{"daily_xp_cap":PILOT_DAILY_XP_CAP,"daily_game_cap":PILOT_DAILY_GAME_START_CAP,"daily_tts_cap":PILOT_DAILY_TTS_CAP,"daily_practice_cap":PILOT_DAILY_PRACTICE_CAP,
                            "groups_total":int(db.scalar(select(func.count()).select_from(PilotGroup)) or 0),
                            "groups_active":sum(1 for g in db.scalars(select(PilotGroup)).all() if _group_active(g)),
                            "feature_flags":len(feature_flag_catalog(db))},
        "chinese_learning_standard":"Путунхуа (普通话) — стандартный китайский; диалекты справочно",
        "ready":ok, "checks":checks, "configuration_findings":findings,
    }


app.mount("/", StaticFiles(directory=STATIC_DIR, html=True), name="static")
